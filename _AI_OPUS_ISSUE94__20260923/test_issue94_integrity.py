"""Issue #94 fail-closed integrity regressions (mandatory cases 1-4).

Synthetic fixtures only. Every workbook and document below is generated in a
pytest tmp dir; nothing reads client, P10, or report bytes.

  1. <c t="e"><f>...</f><v>#REF!</v></c> never emerges as a non-error literal
     and never converges/promotes.
  2. Malformed worksheet XML is a hard defect: zero output promotion, and no
     row/cell loss is accepted (including from an injected lossy fixer).
  3. Effective/execution date without a recording label -> blank
     recording_date plus a review-required row.
  4. 0.5 + 0.5 parses and reconciles to 1.0; incomplete owner sets never
     assert a sum (true or false).

Copy to tests/ unchanged; the repo root is resolved from this file's parent.
"""

from __future__ import annotations

import csv
import re
import sys
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pytest
from lxml import etree

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import grocery_report_pipeline as grp  # noqa: E402
from horizon.audit import AuditLog  # noqa: E402
from horizon.config import HorizonConfig  # noqa: E402
from horizon.models import ReportModel, TitleRow  # noqa: E402
from horizon.orchestrator import Orchestrator  # noqa: E402
from horizon.repair import repair_workbook, restore_formula_from_template  # noqa: E402
from horizon.validation import Requirements  # noqa: E402

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
EXCEL_ERROR_TOKENS = (
    "#REF!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#VALUE!",
    "#SPILL!", "#CALC!", "#GETTING_DATA", "#FIELD!", "#BLOCKED!", "#UNKNOWN!",
)
MEDIA_PART = "xl/media/plat1.png"
MEDIA_BYTES = b"\x89PNG\r\n\x1a\nSYNTHETIC-ISSUE94-PLAT"
BASE = "31-12N-24W_report"

# Phrase that proves an owner set is complete. If production adopts a
# different proof marker, change it here and in make_synthetic_corpus.
COMPLETE_OWNER_SET_MARKER = "Complete owner set"


# ---------------------------------------------------------------------------
# Synthetic workbook builders
# ---------------------------------------------------------------------------
def _sheet_xml(sheet_data_inner: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{MAIN_NS}"><sheetData>{sheet_data_inner}</sheetData>'
        '<pageMargins left="0.75" right="0.75" top="1" bottom="1" header="0.5" '
        'footer="0.5"/></worksheet>'
    ).encode("utf-8")


def _rows(middle_cell: str) -> str:
    return (
        '<row r="1"><c r="A1" t="inlineStr"><is><t>SYNTHETIC</t></is></c></row>'
        f'<row r="2">{middle_cell}</row>'
        '<row r="3"><c r="A3" t="n"><v>3</v></c></row>'
    )


def _build_xlsx(path: Path, sheet_xml: bytes, media: bool = True) -> Path:
    """openpyxl-valid package whose sheet1.xml is replaced by ``sheet_xml``."""
    import openpyxl

    scratch = path.with_suffix(".scratch.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "SYNTHETIC"
    wb.save(scratch)
    wb.close()
    with zipfile.ZipFile(scratch) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = sheet_xml
            zout.writestr(item, data)
        if media:
            zout.writestr(MEDIA_PART, MEDIA_BYTES)
    scratch.unlink()
    return path


CellMap = Dict[Tuple[str, str], Tuple[Optional[str], Optional[str], Optional[str]]]


def _cells(xlsx: Path) -> CellMap:
    """(part, ref) -> (t, formula, cached). Strict parse: a malformed output
    part raises instead of being tolerated."""
    strict = etree.XMLParser(recover=False)
    out: CellMap = {}
    with zipfile.ZipFile(xlsx) as zf:
        for name in zf.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(zf.read(name), parser=strict)
            for c in root.iter(f"{{{MAIN_NS}}}c"):
                f = c.find(f"{{{MAIN_NS}}}f")
                v = c.find(f"{{{MAIN_NS}}}v")
                out[(name, c.get("r", "?"))] = (
                    c.get("t"),
                    None if f is None else (f.text or ""),
                    None if v is None else (v.text or ""),
                )
    return out


def _downgraded_error_literals(xlsx: Path) -> List[str]:
    """Cells whose cached value is an Excel error but are not typed t="e"."""
    bad = []
    for (part, ref), (t, _f, v) in _cells(xlsx).items():
        if v and v.strip().upper() in EXCEL_ERROR_TOKENS and t != "e":
            bad.append(f"{part}!{ref} t={t!r} v={v!r}")
    return bad


def _error_cells(xlsx: Path) -> List[str]:
    hits = []
    for (part, ref), (t, f, v) in _cells(xlsx).items():
        if t == "e" or (f or "").lstrip("=").startswith("#") or \
                (v or "").strip().upper() in EXCEL_ERROR_TOKENS:
            hits.append(f"{part}!{ref}")
    return hits


def _assert_refused(result, src: Path, src_bytes: bytes, dest: Path) -> None:
    assert src.read_bytes() == src_bytes, "source workbook bytes changed"
    assert result.output is None, f"defective workbook produced output {result.output}"
    assert result.repaired is False
    assert result.error, "refusal must carry a defect/review receipt (error text)"
    assert getattr(result, "promoted", False) is False
    assert not dest.exists(), "defective candidate left on disk (could be promoted)"
    siblings = dest.parent.iterdir() if dest.parent.exists() else []
    leftovers = [p.name for p in siblings
                 if p != src and p.suffix in {".xlsx", ".repairing", ".tmp"}]
    assert not leftovers, f"partial/temporary outputs left behind: {leftovers}"


# ---------------------------------------------------------------------------
# Regression 1 -- error formula never becomes a literal, never promotes
# ---------------------------------------------------------------------------
ERROR_CELL_VARIANTS = {
    "issue_exact_ref": '<c r="A2" t="e"><f>#REF!+1</f><v>#REF!</v></c>',
    "valid_body_ref_value": '<c r="A2" t="e"><f>Sheet9!A1*2</f><v>#REF!</v></c>',
    "div0": '<c r="A2" t="e"><f>1/0</f><v>#DIV/0!</v></c>',
    "shared_child": '<c r="A2" t="e"><f t="shared" si="0"/><v>#REF!</v></c>',
    "openpyxl_style_no_type": '<c r="A2"><f>#REF!</f><v></v></c>',
    "cached_error_type_missing": '<c r="A2"><f>A1*2</f><v>#REF!</v></c>',
}


@pytest.mark.parametrize("cell_xml", ERROR_CELL_VARIANTS.values(), ids=ERROR_CELL_VARIANTS.keys())
def test_error_formula_repair_is_refused_not_downgraded(tmp_path, cell_xml):
    src = _build_xlsx(tmp_path / "synthetic_src.xlsx", _sheet_xml(_rows(cell_xml)))
    src_bytes = src.read_bytes()
    dest = tmp_path / "out" / "synthetic_v002.xlsx"

    result = repair_workbook(src, dest)

    _assert_refused(result, src, src_bytes, dest)
    for produced in tmp_path.rglob("*.xlsx"):
        if produced != src:
            assert not _downgraded_error_literals(produced), produced


def test_template_restore_leaves_no_cached_error_literal(tmp_path):
    """The only approved non-refusal path: exact template formula, no cached
    value, so native recalculation is required before any value is trusted."""
    staged = _build_xlsx(tmp_path / "staged.xlsx",
                         _sheet_xml(_rows(ERROR_CELL_VARIANTS["issue_exact_ref"])))
    template = _build_xlsx(tmp_path / "template.xlsx",
                           _sheet_xml(_rows('<c r="A2"><f>A3*2</f><v>6</v></c>')))

    result = restore_formula_from_template(staged, template, "Sheet", "A2")

    assert result.repaired and not result.error, result.error
    t, formula, cached = _cells(staged)[("xl/worksheets/sheet1.xml", "A2")]
    assert formula == "A3*2"
    assert cached is None, "restored formula must not carry any cached value"
    assert t != "e" and not _downgraded_error_literals(staged)
    with zipfile.ZipFile(staged) as zf:
        assert zf.read(MEDIA_PART) == MEDIA_BYTES


@pytest.fixture()
def horizon_cfg(tmp_path):
    cfg = HorizonConfig(root=tmp_path / "Horizon")
    cfg.ensure_workspace()
    return cfg


def _versions(cfg) -> List[Path]:
    return sorted(cfg.final_reports.glob(f"{BASE}_v*.xlsx"))


def test_errored_workbook_failing_validation_never_promotes(horizon_cfg):
    seed = _build_xlsx(horizon_cfg.final_reports / f"{BASE}_v001.xlsx",
                       _sheet_xml(_rows(ERROR_CELL_VARIANTS["issue_exact_ref"])))
    seed_bytes = seed.read_bytes()
    report = ReportModel(section="31-12N-24W", rows=[
        TitleRow(grantor="A", grantee="B", instrument_number="100")])
    audit = AuditLog(path=None, echo=False)

    result = Orchestrator(horizon_cfg, audit).run(
        report, BASE, Requirements(required_instruments={"999"}))

    assert not result.converged and result.final_version is None
    assert _versions(horizon_cfg) == [seed], "a defective workbook was versioned"
    assert seed.read_bytes() == seed_bytes
    assert audit.count("ESCALATED") >= 1


def test_errored_workbook_cannot_converge_even_if_model_validates(horizon_cfg):
    """emit_version() templates on the ingested workbook (to keep plats), so a
    passing in-memory model would otherwise carry the #REF! sheet forward."""
    seed = _build_xlsx(horizon_cfg.final_reports / f"{BASE}_v001.xlsx",
                       _sheet_xml(_rows(ERROR_CELL_VARIANTS["issue_exact_ref"])))
    report = ReportModel(section="31-12N-24W", rows=[
        TitleRow(grantor="A", grantee="B", instrument_number="100",
                 conveyed_interest="1/2", retained_interest="1/2")])

    result = Orchestrator(horizon_cfg, AuditLog(path=None, echo=False)).run(
        report, BASE, Requirements())

    assert not result.converged, "converged on a workbook carrying an Excel error"
    assert result.final_version is None
    assert _versions(horizon_cfg) == [seed]


# ---------------------------------------------------------------------------
# Regression 2 -- malformed worksheet XML is a hard defect; no cell loss
# ---------------------------------------------------------------------------
_VALID_ROWS = _rows('<c r="A2" t="n"><v>2</v></c>')
MALFORMED_SHEETS = {
    "unclosed_row": _sheet_xml(
        '<row r="1"><c r="A1" t="n"><v>1</v></c>'
        '<row r="2"><c r="A2" t="n"><v>2</v></c></row>'
        '<row r="3"><c r="A3" t="n"><v>3</v></c></row>'),
    "mismatched_close": _sheet_xml(_rows('<c r="A2" t="n"><v>2</v></x>')),
    "bare_ampersand": _sheet_xml(_rows('<c r="A2" t="inlineStr"><is><t>AT&T</t></is></c>')),
    "undefined_entity": _sheet_xml(_rows('<c r="A2" t="inlineStr"><is><t>a&nbsp;b</t></is></c>')),
    "duplicate_attribute": _sheet_xml(_rows('<c r="A2" r="A9" t="n"><v>2</v></c>')),
    "truncated_part": _sheet_xml(_VALID_ROWS)[: len(_sheet_xml(_VALID_ROWS)) * 2 // 3],
}


@pytest.mark.parametrize("sheet_xml", MALFORMED_SHEETS.values(), ids=MALFORMED_SHEETS.keys())
def test_malformed_worksheet_is_hard_defect_with_zero_promotion(tmp_path, sheet_xml):
    src = _build_xlsx(tmp_path / "synthetic_src.xlsx", sheet_xml)
    src_bytes = src.read_bytes()
    dest = tmp_path / "out" / "synthetic_v002.xlsx"

    result = repair_workbook(src, dest)

    _assert_refused(result, src, src_bytes, dest)


def _drop_last_row(xml: bytes, fixes: List[str]) -> bytes:
    root = etree.fromstring(xml)
    sheet_data = root.find(f"{{{MAIN_NS}}}sheetData")
    sheet_data.remove(sheet_data[-1])
    fixes.append("normalized sheet")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _drop_one_cell(xml: bytes, fixes: List[str]) -> bytes:
    root = etree.fromstring(xml)
    cell = root.find(f".//{{{MAIN_NS}}}c[@r='A2']")
    cell.getparent().remove(cell)
    fixes.append("normalized cell")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


@pytest.mark.parametrize("fixer", [_drop_last_row, _drop_one_cell], ids=["drop_row", "drop_cell"])
def test_repair_rejects_any_row_or_cell_loss(tmp_path, fixer):
    src = _build_xlsx(tmp_path / "synthetic_src.xlsx", _sheet_xml(_VALID_ROWS))
    src_bytes = src.read_bytes()
    dest = tmp_path / "out" / "synthetic_v002.xlsx"

    result = repair_workbook(src, dest, worksheet_fixer=fixer)

    _assert_refused(result, src, src_bytes, dest)


def test_clean_workbook_repair_preserves_every_cell_and_media(tmp_path):
    src = _build_xlsx(tmp_path / "synthetic_src.xlsx", _sheet_xml(_VALID_ROWS))
    dest = tmp_path / "out" / "synthetic_v002.xlsx"

    result = repair_workbook(src, dest)

    assert not result.error, result.error
    if result.output is not None:
        assert _cells(result.output) == _cells(src)
        with zipfile.ZipFile(result.output) as zf:
            assert zf.read(MEDIA_PART) == MEDIA_BYTES


def test_malformed_workbook_never_versioned_by_orchestrator(horizon_cfg):
    seed = _build_xlsx(horizon_cfg.final_reports / f"{BASE}_v001.xlsx",
                       MALFORMED_SHEETS["unclosed_row"])
    seed_bytes = seed.read_bytes()
    report = ReportModel(section="31-12N-24W", rows=[
        TitleRow(grantor="A", grantee="B", instrument_number="100")])

    result = Orchestrator(horizon_cfg, AuditLog(path=None, echo=False)).run(
        report, BASE, Requirements(required_instruments={"999"}))

    assert not result.converged and result.final_version is None
    assert _versions(horizon_cfg) == [seed]
    assert seed.read_bytes() == seed_bytes


def test_template_restore_refuses_malformed_candidate(tmp_path):
    staged = _build_xlsx(tmp_path / "staged.xlsx", MALFORMED_SHEETS["unclosed_row"])
    staged_bytes = staged.read_bytes()
    template = _build_xlsx(tmp_path / "template.xlsx",
                           _sheet_xml(_rows('<c r="A2"><f>A3*2</f></c>')))

    result = restore_formula_from_template(staged, template, "Sheet", "A2")

    assert result.output is None and not result.repaired and result.error
    assert staged.read_bytes() == staged_bytes
    assert not list(tmp_path.glob("*.repairing"))


# ---------------------------------------------------------------------------
# Regressions 3 and 4 -- grocery pipeline on a synthetic corpus
# ---------------------------------------------------------------------------
_HDR = "SYNTHETIC TEST DOCUMENT -- NOT REAL TITLE DATA\n"
CORPUS = {
    # Regression 3: recording date must never be inferred.
    "r1_effective_execution_only.txt": (
        _HDR + "MINERAL DEED\nGrantor: Alpha Synthetic\nGrantee: Beta Synthetic\n"
        "Effective Date: 2019-03-15\nExecuted: 2019-03-10\n"
        "Legal: Section 31, T7N, R63W\n"),
    "r2_instrument_but_no_recording_label.txt": (
        _HDR + "MINERAL DEED\nGrantor: Gamma Synthetic\nGrantee: Delta Synthetic\n"
        "Effective Date: 2019-05-01\nInstrument No 20190001234\n"
        "Notary: my commission expires 2027-01-01\n"
        "Legal: Section 32, T7N, R63W\n"),
    "r3_unrecorded_substring_trap.txt": (
        _HDR + "MEMORANDUM\nLessor: Epsilon Synthetic\nLessee: Zeta Synthetic\n"
        "Memorandum of unrecorded lease dated 2018-07-04.\n"
        "Legal: Section 33, T7N, R63W\n"),
    "r4_labeled_recording_control.txt": (
        _HDR + "MINERAL DEED\nGrantor: Eta Synthetic\nGrantee: Theta Synthetic\n"
        "Effective Date: 2020-01-15\nRecorded: 2020-02-02 Book 9 Page 9\n"
        "Legal: Section 34, T7N, R63W\n"),
    # Regression 4: decimal precision and owner-set completeness.
    "d1_half_half_complete.txt": (
        _HDR + f"DIVISION ORDER OWNER SCHEDULE -- {COMPLETE_OWNER_SET_MARKER}\n"
        "Owner Alpha Synthetic decimal interest 0.5\n"
        "Owner Beta Synthetic decimal interest 0.5\n"
        "Legal: Section 22, T7N, R63W\n"),
    "d2_quarter_eighth_complete.txt": (
        _HDR + f"DIVISION ORDER OWNER SCHEDULE -- {COMPLETE_OWNER_SET_MARKER}\n"
        "Owner A decimal interest 0.25\nOwner B decimal interest 0.125\n"
        "Owner C decimal interest .125\nOwner D decimal interest 0.50000000\n"
        "Legal: Section 23, T7N, R63W\n"),
    "d3_partial_owner_set.txt": (
        _HDR + "PARTIAL owner schedule\n"
        "Owner Alpha Synthetic decimal interest 0.5\n"
        "Legal: Section 24, T7N, R63W\n"),
    "d4_ownership_excerpt_unproved.txt": (
        _HDR + "OWNERSHIP schedule (excerpt from page 2 of 3)\n"
        "Owner Alpha Synthetic decimal interest 0.5\n"
        "Owner Beta Synthetic decimal interest 0.25\n"
        "Legal: Section 25, T7N, R63W\n"),
    "d5_complete_but_short.txt": (
        _HDR + f"DIVISION ORDER OWNER SCHEDULE -- {COMPLETE_OWNER_SET_MARKER}\n"
        "Owner Alpha Synthetic decimal interest 0.5\n"
        "Owner Beta Synthetic decimal interest 0.25\n"
        "Legal: Section 26, T7N, R63W\n"),
    "d6_complete_schedule.txt": (
        _HDR + f"DIVISION ORDER OWNER SCHEDULE -- {COMPLETE_OWNER_SET_MARKER}\n"
        "Owner Alpha Synthetic decimal interest 0.5\n"
        "Owner Beta Synthetic decimal interest 0.5\n"
        "Legal: Section 27, T7N, R63W\n"),
}
CORPUS["d6_complete_schedule_COPY.txt"] = CORPUS["d6_complete_schedule.txt"]


def _read_csv(path: Path) -> List[Dict[str, str]]:
    with path.open(encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


@pytest.fixture(scope="module")
def pipeline_out(tmp_path_factory) -> Path:
    base = tmp_path_factory.mktemp("issue94_grocery")
    corpus = base / "corpus"
    corpus.mkdir()
    for name, body in CORPUS.items():
        (corpus / name).write_text(body, encoding="utf-8")
    out = base / "output"
    grp.run_pipeline(corpus, out, "Issue94_Synthetic", apply_quar=False, log=grp.BuildLog())
    return out


@pytest.fixture(scope="module")
def facts(pipeline_out) -> Dict[str, Dict[str, str]]:
    return {r["source_file"]: r for r in _read_csv(pipeline_out / "extracted_facts.csv")}


@pytest.fixture(scope="module")
def review_rows(pipeline_out) -> List[Dict[str, str]]:
    return _read_csv(pipeline_out / "review_required.csv")


@pytest.fixture(scope="module")
def decimal_calc(pipeline_out) -> Dict[str, Dict[str, object]]:
    import openpyxl

    wb = openpyxl.load_workbook(pipeline_out / "reconciliation_table.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["acreage_decimal_calc"].iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in rows[0]]
    return {str(r[0]).upper(): dict(zip(header, r)) for r in rows[1:] if r and r[0]}


def _section(section: int) -> str:
    return f"SECTION {section}, T7N, R63W"


def _recording_review(review_rows, source_file: str) -> List[Dict[str, str]]:
    return [r for r in review_rows
            if "recording" in r["rule"].lower()
            and source_file in (r["subject"], r["source"])]


def _decimal_sum_issues(review_rows, section: int) -> List[Dict[str, str]]:
    return [r for r in review_rows
            if r["rule"] == "decimal-sum" and r["subject"].upper() == _section(section)]


RECORDING_DATE_UNPROVED = [
    "r1_effective_execution_only.txt",
    "r2_instrument_but_no_recording_label.txt",
    "r3_unrecorded_substring_trap.txt",
]


@pytest.mark.parametrize("source_file", RECORDING_DATE_UNPROVED)
def test_recording_date_blank_without_recording_label(facts, source_file):
    row = facts[source_file]
    assert row["recording_date"] == "", (
        f"recording_date fabricated as {row['recording_date']!r}")


@pytest.mark.parametrize("source_file", RECORDING_DATE_UNPROVED)
def test_missing_recording_date_is_review_required(review_rows, source_file):
    assert _recording_review(review_rows, source_file), (
        f"{source_file} has no recording evidence but no recording review row")


def test_effective_and_execution_dates_are_kept_in_their_own_fields(facts):
    row = facts["r1_effective_execution_only.txt"]
    assert row["effective_date"] == "2019-03-15"
    assert row["execution_date"] == "2019-03-10"


def test_labeled_recording_date_is_still_captured(facts, review_rows):
    row = facts["r4_labeled_recording_control.txt"]
    assert row["recording_date"] == "2020-02-02"
    assert not _recording_review(review_rows, "r4_labeled_recording_control.txt")


@pytest.mark.parametrize("text, expected", [
    ("decimal interest 0.5", 0.5),
    ("Decimal Interest: 0.5", 0.5),
    ("decimal of 0.25", 0.25),
    ("decimal interest 0.125", 0.125),
    ("decimal interest .125", 0.125),
    ("decimal interest 0.12500000", 0.125),
    ("decimal interest 0.00390625", 0.00390625),
    ("decimal interest 1", 1.0),
    ("decimal interest 1.00000000", 1.0),
])
def test_decimal_parses_any_valid_precision(text, expected):
    assert [float(x) for x in grp._DECIMAL_RX.findall(text)] == [expected]


@pytest.mark.parametrize("text", [
    "decimal interest 1.5",
    "decimal interest 10.5",
    "decimal interest 12",
])
def test_out_of_range_decimal_is_not_coerced_to_one(text):
    values = [float(x) for x in grp._DECIMAL_RX.findall(text)]
    assert 1.0 not in values, f"{text!r} was silently read as 1.0: {values}"


def test_half_plus_half_extracted_and_reconciles_to_one(facts, decimal_calc, review_rows):
    assert facts["d1_half_half_complete.txt"]["decimal_interest"] in {"0.5", ".5"}
    row = decimal_calc[_section(22)]
    assert float(row["decimal_sum"]) == pytest.approx(1.0, abs=1e-9)
    assert str(row["decimal_check"]).upper().startswith("OK"), row
    assert not _decimal_sum_issues(review_rows, 22)


def test_mixed_precision_complete_set_reconciles_to_one(decimal_calc, review_rows):
    row = decimal_calc[_section(23)]
    assert float(row["decimal_sum"]) == pytest.approx(1.0, abs=1e-9)
    assert str(row["decimal_check"]).upper().startswith("OK"), row
    assert not _decimal_sum_issues(review_rows, 23)


@pytest.mark.parametrize("section", [24, 25], ids=["partial_marker", "ownership_excerpt"])
def test_incomplete_owner_set_asserts_no_sum(decimal_calc, review_rows, section):
    check = str(decimal_calc[_section(section)]["decimal_check"])
    assert not check.upper().startswith("OK"), f"incomplete set reported balanced: {check}"
    assert "expected 1.0" not in check, f"incomplete set asserted imbalance: {check}"
    assert not _decimal_sum_issues(review_rows, section)


def test_complete_owner_set_imbalance_is_still_flagged(review_rows):
    issues = _decimal_sum_issues(review_rows, 26)
    assert issues and issues[0]["severity"] == "red"
    assert "0.75" in issues[0]["detail"]


def test_exact_duplicate_schedule_does_not_double_count(review_rows):
    assert not _decimal_sum_issues(review_rows, 27), (
        "an exact duplicate of a complete 1.0 schedule produced a false conflict")


def test_synthetic_corpus_marks_its_owner_schedule_complete(tmp_path):
    """The shipped self-test corpus must prove completeness explicitly so its
    0.95 imbalance is flagged without a heuristic."""
    grp.make_synthetic_corpus(tmp_path)
    text = (tmp_path / "04_ownership_note.txt").read_text(encoding="utf-8")
    assert re.search(re.escape(COMPLETE_OWNER_SET_MARKER), text, re.I)
