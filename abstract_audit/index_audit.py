"""Audit a Penterra abstract index workbook against the Section 15 standard.

Reports row counts, per-column fill rates, blank required cells,
duplicate document numbers, and header-block conformance, so a
section can be judged turn-in ready without opening Excel.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from schema import COLUMNS, HEADER_KEYS, REQUIRED_COLUMNS

#: Matches both legacy sequential numbers (1043044) and modern
#: e-recording numbers (2026-03115).
DOC_NO_RE = re.compile(r"^(?:\d{4,8}|\d{4}-\d{4,5})$")

#: Book-Page as used by the Campbell County clerk, e.g. 3145-0695 or 20LIEN-0432.
BOOK_PAGE_RE = re.compile(r"^[0-9A-Z]{3,8}-\d{3,4}$")

DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/(?:\d{2}|\d{4})$")


@dataclass
class ColumnStat:
    name: str
    filled: int = 0
    blank: int = 0

    @property
    def fill_rate(self) -> float:
        total = self.filled + self.blank
        return 1.0 if total == 0 else self.filled / total


@dataclass
class IndexReport:
    path: str
    sheet: str
    header: dict[str, str] = field(default_factory=dict)
    missing_header_keys: list[str] = field(default_factory=list)
    column_order_ok: bool = True
    found_columns: list[str] = field(default_factory=list)
    row_count: int = 0
    columns: dict[str, ColumnStat] = field(default_factory=dict)
    blank_required: list[dict[str, object]] = field(default_factory=list)
    duplicate_doc_nos: dict[str, list[int]] = field(default_factory=dict)
    malformed_doc_nos: list[dict[str, object]] = field(default_factory=list)
    malformed_dates: list[dict[str, object]] = field(default_factory=list)

    @property
    def completeness(self) -> float:
        """Fraction of required cells that are populated."""
        req = [c for n, c in self.columns.items() if n in REQUIRED_COLUMNS]
        total = sum(c.filled + c.blank for c in req)
        return 1.0 if total == 0 else sum(c.filled for c in req) / total

    def to_dict(self) -> dict[str, object]:
        d = asdict(self)
        d["columns"] = {
            n: {"filled": c.filled, "blank": c.blank, "fill_rate": round(c.fill_rate, 4)}
            for n, c in self.columns.items()
        }
        d["completeness"] = round(self.completeness, 4)
        return d


def _text(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%-m/%-d/%Y")
    return str(value).strip()


def _find_header_row(ws) -> int | None:
    """Return the 1-based row index of the document-table header."""
    for row in ws.iter_rows(min_row=1, max_row=40):
        values = [_text(c.value) for c in row]
        if values and values[0] == COLUMNS[0]:
            return row[0].row
    return None


def audit_workbook(path: Path) -> IndexReport:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    report = IndexReport(path=str(path), sheet=ws.title)

    header_row = _find_header_row(ws)
    if header_row is None:
        raise ValueError(f"{path}: no '{COLUMNS[0]}' header row found")

    # Header block above the table: "Key:,Value"
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        cells = [_text(c.value) for c in row]
        if not cells or not cells[0]:
            continue
        key = cells[0].rstrip(":").strip()
        value = next((c for c in cells[1:] if c), "")
        report.header[key] = value
    report.missing_header_keys = [k for k in HEADER_KEYS if k not in report.header]

    hdr_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
    report.found_columns = [_text(c.value) for c in hdr_cells if _text(c.value)]
    report.column_order_ok = report.found_columns[: len(COLUMNS)] == list(COLUMNS)

    for name in COLUMNS:
        report.columns[name] = ColumnStat(name=name)

    seen: dict[str, list[int]] = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        values = [_text(c.value) for c in row]
        if not any(values):
            continue
        report.row_count += 1
        rownum = row[0].row
        record = dict(zip(COLUMNS, values + [""] * len(COLUMNS)))

        for name in COLUMNS:
            stat = report.columns[name]
            if record[name]:
                stat.filled += 1
            else:
                stat.blank += 1
                if name in REQUIRED_COLUMNS:
                    report.blank_required.append(
                        {"row": rownum, "column": name,
                         "doc_no": record["Doc No"],
                         "document_type": record["Document Type"]}
                    )

        doc_no = record["Doc No"]
        if doc_no:
            seen.setdefault(doc_no, []).append(rownum)
            if not DOC_NO_RE.match(doc_no):
                report.malformed_doc_nos.append({"row": rownum, "value": doc_no})

        for date_col in ("Date of Doc", "Rec Date"):
            val = record[date_col]
            if val and not DATE_RE.match(val):
                report.malformed_dates.append(
                    {"row": rownum, "column": date_col, "value": val}
                )

    report.duplicate_doc_nos = {k: v for k, v in seen.items() if len(v) > 1}
    wb.close()
    return report


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbooks", nargs="+", type=Path)
    ap.add_argument("--json", type=Path, help="write full report JSON here")
    args = ap.parse_args(argv)

    reports = []
    for wb_path in args.workbooks:
        rep = audit_workbook(wb_path)
        reports.append(rep)
        print(f"\n=== {wb_path.name} ===")
        print(f"  rows                 : {rep.row_count}")
        print(f"  required completeness: {rep.completeness:.1%}")
        print(f"  column order ok      : {rep.column_order_ok}")
        if rep.missing_header_keys:
            print(f"  MISSING HEADER KEYS  : {', '.join(rep.missing_header_keys)}")
        for name, stat in rep.columns.items():
            flag = "  <-- REQUIRED" if name in REQUIRED_COLUMNS and stat.blank else ""
            print(f"    {name:<20} filled={stat.filled:<4} blank={stat.blank:<4}"
                  f" ({stat.fill_rate:.0%}){flag}")
        if rep.duplicate_doc_nos:
            print(f"  DUPLICATE Doc No     : {len(rep.duplicate_doc_nos)}")
            for doc, rows in list(rep.duplicate_doc_nos.items())[:10]:
                print(f"    {doc} -> rows {rows}")
        if rep.malformed_doc_nos:
            print(f"  MALFORMED Doc No     : {len(rep.malformed_doc_nos)}")
        if rep.malformed_dates:
            print(f"  MALFORMED dates      : {len(rep.malformed_dates)}")

    if args.json:
        args.json.write_text(
            json.dumps([r.to_dict() for r in reports], indent=2), encoding="utf-8"
        )
        print(f"\nwrote {args.json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
