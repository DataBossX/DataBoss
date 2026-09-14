"""Referee checks: defect classes that survive a per-column audit.

A workbook can be 100% populated and still be wrong. These checks catch
the two defect classes that blocked the Section 15 reference set:

1. A recording date that precedes its own document date — impossible,
   since a document must exist before the clerk records it.
2. The same party spelled with a different entity suffix across
   workbooks ("Jibber Resources, LLC" vs "Jibber Resources Inc."),
   which silently changes who holds an interest.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from index_audit import _find_header_row, _text
from schema import COLUMNS

#: Entity suffixes that distinguish otherwise identical party names.
SUFFIX_RE = re.compile(
    r"[,\s]+(?:an?\s+)?(LLC|L\.L\.C\.|L\.L\.L\.P\.|LLLP|LP|L\.P\.|INC|INC\.|"
    r"CORP|CORP\.|CORPORATION|COMPANY|CO|CO\.|TRUST|LTD|LTD\.)\s*$",
    re.I,
)

_SUFFIX_CANON = {
    "LLC": "LLC", "L.L.C.": "LLC",
    "LLLP": "LLLP", "L.L.L.P.": "LLLP",
    "LP": "LP", "L.P.": "LP",
    "INC": "INC", "INC.": "INC",
    "CORP": "CORP", "CORP.": "CORP", "CORPORATION": "CORP",
    "COMPANY": "CO", "CO": "CO", "CO.": "CO",
    "TRUST": "TRUST", "LTD": "LTD", "LTD.": "LTD",
}


@dataclass
class DateDefect:
    workbook: str
    row: int
    doc_no: str
    date_of_doc: str
    rec_date: str

    def describe(self) -> str:
        return (
            f"{self.workbook} row {self.row} (Doc {self.doc_no}): "
            f"recorded {self.rec_date} precedes document date {self.date_of_doc}"
        )


@dataclass
class NameDefect:
    stem: str
    variants: dict[str, list[str]]

    def describe(self) -> str:
        parts = "; ".join(
            f"{v} in {', '.join(sorted(set(w)))}" for v, w in sorted(self.variants.items())
        )
        return f"'{self.stem}' spelled inconsistently: {parts}"


def _parse_date(value: str) -> date | None:
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def split_entity(name: str) -> tuple[str, str] | None:
    """Split 'Jibber Resources, LLC' into ('jibber resources', 'LLC')."""
    m = SUFFIX_RE.search(name)
    if not m:
        return None
    stem = name[: m.start()].strip().rstrip(",").lower()
    suffix = _SUFFIX_CANON.get(m.group(1).upper(), m.group(1).upper())
    return (stem, suffix) if stem else None


def read_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = _find_header_row(ws)
    if header_row is None:
        wb.close()
        raise ValueError(f"{path}: no '{COLUMNS[0]}' header row found")
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1):
        values = [_text(c.value) for c in row]
        if not any(values):
            continue
        record = dict(zip(COLUMNS, values + [""] * len(COLUMNS)))
        record["_row"] = str(row[0].row)
        rows.append(record)
    wb.close()
    return rows


def check_date_order(path: Path, rows: list[dict[str, str]]) -> list[DateDefect]:
    """Flag rows where Rec Date precedes Date of Doc."""
    defects = []
    for record in rows:
        doc_d = _parse_date(record["Date of Doc"])
        rec_d = _parse_date(record["Rec Date"])
        if doc_d and rec_d and rec_d < doc_d:
            defects.append(
                DateDefect(
                    workbook=path.name, row=int(record["_row"]),
                    doc_no=record["Doc No"],
                    date_of_doc=record["Date of Doc"], rec_date=record["Rec Date"],
                )
            )
    return defects


def check_entity_consistency(
    per_workbook: dict[str, list[dict[str, str]]],
) -> list[NameDefect]:
    """Flag a party name carrying different entity suffixes across workbooks."""
    stems: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for wb_name, rows in per_workbook.items():
        for record in rows:
            for col in ("Grantor", "Grantee"):
                for part in record[col].split(";"):
                    part = part.strip()
                    if not part:
                        continue
                    if split := split_entity(part):
                        stem, suffix = split
                        stems[stem][suffix].append(wb_name)

    return [
        NameDefect(stem=stem, variants=dict(variants))
        for stem, variants in sorted(stems.items())
        if len(variants) > 1
    ]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbooks", nargs="+", type=Path)
    ap.add_argument("--json", type=Path)
    args = ap.parse_args(argv)

    per_workbook: dict[str, list[dict[str, str]]] = {}
    date_defects: list[DateDefect] = []
    for path in args.workbooks:
        rows = read_rows(path)
        per_workbook[path.name] = rows
        date_defects.extend(check_date_order(path, rows))

    name_defects = check_entity_consistency(per_workbook)

    print(f"DATE-ORDER defects: {len(date_defects)}")
    for d in date_defects:
        print(f"  {d.describe()}")
    print(f"ENTITY-NAME defects: {len(name_defects)}")
    for n in name_defects:
        print(f"  {n.describe()}")

    if args.json:
        args.json.write_text(
            json.dumps(
                {
                    "date_order": [d.__dict__ for d in date_defects],
                    "entity_names": [
                        {"stem": n.stem, "variants": n.variants} for n in name_defects
                    ],
                },
                indent=2,
            ),
            encoding="utf-8",
        )
        print(f"wrote {args.json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
