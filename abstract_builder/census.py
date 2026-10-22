"""Read any existing abstract index workbook (xlsx or ods) into a uniform structure and count it."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from .normalize import is_blank, is_unknown

COUNTY_KEY = ("Document Type", "Grantor", "Grantee", "Doc No", "Book-Page")
BLM_KEY = ("Document Type", "Grantor", "Grantee", "File No", "Page")


@dataclass
class SheetCensus:
    path: str
    sheet: str
    layout: str                       # county | blm | unknown
    header_row: int | None
    meta: dict[str, Any] = field(default_factory=dict)
    rows: list[list[Any]] = field(default_factory=list)
    n_rows: int = 0
    n_blank_substantive: int = 0
    n_unknown: int = 0
    n_missing_image: int = 0
    first_font: str | None = None
    sha256: str | None = None
    bytes: int = 0

    def as_dict(self) -> dict[str, Any]:
        d = self.__dict__.copy()
        d.pop("rows")
        d["meta"] = json.dumps(self.meta, default=str)
        return d


def to_xlsx(path: Path) -> Path:
    """Convert .ods (or anything LibreOffice reads) to .xlsx in a temp dir; returns the xlsx path."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return path
    from .render import convert
    out = Path(tempfile.mkdtemp(prefix="ab_census_"))
    cand = convert(path, "xlsx", out)
    if not cand.exists():
        raise RuntimeError(f"conversion failed for {path}")
    return cand


def _norm_header(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def detect_layout(ws) -> tuple[str, int | None]:
    for r in range(1, min(ws.max_row, 30) + 1):
        a = _norm_header(ws.cell(r, 1).value)
        if a == "document type":
            d = _norm_header(ws.cell(r, 4).value)
            if d == "file no":
                return "blm", r
            if d == "doc no":
                return "county", r
            return "unknown", r
    return "unknown", None


def census_workbook(path: Path, group: str = "") -> list[SheetCensus]:
    path = Path(path)
    raw = path.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    xp = to_xlsx(path)
    wb = openpyxl.load_workbook(xp, data_only=True)
    out = []
    for ws in wb.worksheets:
        layout, hr = detect_layout(ws)
        sc = SheetCensus(str(path), ws.title, layout, hr, sha256=sha, bytes=len(raw))
        ncols = 9 if layout == "county" else 10 if layout == "blm" else ws.max_column
        if hr:
            for r in range(1, hr):
                k = ws.cell(r, 1).value
                if k is not None and str(k).strip():
                    sc.meta[str(k).strip().rstrip(":")] = ws.cell(r, 2).value
            subs = [1, 2, 3, 4, 5, 6, 7, 8] if layout == "county" else [1, 2, 3, 5, 6, 7, 9, 10]
            for r in range(hr + 1, ws.max_row + 1):
                vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
                if all(is_blank(v) for v in vals):
                    continue
                sc.rows.append(vals)
                if sc.first_font is None:
                    sc.first_font = ws.cell(r, 1).font.name
                for c in subs:
                    v = vals[c - 1]
                    if is_blank(v):
                        sc.n_blank_substantive += 1
                    elif is_unknown(v):
                        sc.n_unknown += 1
                cm = vals[8] if layout == "county" else None
                if isinstance(cm, str) and cm.strip().lower() == "missing image":
                    sc.n_missing_image += 1
            sc.n_rows = len(sc.rows)
        out.append(sc)
    return out
