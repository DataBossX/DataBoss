"""Format-compliance gates for a rendered index workbook (gate C 'fields & donor format')."""
from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from . import donor as D
from .normalize import is_blank, is_unknown


def custom_height_rows(xlsx: Path) -> int:
    """Number of data rows (r >= 10) whose height is pinned (customHeight=1). Auto-fit rows carry ht but customHeight=0,
    which openpyxl cannot tell apart, so the raw sheet XML is read."""
    with zipfile.ZipFile(xlsx) as z:
        name = next(n for n in z.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))
        x = z.read(name).decode("utf-8", "replace")
    n = 0
    for tag in re.findall(r"<(?:\w+:)?row\b[^>]*>", x):
        m = re.search(r'\br="(\d+)"', tag)
        if m and int(m.group(1)) >= 10 and re.search(r'customHeight="(?:1|true)"', tag):
            n += 1
    return n


def check_index(xlsx: Path, allow_text_doc_no: bool = False) -> list[str]:
    """Return a list of failures (empty == pass). allow_text_doc_no accepts year-sequence recorder IDs such as
    2020-01337 (Campbell County) that the notes require to stay text because they are not plain integers."""
    fails: list[str] = []
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    layout = "blm" if str(ws.cell(9, 4).value).strip() == "File No" else "county"
    headers = D.BLM_HEADERS if layout == "blm" else D.COUNTY_HEADERS
    ncols = len(headers)
    for c, h in enumerate(headers, start=1):
        if ws.cell(9, c).value != h:
            fails.append(f"HEADER_MISMATCH:{ws.cell(9, c).coordinate}={ws.cell(9, c).value!r} expected {h!r}")
    for i, lab in enumerate(D.META_LABELS, start=1):
        if ws.cell(i, 1).value != lab:
            fails.append(f"META_LABEL:{i}={ws.cell(i,1).value!r}")
    if any(not is_blank(ws.cell(8, c).value) for c in range(1, ncols + 1)):
        fails.append("ROW8_NOT_BLANK")
    if ws.page_setup.orientation != "landscape":
        fails.append("NOT_LANDSCAPE")
    if not (ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1):
        fails.append("NOT_FIT_ONE_PAGE_WIDE")
    pm = ws.page_margins
    if (round(pm.left, 2), round(pm.right, 2), round(pm.top, 2), round(pm.bottom, 2)) != (0.7, 0.7, 0.3, 0.3):
        fails.append(f"MARGINS:{pm.left},{pm.right},{pm.top},{pm.bottom}")
    date_cols = D.BLM_DATE_COLS if layout == "blm" else D.COUNTY_DATE_COLS
    for r in range(10, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if all(is_blank(v) for v in vals):
            continue
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if cell.font.name != D.FONT_NAME or (cell.font.size or 0) != D.FONT_SIZE or cell.font.bold:
                fails.append(f"FONT:{cell.coordinate}")
            if cell.fill.fill_type not in (None, "none"):
                fails.append(f"FILL:{cell.coordinate}")
            b = cell.border
            if any(getattr(b, s).style != D.BORDER_STYLE for s in ("left", "right", "top", "bottom")):
                fails.append(f"BORDER:{cell.coordinate}")
            if not cell.alignment.wrap_text or cell.alignment.vertical != "top":
                fails.append(f"ALIGN:{cell.coordinate}")
            if is_unknown(v):
                fails.append(f"UNKNOWN_IN_CLIENT_CELL:{cell.coordinate}")
            if c in date_cols and not is_blank(v):
                if not isinstance(v, (date, datetime)):
                    fails.append(f"DATE_NOT_REAL:{cell.coordinate}={v!r}")
                elif cell.number_format != D.DATE_FORMAT:
                    fails.append(f"DATE_FORMAT:{cell.coordinate}")
            if layout == "county" and c == 4 and not is_blank(v) and not isinstance(v, int) and not (allow_text_doc_no and isinstance(v, str)):
                fails.append(f"DOC_NO_NOT_INT:{cell.coordinate}={v!r}")
            if layout == "county" and c == 9 and not is_blank(v) and v != "missing image":
                fails.append(f"COMMENT_TEXT:{cell.coordinate}={v!r}")
            if layout == "blm" and c in (4, 8) and not is_blank(v):
                fails.append(f"BLM_FILE_NO_OR_COMMENT_NOT_BLANK:{cell.coordinate}")
            if layout == "blm" and c in (5, 10) and not is_blank(v) and not isinstance(v, int):
                fails.append(f"BLM_PAGE_PART_NOT_INT:{cell.coordinate}={v!r}")
    n_fixed = custom_height_rows(xlsx)
    if n_fixed:
        fails.append(f"FIXED_ROW_HEIGHT:{n_fixed}_rows_pinned")
    if ws.freeze_panes:
        fails.append(f"FREEZE_PANES:{ws.freeze_panes}")
    widths = D.BLM_WIDTHS if layout == "blm" else D.COUNTY_WIDTHS
    from openpyxl.utils import get_column_letter
    for c, w in enumerate(widths, start=1):
        got = ws.column_dimensions[get_column_letter(c)].width
        if got is None or abs(got - w) > 1.0:
            fails.append(f"COL_WIDTH:{get_column_letter(c)}={got if got is None else round(got, 2)} donor {w}")
    return sorted(set(fails))
