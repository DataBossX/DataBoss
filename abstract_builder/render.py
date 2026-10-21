"""Render county / BLM index workbooks, the 3-band checklist and the certification letter in the measured
Section 11 Overtime donor geometry. ODS and PDF are produced through LibreOffice (soffice --headless)."""
from __future__ import annotations

import copy
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import donor as D


@dataclass
class IndexMeta:
    county: str                      # "Johnson"
    lands: str                       # "T47N-R76W Section 6: All"  or "WYW-134914"
    date: date                       # preparation date
    starting_date: str | date | None # "Inception" or None
    date_posted_thru: date | None
    indexed_by: str = ""
    project: str = ""


def _font(bold=False) -> Font:
    return Font(name=D.FONT_NAME, size=D.FONT_SIZE, bold=bold, color="FF000000")


def _border() -> Border:
    s = Side(style=D.BORDER_STYLE, color="FF000000")
    return Border(left=s, right=s, top=s, bottom=s)


_NOFILL = PatternFill(fill_type=None)


def _write_meta(ws, meta: IndexMeta, ncols: int):
    vals = [meta.county, meta.lands, meta.date, meta.starting_date, meta.date_posted_thru, meta.indexed_by, meta.project]
    for i, (label, v) in enumerate(zip(D.META_LABELS, vals), start=1):
        a = ws.cell(i, 1, label); a.font = _font(True); a.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        b = ws.cell(i, 2, v); b.font = _font(); b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if isinstance(v, date):
            b.number_format = D.DATE_FORMAT
        ws.row_dimensions[i].height = D.ROW_HEIGHTS_META
    ws.row_dimensions[8].height = D.ROW_HEIGHT_BLANK8


def _page_setup(ws, cfg: dict, fit_width: bool = True, print_titles: str | None = None):
    ws.page_setup.orientation = cfg["orientation"]
    ws.page_setup.paperSize = cfg.get("paperSize", 1)
    if fit_width:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = cfg.get("fitToWidth", 1)
        ws.page_setup.fitToHeight = cfg.get("fitToHeight", 0)
    pm = ws.page_margins
    pm.left, pm.right, pm.top, pm.bottom = cfg["left"], cfg["right"], cfg["top"], cfg["bottom"]
    pm.header, pm.footer = cfg.get("header", 0.3), cfg.get("footer", 0.3)
    if print_titles:
        ws.print_title_rows = print_titles


def write_index(rows: list[list[Any]], meta: IndexMeta, layout: str, out_xlsx: Path, repeat_header: bool = True) -> Path:
    """rows: already-normalized value lists (9 for county, 10 for blm). Dates must be datetime.date, ints int."""
    assert layout in ("county", "blm")
    headers = D.COUNTY_HEADERS if layout == "county" else D.BLM_HEADERS
    widths = D.COUNTY_WIDTHS if layout == "county" else D.BLM_WIDTHS
    num_cols = D.COUNTY_NUMERIC_COLS if layout == "county" else D.BLM_NUMERIC_COLS
    date_cols = D.COUNTY_DATE_COLS if layout == "county" else D.BLM_DATE_COLS
    int_cols = D.COUNTY_INT_COLS if layout == "county" else D.BLM_INT_COLS
    ncols = len(headers)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Index"
    _write_meta(ws, meta, ncols)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(9, c, h)
        cell.font = _font(True); cell.border = _border(); cell.fill = _NOFILL
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.row_dimensions[9].height = D.ROW_HEIGHT_HEADER_COUNTY if layout == "county" else D.ROW_HEIGHT_HEADER_BLM
    for i, row in enumerate(rows, start=10):
        row = list(row) + [None] * (ncols - len(row))
        for c in range(1, ncols + 1):
            v = row[c - 1]
            cell = ws.cell(i, c, v)
            cell.font = _font(); cell.border = _border(); cell.fill = _NOFILL
            centre = D.DONOR_CENTER_NUMERIC and c in num_cols
            cell.alignment = Alignment(horizontal="center" if centre else "left", vertical="top", wrap_text=True)
            if c in date_cols:
                cell.number_format = D.DATE_FORMAT
            elif c in int_cols and isinstance(v, int):
                cell.number_format = D.INT_FORMAT
        # no explicit height => Excel/LibreOffice auto-fit (the double-click behaviour)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    _page_setup(ws, D.PAGE, True, "9:9" if repeat_header else None)
    out_xlsx = Path(out_xlsx); out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx


@dataclass
class ChecklistData:
    band1: list[str]        # 11 values under CHECKLIST_BAND1
    band2: list[Any]        # 8 values
    band3: list[Any]        # 7 values


def write_checklist(data: ChecklistData, out_xlsx: Path) -> Path:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = D.CHECKLIST_SHEET
    white = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    def put(r, vals, bold):
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.font = _font(bold); cell.border = _border(); cell.fill = white
            left = (r == 9 and c in D.CHECKLIST_LEFT_ALIGNED_ROW9)
            cell.alignment = Alignment(horizontal="left" if left else "center", vertical="center", wrap_text=True)
            if isinstance(v, date):
                cell.number_format = D.DATE_FORMAT
    put(2, D.CHECKLIST_BAND1, True); put(3, data.band1, False)
    put(5, D.CHECKLIST_BAND2, True); put(6, data.band2, False)
    put(8, D.CHECKLIST_BAND3, True); put(9, data.band3, False)
    for r, h in D.CHECKLIST_ROW_HEIGHTS.items():
        ws.row_dimensions[r].height = h
    for c, w in enumerate(D.CHECKLIST_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    _page_setup(ws, dict(D.CHECKLIST_PAGE, paperSize=1), fit_width=True)
    out_xlsx = Path(out_xlsx); out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx


@dataclass
class LetterData:
    twp: int; rge: int; section: int
    county: str                          # "Johnson"
    county_cutoff_text: str              # e.g. "from Inception to last Indexed Instrument Recorded October 6, 2026."
    assessor_text: str = "N/A"
    federal_serials: list[str] = field(default_factory=list)   # ["WYW-149799"]
    federal_through_text: str | None = None                      # "October 13, 2026" or None
    client: str = ""
    prepared_on: date = date(2026, 10, 20)
    preparer: str = ""
    qualified_note: str | None = None    # appended paragraph when not final


def _serials_phrase(serials: list[str]) -> str:
    if not serials:
        return ""
    if len(serials) == 1:
        return serials[0]
    return ", ".join(serials[:-1]) + " and " + serials[-1]


def _long_date(d: date) -> str:
    return f"{d.strftime('%B')} {d.day}, {d.year}"


def write_letter(data: LetterData, donor_docx: Path, out_docx: Path) -> Path:
    """Build the letter from the donor DOCX so fonts, margins and spacing carry over exactly.
    The donor paragraph structure (P0 title, P1 'Contents of abstract:', P2 index, P3 assessor, P4 lease files,
    P5 SRPs, P6 lien search, P7 disclaimer, P8 prepared, P9 line, P10/P11 signature) is preserved; text is replaced
    run-by-run so the title keeps its 14pt/9pt-superscript 'th'."""
    import docx
    d = docx.Document(str(donor_docx))
    P = d.paragraphs
    def set_text(p, text, keep_bold=None):
        runs = p.runs
        if not runs:
            p.add_run(text); return
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    # P0 title: runs = ['Township 45 North, Range 76 West, 6', 'th', ' ', 'P.M.', ...'\nSection 11: All Campbell County, Wyoming']
    r = P[0].runs
    r[0].text = f"Township {data.twp} North, Range {data.rge} West, 6"
    # find the run carrying the Section line (contains 'Section')
    for run in r[1:]:
        if "Section" in run.text:
            run.text = re.sub(r"Section \d+: All .* County, Wyoming",
                              f"Section {data.section}: All {data.county} County, Wyoming", run.text)
    set_text(P[2], f"\nIndex of Documents reflecting on the entries posted in the tract indices located and maintained\n"
                   f"by the {data.county} County Clerk's office, {data.county_cutoff_text}")
    set_text(P[3], f"\n{data.county} County Assessor and/or Treasurer office reports, detailing surface owner property "
                   f"accounts and the state of any taxes due, {data.assessor_text}.")
    if data.federal_serials:
        ser = _serials_phrase(data.federal_serials)
        thru = f" from date of lease sale to {data.federal_through_text}." if data.federal_through_text else \
               " from date of lease sale through the dates shown on the Serial Register Pages."
        pl = "files" if len(data.federal_serials) > 1 else "file"
        set_text(P[4], f"Images of Federal Lease {pl} {ser}{thru}")
        set_text(P[5], f"\nBureau of Land Management (MLRS) Serial Register Pages {ser}{thru}")
    else:
        set_text(P[4], "Images of Federal Lease files, N/A.")
        set_text(P[5], "\nBureau of Land Management (MLRS) Serial Register Pages, N/A.")
    set_text(P[7], f"\nThis abstract was prepared and made solely for the benefit of {data.client}\n"
                   f"and is not to be construed as a title opinion, nor a guaranty of title, and the information\n"
                   f"contained herein is conditioned upon the accuracy of the tract indices of {data.county} County,\n"
                   f"including any errors or omissions contained therein. This abstract is limited to the materials\n"
                   f"reviewed and searched and described herein. The instruments referred to have not been examined\n"
                   f"to determine their legal sufficiency.")
    set_text(P[8], f"Prepared on {_long_date(data.prepared_on)}, through the source dates shown above by: {data.preparer}.")
    if data.qualified_note:
        P[8].runs[0].text += f" {data.qualified_note}"
    out_docx = Path(out_docx); out_docx.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_docx))
    return out_docx


def convert(path: Path, fmt: str, outdir: Path | None = None) -> Path:
    """LibreOffice conversion: fmt in {'ods','pdf','xlsx','docx'}."""
    path = Path(path); outdir = Path(outdir or path.parent)
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=300)
    out = outdir / (path.stem + "." + fmt)
    if not out.exists():
        raise RuntimeError(f"conversion to {fmt} failed for {path}")
    return out
