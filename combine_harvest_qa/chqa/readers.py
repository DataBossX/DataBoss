"""Native readers: ODS (lxml, no LibreOffice needed) and XLSX (openpyxl) -> model.Workbook."""
from __future__ import annotations

import datetime as dt
import re
import zipfile
from typing import Optional

from lxml import etree

from .model import Cell, CellStyle, PrintSetup, Sheet, Workbook, col_letter

NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "number": "urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "meta": "urn:oasis:names:tc:opendocument:xmlns:meta:1.0",
    "dc": "http://purl.org/dc/elements/1.1/",
}
MAX_EMPTY_REPEAT = 64  # repeated empty cells/rows beyond this are padding, not data


def _q(prefix: str, name: str) -> str:
    return "{%s}%s" % (NS[prefix], name)


def _len_to_in(v: Optional[str]) -> Optional[float]:
    if not v:
        return None
    m = re.match(r"^(-?[\d.]+)\s*(cm|mm|in|pt)?$", v.strip())
    if not m:
        return None
    n, u = float(m.group(1)), m.group(2) or "in"
    return {"cm": n / 2.54, "mm": n / 25.4, "in": n, "pt": n / 72.0}[u]


def _len_to_pt(v: Optional[str]) -> Optional[float]:
    i = _len_to_in(v)
    return None if i is None else i * 72.0


def _border_name(spec: str) -> str:
    tok = spec.split()[0]
    m = re.match(r"^([\d.]+)(pt|cm|mm|in)?$", tok)
    if not m:
        return tok            # already a keyword (thin/medium/hair)
    pt = float(m.group(1)) * {"pt": 1.0, "cm": 28.3465, "mm": 2.83465, "in": 72.0, None: 1.0}[m.group(2)]
    return "hair" if pt < 0.5 else "thin" if pt < 1.1 else "medium" if pt < 2.0 else "thick"


def _cell_text(cell: etree._Element) -> str:
    paras = []
    for p in cell.findall(_q("text", "p")):
        buf = []
        for node in p.iter():
            if node.tag == _q("text", "line-break"):
                buf.append("\n")
            elif node.tag == _q("text", "s"):
                buf.append(" " * int(node.get(_q("text", "c"), "1")))
            elif node.tag == _q("text", "tab"):
                buf.append("\t")
            if node.tag != p.tag and node.tail:
                buf.append(node.tail)
            if node is p and p.text:
                buf.append(p.text)
            elif node is not p and node.text and node.tag in (_q("text", "span"), _q("text", "a")):
                buf.append(node.text)
        paras.append("".join(buf))
    return "\n".join(paras)


def _parse_style_props(el: etree._Element) -> dict:
    out: dict = {}
    for child in el:
        tag = etree.QName(child).localname
        for k, v in child.attrib.items():
            out[f"{tag}.{etree.QName(k).localname}"] = v
    return out


def read_ods(path: str) -> Workbook:
    with zipfile.ZipFile(path) as z:
        content = etree.fromstring(z.read("content.xml"))
        styles = etree.fromstring(z.read("styles.xml")) if "styles.xml" in z.namelist() else None
        meta = etree.fromstring(z.read("meta.xml")) if "meta.xml" in z.namelist() else None

    cell_styles: dict[str, dict] = {}
    col_styles: dict[str, dict] = {}
    row_styles: dict[str, dict] = {}
    parents: dict[str, str] = {}
    num_formats: dict[str, str] = {}

    def harvest(root: etree._Element, xpath: str) -> None:
        for s in root.xpath(xpath, namespaces=NS):
            name = s.get(_q("style", "name"))
            fam = s.get(_q("style", "family"))
            props = _parse_style_props(s)
            if s.get(_q("style", "parent-style-name")):
                parents[name] = s.get(_q("style", "parent-style-name"))
            if s.get(_q("style", "data-style-name")):
                props["data-style"] = s.get(_q("style", "data-style-name"))
            if fam == "table-cell":
                cell_styles[name] = props
            elif fam == "table-column":
                col_styles[name] = props
            elif fam == "table-row":
                row_styles[name] = props

    harvest(content, "//office:automatic-styles/style:style")
    if styles is not None:
        harvest(styles, "//office:styles/style:style")
        harvest(styles, "//office:automatic-styles/style:style")
    for root in [content] + ([styles] if styles is not None else []):
        for nf in root.xpath("//number:date-style|//number:number-style|//number:text-style|//number:percentage-style", namespaces=NS):
            name = nf.get(_q("style", "name"))
            parts = []
            for ch in nf:
                ln = etree.QName(ch).localname
                if ln in ("year", "month", "day", "hours", "minutes", "seconds"):
                    style_ = ch.get(_q("number", "style"), "short")
                    sym = {"year": "yyyy" if style_ == "long" else "yy",
                           "month": "mm" if style_ == "long" else "m",
                           "day": "dd" if style_ == "long" else "d",
                           "hours": "hh", "minutes": "mm", "seconds": "ss"}[ln]
                    parts.append(sym)
                elif ln == "text":
                    parts.append(ch.text or "")
                elif ln == "number":
                    mind = int(ch.get(_q("number", "min-integer-digits"), "1"))
                    dec = int(ch.get(_q("number", "decimal-places"), "0"))
                    parts.append(("0" * mind) + (("." + "0" * dec) if dec else ""))
            num_formats[name] = "".join(parts)

    default_font = None
    if styles is not None:
        for d in styles.xpath("//style:default-style[@style:family='table-cell']/style:text-properties", namespaces=NS):
            fn = d.get(_q("style", "font-name")) or d.get(_q("fo", "font-family"))
            fs = _len_to_pt(d.get(_q("fo", "font-size")))
            default_font = (fn, fs)

    if default_font is None and "Default" in cell_styles:
        d = cell_styles["Default"]
        fn = d.get("text-properties.font-name") or d.get("text-properties.font-family")
        default_font = (fn.strip("'\"") if fn else None, _len_to_pt(d.get("text-properties.font-size")))
        wb_default_font = default_font
    else:
        wb_default_font = default_font

    def resolve(name: Optional[str]) -> dict:
        chain, seen = [], set()
        if name and name != "Default" and name not in parents:
            parents[name] = "Default"
        while name and name not in seen:
            seen.add(name)
            chain.append(cell_styles.get(name, {}))
            name = parents.get(name)
        merged: dict = {}
        for d in reversed(chain):
            merged.update(d)
        return merged

    def to_style(p: dict) -> CellStyle:
        st = CellStyle()
        fn = p.get("text-properties.font-name") or p.get("text-properties.font-family")
        st.font_name = fn.strip("'\"") if fn else (default_font[0] if default_font else None)
        fs = _len_to_pt(p.get("text-properties.font-size"))
        st.font_size = fs if fs else (default_font[1] if default_font else None)
        st.bold = p.get("text-properties.font-weight") == "bold"
        st.italic = p.get("text-properties.font-style") == "italic"
        st.underline = p.get("text-properties.text-underline-style", "none") != "none"
        c = p.get("text-properties.color")
        st.color = c.upper() if c else None
        bg = p.get("table-cell-properties.background-color")
        st.fill = None if bg in (None, "transparent") else bg.upper()
        ta = p.get("paragraph-properties.text-align")
        st.h_align = {"start": "left", "left": "left", "center": "center", "end": "right", "right": "right"}.get(ta)
        va = p.get("table-cell-properties.vertical-align")
        st.v_align = {"top": "top", "middle": "center", "bottom": "bottom"}.get(va)
        st.wrap = p.get("table-cell-properties.wrap-option") == "wrap"
        st.shrink = p.get("table-cell-properties.shrink-to-fit") == "true"
        sides = {}
        for side in ("", "-top", "-bottom", "-left", "-right"):
            b = p.get(f"table-cell-properties.border{side}")
            if b and b != "none":
                sides[side or "all"] = _border_name(b)
        if "all" in sides and len(sides) == 1:
            st.border = sides["all"]
        elif sides:
            st.border = ",".join(f"{k.strip('-') or 'all'}:{v}" for k, v in sorted(sides.items()))
        ds = p.get("data-style")
        st.number_format = num_formats.get(ds) if ds else None
        return st

    wb = Workbook(path=path, fmt="ods", default_font=wb_default_font)
    if meta is not None:
        for tag in ("generator", "creation-date"):
            for e in meta.xpath(f"//meta:{tag}", namespaces=NS):
                wb.properties[tag] = e.text or ""
        for e in meta.xpath("//dc:date|//dc:creator|//dc:title", namespaces=NS):
            wb.properties[etree.QName(e).localname] = e.text or ""

    layouts: dict[str, dict] = {}
    masters: list[tuple[str, str]] = []
    if styles is not None:
        for pl in styles.xpath("//style:page-layout", namespaces=NS):
            pr = pl.find(_q("style", "page-layout-properties"))
            layouts[pl.get(_q("style", "name"))] = dict(pr.attrib) if pr is not None else {}
        for mp in styles.xpath("//style:master-page", namespaces=NS):
            masters.append((mp.get(_q("style", "name")), mp.get(_q("style", "page-layout-name"))))
    table_master: dict[str, str] = {}
    for ts in content.xpath("//office:automatic-styles/style:style[@style:family='table']", namespaces=NS):
        if ts.get(_q("style", "master-page-name")):
            table_master[ts.get(_q("style", "name"))] = ts.get(_q("style", "master-page-name"))

    def layout_for(tbl_style: Optional[str]) -> dict:
        mname = table_master.get(tbl_style or "")
        for name, lay in masters:
            if mname is None or name == mname:
                return layouts.get(lay, {})
        return {}

    for tbl in content.xpath("//table:table", namespaces=NS):
        sh = Sheet(name=tbl.get(_q("table", "name")))
        col_default: dict[int, str] = {}
        c = 0
        for colnode in tbl.xpath("./table:table-column|./table:table-column-group/table:table-column|./table:table-header-columns/table:table-column", namespaces=NS):
            rep = int(colnode.get(_q("table", "number-columns-repeated"), "1"))
            w = _len_to_in(col_styles.get(colnode.get(_q("table", "style-name")), {}).get("table-column-properties.column-width"))
            dflt = colnode.get(_q("table", "default-cell-style-name"))
            if rep > MAX_EMPTY_REPEAT:
                break
            for _ in range(rep):
                c += 1
                if dflt:
                    col_default[c] = dflt
                if w is not None:
                    sh.col_widths_in[c] = w
        r = 0
        for rownode in tbl.xpath("./table:table-row|./table:table-header-rows/table:table-row|./table:table-row-group/table:table-row", namespaces=NS):
            rrep = int(rownode.get(_q("table", "number-rows-repeated"), "1"))
            rh = _len_to_pt(row_styles.get(rownode.get(_q("table", "style-name")), {}).get("table-row-properties.row-height"))
            row_default = rownode.get(_q("table", "default-cell-style-name"))
            cells = []
            cc = 0
            for cn in rownode.xpath("./table:table-cell|./table:covered-table-cell", namespaces=NS):
                crep = int(cn.get(_q("table", "number-columns-repeated"), "1"))
                vt = cn.get(_q("office", "value-type"))
                txt = _cell_text(cn)
                is_empty = vt is None and not txt
                explicit = cn.get(_q("table", "style-name"))
                if is_empty and crep > MAX_EMPTY_REPEAT:
                    # padding run: LibreOffice folds styled-but-blank columns into it; honour column defaults (first 64 cols)
                    for k in range(cc + 1, min(cc + crep, 64) + 1):
                        dprop = resolve(explicit or row_default or col_default.get(k))
                        if any(key.startswith("table-cell-properties.border") for key in dprop):
                            cells.append(Cell(row=0, col=k, style=to_style(dprop)))
                    cc += crep
                    continue
                for _ in range(crep):
                    cc += 1
                    sname = explicit or row_default or col_default.get(cc)
                    prop = resolve(sname)
                    if is_empty and not prop.get("table-cell-properties.border") and not any(k.startswith("table-cell-properties.border") for k in prop):
                        continue
                    cell = Cell(row=0, col=cc, style=to_style(prop))
                    if vt == "float":
                        f = float(cn.get(_q("office", "value")))
                        cell.value, cell.kind = (int(f), "int") if f == int(f) else (f, "float")
                    elif vt == "date":
                        raw = cn.get(_q("office", "date-value"))
                        cell.value = dt.datetime.fromisoformat(raw).date() if "T" not in raw or raw.endswith("T00:00:00") else dt.datetime.fromisoformat(raw)
                        cell.kind = "date"
                    elif vt == "string":
                        cell.value, cell.kind = txt, "string"
                    elif vt == "boolean":
                        cell.value, cell.kind = cn.get(_q("office", "boolean-value")) == "true", "bool"
                    cell.text = txt
                    anno = cn.find(_q("office", "annotation"))
                    if anno is not None:
                        cell.comment = "".join(anno.itertext()).strip()
                    cells.append(cell)
            all_empty = all(x.kind == "empty" for x in cells)
            if all_empty and rrep > MAX_EMPTY_REPEAT:
                r += rrep
                continue
            for _ in range(rrep):
                r += 1
                if rh:
                    sh.row_heights_pt[r] = rh
                for cell in cells:
                    cp = Cell(row=r, col=cell.col, value=cell.value, kind=cell.kind, text=cell.text,
                              style=cell.style, comment=cell.comment)
                    sh.cells[(r, cp.col)] = cp
                    sh.max_row = max(sh.max_row, r) if cp.kind != "empty" else sh.max_row
                    sh.max_col = max(sh.max_col, cp.col) if cp.kind != "empty" else sh.max_col
        page_props = layout_for(tbl.get(_q("table", "style-name")))
        ps = sh.print
        ps.orientation = page_props.get(_q("style", "print-orientation"))
        ps.margin_top = _len_to_in(page_props.get(_q("fo", "margin-top")))
        ps.margin_bottom = _len_to_in(page_props.get(_q("fo", "margin-bottom")))
        ps.margin_left = _len_to_in(page_props.get(_q("fo", "margin-left")))
        ps.margin_right = _len_to_in(page_props.get(_q("fo", "margin-right")))
        ps.paper_w_in = _len_to_in(page_props.get(_q("fo", "page-width")))
        ps.paper_h_in = _len_to_in(page_props.get(_q("fo", "page-height")))
        sx = page_props.get(_q("style", "scale-to-X"))
        sy = page_props.get(_q("style", "scale-to-Y"))
        sp = page_props.get(_q("style", "scale-to"))
        ps.fit_width = int(sx) if sx else None
        ps.fit_height = int(sy) if sy else None
        ps.scale_pct = int(float(sp.rstrip("%")) * (100 if "%" not in sp else 1)) if sp else None
        rep_rows = tbl.get(_q("table", "print-ranges"))
        hdr = tbl.find(_q("table", "table-header-rows"))
        if hdr is not None:
            ps.title_rows = "header-rows"
        wb.sheets.append(sh)
    return wb


def _side_style(border, side: str):
    sd = getattr(border, side, None) if border is not None else None
    return getattr(sd, "style", None) if sd is not None else None


def read_xlsx(path: str) -> Workbook:
    import openpyxl
    from openpyxl.utils import get_column_letter

    ox = openpyxl.load_workbook(path)
    wb = Workbook(path=path, fmt="xlsx")
    wb.properties = {"creator": ox.properties.creator or "", "modified": str(ox.properties.modified or "")}
    for ws in ox.worksheets:
        sh = Sheet(name=ws.title)
        default_w = ws.sheet_format.defaultColWidth or 8.43
        for key, dim in ws.column_dimensions.items():
            lo = dim.min or openpyxl.utils.column_index_from_string(key)
            hi = dim.max or lo
            for ci in range(lo, min(hi, 60) + 1):
                if dim.width:
                    sh.col_widths_in[ci] = (float(dim.width) * 7.0 + 5.0) / 96.0  # Excel width units -> inches
        for ri, dim in ws.row_dimensions.items():
            if dim.height:
                sh.row_heights_pt[ri] = float(dim.height)
        for row in ws.iter_rows():
            for c in row:
                has_style = any(_side_style(c.border, s) for s in ("left", "right", "top", "bottom"))
                if c.value is None and not has_style:
                    continue
                st = CellStyle()
                st.font_name, st.font_size = c.font.name, float(c.font.sz) if c.font.sz else None
                st.bold, st.italic = bool(c.font.b), bool(c.font.i)
                st.underline = bool(c.font.u)
                st.color = ("#" + c.font.color.rgb[-6:]) if c.font.color is not None and isinstance(c.font.color.rgb, str) else None
                if c.fill and c.fill.fill_type == "solid" and isinstance(c.fill.fgColor.rgb, str):
                    st.fill = "#" + c.fill.fgColor.rgb[-6:]
                st.h_align = c.alignment.horizontal if c.alignment.horizontal not in (None, "general") else None
                st.v_align = c.alignment.vertical
                st.wrap = bool(c.alignment.wrap_text)
                st.shrink = bool(c.alignment.shrink_to_fit)
                sides = [_side_style(c.border, s) for s in ("left", "right", "top", "bottom")]
                if all(s == sides[0] and s for s in sides):
                    st.border = {"thin": "thin", "hair": "hair"}.get(sides[0], sides[0])
                elif any(sides):
                    st.border = ",".join(f"{n}:{s}" for n, s in zip("LRTB", sides) if s)
                st.number_format = c.number_format if c.number_format != "General" else None
                cell = Cell(row=c.row, col=c.column, style=st)
                v = c.value
                if isinstance(v, bool):
                    cell.value, cell.kind = v, "bool"
                elif isinstance(v, dt.datetime):
                    cell.value, cell.kind = (v.date() if v.time() == dt.time(0) else v), "date"
                elif isinstance(v, dt.date):
                    cell.value, cell.kind = v, "date"
                elif isinstance(v, int):
                    cell.value, cell.kind = v, "int"
                elif isinstance(v, float):
                    cell.value, cell.kind = (int(v), "int") if v == int(v) else (v, "float")
                elif v is not None:
                    cell.value, cell.kind = str(v), "string"
                cell.text = "" if v is None else str(v)
                if c.comment:
                    cell.comment = c.comment.text
                sh.cells[(c.row, c.column)] = cell
                if cell.kind != "empty":
                    sh.max_row, sh.max_col = max(sh.max_row, c.row), max(sh.max_col, c.column)
        ps = sh.print
        ps.orientation = ws.page_setup.orientation
        ps.margin_top, ps.margin_bottom = ws.page_margins.top, ws.page_margins.bottom
        ps.margin_left, ps.margin_right = ws.page_margins.left, ws.page_margins.right
        ps.fit_width = ws.page_setup.fitToWidth if ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage else None
        ps.fit_height = ws.page_setup.fitToHeight if ps.fit_width is not None else None
        ps.scale_pct = ws.page_setup.scale
        ps.title_rows = ws.print_title_rows
        ps.freeze = ws.freeze_panes
        sh.merged = [str(m) for m in ws.merged_cells.ranges]
        wb.sheets.append(sh)
    return wb


def read_workbook(path: str) -> Workbook:
    p = path.lower()
    if p.endswith(".ods"):
        return read_ods(path)
    if p.endswith((".xlsx", ".xlsm")):
        return read_xlsx(path)
    raise ValueError(f"unsupported workbook type: {path}")


def select_index_sheet(wb: Workbook, name: str | None = None) -> Sheet:
    """Pick the client index sheet: explicit name, else a sheet named 'Index' holding the header row, else the first with the header."""
    def has_header(sh: Sheet) -> bool:
        return any(str(sh.get(r, 1).value or "").strip() == "Document Type" for r in range(1, 30))
    if name:
        for sh in wb.sheets:
            if sh.name == name:
                return sh
        raise ValueError(f"sheet {name!r} not found; have {[x.name for x in wb.sheets]}")
    for sh in wb.sheets:
        if sh.name == "Index" and has_header(sh):
            return sh
    for sh in wb.sheets:
        if has_header(sh):
            return sh
    return wb.sheets[0]
