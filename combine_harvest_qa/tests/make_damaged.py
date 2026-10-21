"""Build a deliberately damaged XLSX from a donor ODS so the checkers have known defects to catch."""
import sys, datetime as dt
sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.dirname(__file__)))
import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment
from chqa.readers import read_ods

def build(donor_ods: str, out_xlsx: str, rows: int = 40) -> dict:
    sh = read_ods(donor_ods).sheets[0]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"        # defect: sheet name
    for r in range(1, 10 + rows):
        for c in range(1, 10):
            cell = sh.get(r, c); v = cell.value
            ws.cell(r, c, v if cell.kind != "date" else v)
            ws.cell(r, c).font = Font(name="Arial", size=11)               # defect: font
            ws.cell(r, c).border = Border(*(Side(style="hair"),) * 4) if r >= 9 else Border()  # defect: hair borders
            ws.cell(r, c).alignment = Alignment(wrap_text=False)
            if cell.kind == "date": ws.cell(r, c).number_format = "yyyy-mm-dd"  # defect: date format
    ws["F12"] = "2/30/1922"                    # defect: text date
    ws["B13"] = "Unknown"                      # defect: placeholder
    ws["C14"] = None                           # defect: blank substantive
    ws["I15"] = "check this later"             # defect: comment
    ws["D16"] = str(ws["D16"].value)           # defect: Doc No as text
    ws["J9"] = "Extra"; ws["J10"] = "x"        # defect: extra client column
    for c in range(1, 10): ws.cell(12 + 30, c).value = ws.cell(11, c).value  # defect: duplicate identity
    ws.freeze_panes = "A10"                    # defect: freeze
    ws.print_title_rows = "9:9"                # defect: repeated heading
    wb.save(out_xlsx)
    return {"expected_defect_checks": ["sheet_name", "cell_style", "date_format", "invalid_date", "unknown_placeholder",
            "blank_substantive", "unexpected_comment", "docno_stored_as_text", "extra_column", "duplicate_identity",
            "freeze_panes", "print_titles"]}

if __name__ == "__main__":
    print(build(sys.argv[1], sys.argv[2]))
