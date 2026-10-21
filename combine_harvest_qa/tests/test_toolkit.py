"""End-to-end tests on a synthetic donor (no client data). Run: python3 -m pytest -q tests"""
import datetime as dt
import os
import sys
import zipfile

import openpyxl
import pytest
from openpyxl.styles import Alignment, Border, Font, Side

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from chqa.checks_content import check_content  # noqa: E402
from chqa.checks_format import check_format, expected_lines  # noqa: E402
from chqa.checks_package import check_zip, sha256_file  # noqa: E402
from chqa.conform import conform, parse_date  # noqa: E402
from chqa.discover import best_per_role, rank, section_of  # noqa: E402
from chqa.findings import Report  # noqa: E402
from chqa.profile import extract_profile  # noqa: E402

HEAD = ["Document Type", "Grantor", "Grantee", "Doc No", "Book-Page", "Date of Doc", "Rec Date", "Legal Description", "Comments"]
WIDTHS = [3.24, 3.57, 3.6, 0.78, 0.77, 0.81, 0.78, 3.7, 1.0]
THIN = Side(style="thin", color="000000")


def _style(c, bold=False, h="left", wrap=True, shrink=False, border=True):
    c.font = Font(name="Calibri", size=11, bold=bold, color="000000")
    c.alignment = Alignment(horizontal=h, vertical="top", wrap_text=wrap, shrink_to_fit=shrink)
    c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN) if border else Border()


def build_donor(path, rows=6):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Index"
    meta = [("County:", "Campbell"), ("Lands:", "T45N-R76W Section 11: All"), ("Date:", dt.date(2026, 9, 22)),
            ("Starting Date:", "Inception"), ("Date Posted Thru:", dt.date(2026, 4, 22)), ("Indexed By:", "Ryan Gille"),
            ("Project:", "Abstract 4576 Overtime")]
    for r, (k, v) in enumerate(meta, 1):
        ws.cell(r, 1, k); ws.cell(r, 2, v)
        _style(ws.cell(r, 1), bold=True, border=False); _style(ws.cell(r, 2), border=False)
        if isinstance(v, dt.date):
            ws.cell(r, 2).number_format = "m/d/yyyy"
        ws.row_dimensions[r].height = 15
    for i, h in enumerate(HEAD, 1):
        _style(ws.cell(9, i, h), bold=True, h="center"); ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = (WIDTHS[i - 1] * 96 - 5) / 7
    ws.row_dimensions[9].height = 30
    for n in range(rows):
        r = 10 + n
        vals = ["Warranty Deed", "Grantor %d" % n, "Grantee %d" % n, 1000 + n, "0%03d-0%03d" % (n, n), dt.date(1950, 1, 1 + n), dt.date(1950, 2, 1 + n), "NE/4 of 11-45N-76W, aol", None]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            _style(c, h="center" if i in (4, 5, 6, 7) else "left", wrap=i not in (4, 5), shrink=i in (4, 5))
            if isinstance(v, dt.date): c.number_format = "m/d/yyyy"
            if isinstance(v, int): c.number_format = "0"
        ws.row_dimensions[r].height = 15
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.7; ws.page_margins.top = ws.page_margins.bottom = 0.3
    wb.save(path)


@pytest.fixture()
def donor(tmp_path):
    p = str(tmp_path / "donor.xlsx"); build_donor(p)
    return p, extract_profile(p)


def run_checks(path, prof, **kw):
    f, c = Report("format"), Report("content")
    sh = check_format(path, prof, f)
    check_content(path, sh, prof, c, **kw)
    return f, c


def test_donor_passes_its_own_profile(donor):
    p, prof = donor
    f, c = run_checks(p, prof)
    assert f.errors == 0 and c.errors == 0
    assert prof["header_row"] == 9 and [x["header"] for x in prof["columns"]] == HEAD


def damaged(tmp_path, donor_path):
    wb = openpyxl.load_workbook(donor_path); ws = wb.active
    ws["B11"] = "Unknown"; ws["C12"] = None; ws["F13"] = "2/30/1950"; ws["I14"] = "see later"
    ws["D16"] = "1006"; ws["J9"] = "Extra"; ws["J10"] = "x"
    for c in range(1, 10): ws.cell(15, c).value = ws.cell(10, c).value
    ws["A10"].font = Font(name="Arial", size=11); ws.freeze_panes = "A10"
    out = str(tmp_path / "damaged.xlsx"); wb.save(out); return out


def test_checkers_catch_every_planted_defect_with_cell_refs(tmp_path, donor):
    p, prof = donor
    f, c = run_checks(damaged(tmp_path, p), prof)
    cells = {(x.check, x.cell) for x in f.findings + c.findings}
    for want in [("unknown_placeholder", "B11"), ("blank_substantive", "C12"), ("invalid_date", "F13"), ("unexpected_comment", "I14"),
                 ("docno_stored_as_text", "D16"), ("extra_column", "J9"), ("cell_style", "A10"), ("freeze_panes", None),
                 ("duplicate_identity", "row 15")]:
        assert want in cells, want


def test_conform_fixes_style_and_leaves_facts_alone(tmp_path, donor):
    p, prof = donor
    src = damaged(tmp_path, p)
    before = open(src, "rb").read()
    rep = conform(src, prof, str(tmp_path / "out" / "cand.xlsx"), make_ods=False)
    assert open(src, "rb").read() == before                      # input never modified
    f, c = run_checks(rep["xlsx"], prof)
    assert f.errors == 0
    left = {x.check for x in c.errors and c.findings}
    assert left <= {"unknown_placeholder", "blank_substantive", "invalid_date", "duplicate_identity"}   # never invented away
    assert any("Unknown" in n["issue"] for n in rep["needs_human"])
    assert any(ch["kind"] == "docno_text_to_int" for ch in rep["changes"])


def test_ods_roundtrip_parity_when_libreoffice_present(tmp_path, donor):
    import shutil
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice not installed")
    from chqa.checks_parity import check_parity
    p, prof = donor
    rep = conform(damaged(tmp_path, p), prof, str(tmp_path / "o" / "c.xlsx"))
    assert rep["ods"] and os.path.exists(rep["ods"])
    par = Report("parity"); check_parity(rep["xlsx"], rep["ods"], par)
    assert par.errors == 0, par.findings[:3]
    f, _ = run_checks(rep["ods"], prof)
    assert f.errors == 0, [x.message + str(x.cell) for x in f.findings[:3]]


def test_date_parser_rejects_impossible_dates():
    assert parse_date("2/30/1950") is None and parse_date("12/27/2022") == dt.date(2022, 12, 27)


def test_zip_integrity_hash_and_internal_leak(tmp_path):
    z = tmp_path / "pkg.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("S11_Abstract_Index.ods", "x"); zf.writestr("QA_INTERNAL/report_format.json", "{}")
    rep = Report("package"); info = check_zip(str(z), rep, required=[r"Certification_Letter.*\.docx$"])
    checks = {f.check for f in rep.findings}
    assert {"missing_client_file", "internal_in_client_package"} <= checks and info["sha256"] == sha256_file(str(z))
    bad = tmp_path / "bad.zip"; bad.write_bytes(b"not a zip")
    rep2 = Report("package"); check_zip(str(bad), rep2)
    assert rep2.errors == 1


def test_discover_ranks_best_and_flags_foreign_owner():
    files = [{"id": "1", "title": "S13_County.ods", "modifiedTime": "2026-10-17T00:00:00Z", "owner": "me@x"},
             {"id": "2", "title": "S13_County_Abstract_Index_Overtime_100pct", "modifiedTime": "2026-10-20T00:00:00Z", "owner": "other@x"},
             {"id": "3", "title": "S31_Checklist.xlsx", "modifiedTime": "2026-10-18T00:00:00Z", "owner": "me@x"}]
    r = rank(files, {"13", "31"}, me="me@x")
    for t, sec in [("S13_County.ods", "13"), ("S31_Checklist.xlsx", "31"), ("Section_31_T47N-R76W", "31"), ("13-45N-76W_Certification_Letter.docx", "13"),
                   ("47N-76W-31_Johnson_Co_Penterra_Abstract_Index.ods", "31"), ("CH185443 Section 31 — MASTER", "31"), ("Overtime13", None)]:
        assert section_of(t) == sec, t
    assert set(r) == {"13", "31"}
    from chqa.discover import role_of
    assert role_of("S13_BLM_Abstract_Index_Overtime_100pct") == "blm_index" and role_of("S31_Checklist.xlsx") == "checklist"
    stub = rank([{"id": "9", "title": "S13_County_Abstract_Index_100pct", "modifiedTime": "2026-10-21T00:00:00Z", "owner": "me@x", "fileSize": "1294"}], {"13"})
    assert stub["13"][0]["possible_stub"] is True
    b = best_per_role(r)
    assert b["13"]["county_index"]["id"] == "2" and b["13"]["county_index"]["foreign_owner"]


def test_row_height_estimator_basic():
    assert expected_lines("short", 3.0, True) == 1 and expected_lines("a\nb\nc", 3.0, True) == 3
