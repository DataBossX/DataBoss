# -*- coding: utf-8 -*-
import os,sys,re,json,datetime,shutil
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import style
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
from openpyxl.utils import get_column_letter
SCR=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT=os.path.join(SCR,'out'); D=os.path.join(SCR,'src/overtime/Overtime_Examples')
cty=json.load(open(os.path.join(SCR,'qa/S6_County_final_rows.json')))
blm=json.load(open(os.path.join(SCR,'qa/S6_BLM_final_rows.json')))
NC,NB=len(cty),len(blm)
PREP='10/20/2026'; CTY_THRU='10/6/2026'; SRP_THRU='10/13/2026'

# ---------- checklist (donor: 3 bands, 11 cols, Calibri 11, centered, thin borders) ----------
wb=Workbook(); ws=wb.active; ws.title='Abstract checklist'
B=Border(*[Side(style='thin',color='FF000000')]*4)
AL=Alignment(horizontal='center',vertical='center',wrap_text=True)
WH=PatternFill('solid',fgColor='FFFFFFFF')
bands=[(2,['County Index','County Documents','Plat','STR online search','Compare index to Uploads','BLM File','SRP','State File','Lease Report if applicable','Reviewed SRP vs BLM file or State file','Certification Letter'],
          ['Yes','Available','Indexed',CTY_THRU,'%d entries'%NC,'WYW-149799',SRP_THRU,'N/A','N/A','Qualified','Unsigned']),
       (5,['Abstract Excel Sheet','County','Lands Covered','Date Prepared','Starting Date','Date Reviewed','Landman Name','Project / Prospect'],
          ['Yes','Johnson','T47N-R76W Section 6: All',PREP,'Inception',PREP,'Ryan Gille','Combine-Harvest CH185443 / AFE 210275']),
       (8,['Certification Letter','Client','STR & County','Contents for County Records','Contents for BLM / State','SRP Info','Preparation Date'],
          ['Prepared; unsigned','Anschutz Exploration Corporation','T47N-R76W Section 6: All, Johnson County, Wyoming',
           'County index: %d entries; source cutoff %s'%(NC,CTY_THRU),
           'WYW-149799: %d entries; Parts 1-2'%NB,'WYW-149799 SRP: %s'%SRP_THRU,PREP])]
for r0,hdr,val in bands:
    for j,(h,v) in enumerate(zip(hdr,val),1):
        c=ws.cell(r0,j,h); c.font=style.FB; c.alignment=AL; c.border=B; c.fill=WH
        d=style.to_date(v)
        c2=ws.cell(r0+1,j, d if d else v); c2.font=style.F; c2.alignment=AL; c2.border=B; c2.fill=WH
        if d: c2.number_format='M/D/YYYY'
for col,w in zip('ABCDEFGHIJK',[14.0,18.0,29.3,30.1,38.0,32.7,28.0,18.0,17.7,18.9,15.9]):
    ws.column_dimensions[col].width=w
ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.page_margins.left=0.7; ws.page_margins.right=0.7; ws.page_margins.top=0.75; ws.page_margins.bottom=0.75
chk=os.path.join(OUT,'CH185443_S06_Abstract_Checklist_6-47N-76W_20261020.xlsx'); wb.save(chk)
print('checklist saved')

# ---------- certification letter from donor template ----------
import docx
src=os.path.join(D,'S13_13-45N-76W_Certification_Letter.docx')
doc=docx.Document(src)
REP=[('Township 45 North, Range 76 West','Township 47 North, Range 76 West'),
     ('Section 13: All','Section 6: All'),
     ('Campbell County Clerk','Johnson County Clerk'),
     ('Campbell County Assessor','Johnson County Assessor'),
     ('tract indices of Campbell County','tract indices of Johnson County'),
     ('Campbell County, Wyoming','Johnson County, Wyoming'),
     ('WYW-047318','WYW-149799'),
     ('4/22/2026',CTY_THRU),('8/4/2026',SRP_THRU)]
def fix(p):
    t=p.text
    if not t.strip(): return
    n=t
    for a,b in REP: n=n.replace(a,b)
    n=re.sub(r'Prepared on \d{1,2}/\d{1,2}/\d{4}','Prepared on '+PREP,n)
    if n!=t:
        for i,r in enumerate(p.runs):
            r.text = n if i==0 else ''
for p in doc.paragraphs: fix(p)
for tb in doc.tables:
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs: fix(p)
# qualification paragraph before the signature rule
anchor=None
for i,p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('Prepared on'): anchor=p; break
QUAL=('Qualifications: This abstract is delivered UNSIGNED and QUALIFIED pending examiner acceptance. '
      'The county index is posted through %s as shown on the Johnson County tract index; that date is a county '
      'record-coverage date and is not a search-through date to the preparation date. The federal index is posted '
      'through the %s Serial Register Page run date for WYW-149799 / WYWY105544892. Items remaining open at '
      'delivery are itemized in the accompanying QA register and are not represented as resolved.'%(CTY_THRU,SRP_THRU))
if anchor is not None:
    newp=anchor.insert_paragraph_before(QUAL)
    try: newp.style=anchor.style
    except Exception: pass
cert=os.path.join(OUT,'CH185443_S06_6-47N-76W_Certification_Letter_UNSIGNED.docx')
doc.save(cert); print('cert saved')
