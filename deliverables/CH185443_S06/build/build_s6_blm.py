# -*- coding: utf-8 -*-
import json, sys, os, re
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import norm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

SCR=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src=json.load(open(os.path.join(SCR,'qa/S6_BLM_cells.json')))['Index']
headers=[c['text'] for c in src[8]]
data=src[9:]
assert headers==['Document Type','Grantor','Grantee','File No','Page','Date of Doc','Rec Date','Comments','Legal Description','Part'], headers
assert len(data)==61

W=[1.75,2.10,2.10,0.75,0.60,0.95,0.95,0.85,2.85,0.50]
META=[('County:','Johnson'),
      ('Lands:','T47N-R76W Section 6 - Federal - WYW-149799 / WYWY105544892'),
      ('Date:','10/20/2026'),
      ('Starting Date:','2/1/1999'),
      ('Date Posted Thru:','10/13/2026'),
      ('Indexed By:','Ryan Gille'),
      ('Project:','Combine-Harvest CH185443 / AFE 210275 - WYW-149799')]
norm.qa(4,'Starting Date','(blank)','2/1/1999','SRP WYWY105544892: CASE ESTABLISHED 02/01/1999 and lease Effective Date 02/01/1999')
norm.qa(5,'Date Posted Thru','(blank)','10/13/2026','SRP WYWY105544892 Run Date/Time 10/13/2026 16:11 PM - federal records current through SRP run date; county 10/6/2026 date NOT applied to federal file')

def g(r,i): return (r[i]['text'] if len(r)>i and r[i]['text'] else '')

LEGAL_SPLIT=re.compile(r',\s*(?=(?:Lots?|All|N/2|S/2|E/2|W/2|NE/4|NW/4|SE/4|SW/4|SE/4 NW/4|Tract)\b)')
def norm_blm_legal(v,row):
    s=(v or '').strip()
    if not s: return ''
    aol = s.endswith(', aol')
    body = s[:-5] if aol else s
    parts=[p.strip().rstrip(',') for p in re.split(r'\n', body)]
    out=[]
    for p in parts: out.extend([q.strip() for q in LEGAL_SPLIT.split(p) if q.strip()])
    txt=',\n'.join(out)
    if aol: txt+=', aol'
    if txt!=s: norm.qa(row,'Legal Description',s,txt,'each proved legal placed on its own comma-terminated line; ", aol" retained once')
    return txt

rows=[]
for k,r in enumerate(data):
    xl=10+k
    gtee=g(r,2)
    if gtee.strip().lower()=='unknown':
        gtee='US Realm Powder River, LLC'
        norm.qa(xl,'Grantee','Unknown',gtee,
          'literal Unknown replaced from SRP WYWY105544892: US REALM POWDER RIVER LLC shown LESSEE 100.0000%; name change from Moriah Powder River LLC recognized 06/15/2020, so lessee of record on 6/16/2022 notice. SOURCE-DERIVED - confirm against case page 53 image at examiner review')
    rows.append([g(r,0), norm.norm_party(g(r,1),xl,'Grantor'), norm.norm_party(gtee,xl,'Grantee'),
                 None, int(g(r,4)) if g(r,4).isdigit() else g(r,4),
                 norm.norm_date(g(r,5),xl,'Date of Doc'), norm.norm_date(g(r,6),xl,'Rec Date'),
                 None, norm_blm_legal(g(r,8),xl), int(g(r,9)) if g(r,9).isdigit() else g(r,9)])
rows.sort(key=lambda x:(x[9], x[4]))

wb=Workbook(); ws=wb.active; ws.title='Index'
F=Font(name='Calibri',size=11,color='FF000000'); FB=Font(name='Calibri',size=11,bold=True,color='FF000000')
thin=Side(style='thin',color='FF000000'); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
AL=Alignment(horizontal='left',vertical='top',wrap_text=True); NOFILL=PatternFill(fill_type=None)

for i,(lab,val) in enumerate(META,1):
    ws.cell(i,1,lab).font=FB
    ws.cell(i,2,val).font=F
    for col in range(1,11):
        c=ws.cell(i,col); c.alignment=Alignment(horizontal='left',vertical='top',wrap_text=False); c.fill=NOFILL
        if c.font.name!='Calibri': c.font=F
    ws.row_dimensions[i].height=15.0
ws.row_dimensions[8].height=15.0
for col in range(1,11):
    c=ws.cell(8,col); c.font=F; c.fill=NOFILL
for j,h in enumerate(headers,1):
    c=ws.cell(9,j,h); c.font=FB; c.alignment=AL; c.border=BORD; c.fill=NOFILL
ws.row_dimensions[9].height=15.75
for k,row in enumerate(rows):
    xl=10+k
    for j,v in enumerate(row,1):
        c=ws.cell(xl,j, v if v not in ('',None) else None)
        c.font=F; c.alignment=AL; c.border=BORD; c.fill=NOFILL
        if j in (5,10) and isinstance(v,int): c.number_format='0'
    ws.row_dimensions[xl].height=None
for j,w in enumerate(W,1):
    ws.column_dimensions[get_column_letter(j)].width=round(w*96.0/7.0,2)
ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=5
ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
for a in ('left','right','top','bottom'): setattr(ws.page_margins,a,0.3)
ws.page_margins.header=0.0; ws.page_margins.footer=0.0
ws.print_title_rows='1:9'; ws.print_area='A1:J%d'%(9+len(rows)); ws.freeze_panes='A10'
out=os.path.join(SCR,'out/CH185443_S06_BLM_WYW-149799_Abstract_Index_20261020.xlsx')
wb.save(out)
json.dump(norm.QA,open(os.path.join(SCR,'qa/S6_BLM_QA_register.json'),'w'),indent=1)
print('blm rows',len(rows),'qa',len(norm.QA)); print('saved',out)
