# -*- coding: utf-8 -*-
import json, sys, os, math
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import norm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

AUTOFIT=True
SCR=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
cells=json.load(open(os.path.join(SCR,'qa/S6_County_cells.json')))
src=cells['County index']
meta=src[:7]; headers=[c['text'] for c in src[8]]; data=src[9:]
assert headers[0]=='Document Type' and len(headers)==9, headers
assert len(data)==69, len(data)

WIDTH_IN=[1.389,1.9445,1.9445,0.972,1.222,1.028,1.028,2.9165,0.9028]
PREP_DATE='10/20/2026'
POSTED_THRU='10/6/2026'

META=[('County:','Johnson'),
      ('Lands:','T47N-R76W Section 6'),
      ('Date:',PREP_DATE),
      ('Starting Date:',''),            # not proved - QA hold
      ('Date Posted Thru:',POSTED_THRU),# proved: county tract index p.5 "Land Validated Through: 10/6/2026"
      ('Indexed By:','Ryan Gille'),
      ('Project:','Combine-Harvest CH185443 / AFE 210275 - Section 6')]
norm.qa(4,'Starting Date','Unknown','blanked','county inception not established in supplied sources - OPEN RESEARCH ITEM')
norm.qa(5,'Date Posted Thru','Unknown',POSTED_THRU,'proved from original 7-page county tract index PDF, physical page 5: "Land Validated Through: 10/6/2026" (32 of 32 results, T47N R76W Sec 06). 10/8/2026 5:07 PM is print timestamp only')
norm.qa(1,'banner','WORKING DRAFT — INCOMPLETE SOURCE INTAKE AND REVIEW','removed from client grid','not a donor metadata field; qualification carried in unsigned certificate + QA register')
norm.qa(6,'Indexed By','Working draft — review pending','Ryan Gille','donor metadata field value')

def getc(r,i):
    return (r[i]['text'] if len(r)>i and r[i]['text'] else '')

out_rows=[]
for k,r in enumerate(data):
    xl=10+k
    dt   = getc(r,0).replace('--','—') if False else getc(r,0)
    gtor = norm.norm_party(getc(r,1), xl, 'Grantor')
    gtee = norm.norm_party(getc(r,2), xl, 'Grantee')
    dno  = norm.norm_docno(getc(r,3), xl)
    bp   = norm.norm_bookpage(getc(r,4), xl)
    dod  = norm.norm_date(getc(r,5), xl, 'Date of Doc')
    rec  = norm.norm_date(getc(r,6), xl, 'Rec Date')
    leg  = norm.norm_legal(getc(r,7), xl)
    com  = getc(r,8)
    if com and com.strip().lower()!='missing image':
        norm.qa(xl,'Comments',com,'blanked','client Comments permits only blank or exactly "missing image"'); com=''
    out_rows.append([dt,gtor,gtee,dno,bp,dod,rec,leg,com])

wb=Workbook(); ws=wb.active; ws.title='Index'
F=Font(name='Calibri',size=11,color='FF000000')
FB=Font(name='Calibri',size=11,bold=True,color='FF000000')
thin=Side(style='thin',color='FF000000')
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
AL=Alignment(horizontal='left',vertical='top',wrap_text=True)
NOFILL=PatternFill(fill_type=None)

for i,(lab,val) in enumerate(META,1):
    ws.cell(i,1,lab).font=FB
    c=ws.cell(i,2,val); c.font=F
    for col in range(1,10):
        cc=ws.cell(i,col); cc.alignment=Alignment(horizontal='left',vertical='top',wrap_text=False); cc.fill=NOFILL
        if cc.font.name!='Calibri': cc.font=F
    ws.row_dimensions[i].height=15.0
ws.row_dimensions[8].height=15.0
for col in range(1,10):
    c=ws.cell(8,col); c.font=F; c.fill=NOFILL

for j,h in enumerate(headers,1):
    c=ws.cell(9,j,h); c.font=FB; c.alignment=AL; c.border=BORD; c.fill=NOFILL
ws.row_dimensions[9].height=15.75

for k,row in enumerate(out_rows):
    xl=10+k; lines=1
    for j,v in enumerate(row,1):
        c=ws.cell(xl,j, v if v!='' else None)
        c.font=F; c.alignment=AL; c.border=BORD; c.fill=NOFILL
        if j==4 and isinstance(v,int): c.number_format='0'
        lines=max(lines, norm.wrap_lines(v, WIDTH_IN[j-1]))
    if AUTOFIT:
        ws.row_dimensions[xl].height=None
    else:
        ws.row_dimensions[xl].height=round(lines*14.4,1)

for j,w in enumerate(WIDTH_IN,1):
    ws.column_dimensions[get_column_letter(j)].width=round(w*96.0/7.0,2)

ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.page_setup.orientation='landscape'
ws.page_setup.paperSize=5           # Legal 8.5x14
ws.page_setup.fitToWidth=1
ws.page_setup.fitToHeight=0
ws.page_margins.left=0.3; ws.page_margins.right=0.3
ws.page_margins.top=0.3;  ws.page_margins.bottom=0.3
ws.page_margins.header=0.0; ws.page_margins.footer=0.0
ws.print_title_rows='1:9'
ws.print_area='A1:I%d'%(9+len(out_rows))
ws.freeze_panes='A10'

os.makedirs(os.path.join(SCR,'out'),exist_ok=True)
xlsx=os.path.join(SCR,'out/CH185443_S06_County_Abstract_Index_20261020.xlsx')
wb.save(xlsx)
json.dump(norm.QA, open(os.path.join(SCR,'qa/S6_County_QA_register.json'),'w'), indent=1)
json.dump(out_rows, open(os.path.join(SCR,'qa/S6_County_final_rows.json'),'w'), indent=1)
print('rows',len(out_rows),'qa entries',len(norm.QA))
print('saved',xlsx)
