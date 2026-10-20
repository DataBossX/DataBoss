# -*- coding: utf-8 -*-
import os,json,zipfile,hashlib,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
SCR=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT=os.path.join(SCR,'out'); PKG=os.path.join(SCR,'pkg'); os.makedirs(PKG,exist_ok=True)
F=Font(name='Calibri',size=11,color='FF000000'); FB=Font(name='Calibri',size=11,bold=True,color='FF000000')
B=Border(*[Side(style='thin',color='FF000000')]*4)
AL=Alignment(horizontal='left',vertical='top',wrap_text=True)
NF=PatternFill(fill_type=None)

# ---------- QA register workbook (INTERNAL) ----------
qc=json.load(open(os.path.join(SCR,'qa/S6_County_QA_register.json')))
qb=json.load(open(os.path.join(SCR,'qa/S6_BLM_QA_register.json')))
gates=json.load(open(os.path.join(SCR,'qa/QA_GATES.json')))
wb=Workbook()
ws=wb.active; ws.title='Field actions'
hdr=['Index','Row','Field','Original value','Action / new value','Basis']
for j,h in enumerate(hdr,1):
    c=ws.cell(1,j,h); c.font=FB; c.alignment=AL; c.border=B; c.fill=NF
r=2
for tag,src in (('County',qc),('BLM WYW-149799',qb)):
    for e in src:
        for j,v in enumerate([tag,e['row'],e['field'],e['original'],e['action'],e['reason']],1):
            c=ws.cell(r,j,v); c.font=F; c.alignment=AL; c.border=B; c.fill=NF
        r+=1
for col,w in zip('ABCDEF',[16,7,18,46,40,96]): ws.column_dimensions[col].width=w
ws2=wb.create_sheet('QA gates')
for j,h in enumerate(['Gate','Check','Status','Detail'],1):
    c=ws2.cell(1,j,h); c.font=FB; c.alignment=AL; c.border=B; c.fill=NF
for i,g in enumerate(gates,2):
    for j,v in enumerate([g['gate'],g['check'],g['status'],g['detail']],1):
        c=ws2.cell(i,j,v); c.font=F; c.alignment=AL; c.border=B; c.fill=NF
for col,w in zip('ABCD',[7,62,10,60]): ws2.column_dimensions[col].width=w

ws3=wb.create_sheet('Open items')
OPEN=[['OPEN-1','County','Lode location notices (Book A069) - exact aliquot/target allocation not stated on the instrument face',
       '36 rows carry the section call "6-47N-76W" with no aliquot. Section applicability proved; aliquot NOT proved. Not invented.',
       'Re-examine original A069 county images; if aliquot is genuinely absent, retain section-only call'],
      ['OPEN-2','County','Doc 48736 and Doc 100046 - Date of Doc not proved on the supplied image',
       'Literal "Unknown" removed; cell left genuinely blank. Blanking does NOT close the research issue.',
       'Re-pull instrument face for 0309-0228–0235 and 86A39-0218–0219'],
      ['OPEN-3','County','Starting Date recorded as "Inception" per Overtime donor convention',
       'Counts sheet of the source workbook still shows county inception/certified-through as not established.',
       'Confirm county record inception with the Clerk before signature'],
      ['OPEN-4','BLM','Grantee on Part 2 page 53 derived from the SRP, not from the case page image',
       'Filled as "US Realm Powder River, LLC" from SRP lessee of record; flagged source-derived.',
       'Confirm against WYW-149799 Part 2 page 53 original image'],
      ['OPEN-5','BLM','WYW-149799 original case-file PDFs were not re-opened in this pass',
       'BLM rows carried forward from the verified S06_WYW149799_Combined_Parts_1-2 production file (SHA recorded) and '
       'cross-checked against the 10/13/2026 SRP. No independent page-by-page image re-verification was performed here.',
       'Image-verify Parts 1-2 page by page before signature'],
      ['OPEN-6','County','Row count 69 vs county tract index "32 of 32 results"',
       'The 69 rows are the full qualified historical recording set carried from the source workbook; the tract index '
       'screen reports 32 current hits. These count different things and were NOT reconciled in this pass.',
       'Reconcile 69 abstract rows against the 32 tract-index hits'],
      ['OPEN-7','Both','Certification letter is UNSIGNED and QUALIFIED',
       'Per standing owner instruction, no completion or certification claim is made.',
       'Examiner acceptance and signature']]
for j,h in enumerate(['ID','Index','Issue','Status at delivery','Required to close'],1):
    c=ws3.cell(1,j,h); c.font=FB; c.alignment=AL; c.border=B; c.fill=NF
for i,row in enumerate(OPEN,2):
    for j,v in enumerate(row,1):
        c=ws3.cell(i,j,v); c.font=F; c.alignment=AL; c.border=B; c.fill=NF
for col,w in zip('ABCDE',[10,12,56,74,52]): ws3.column_dimensions[col].width=w
qa_xlsx=os.path.join(OUT,'CH185443_S06_QA_REGISTER_INTERNAL_20261020.xlsx'); wb.save(qa_xlsx)

# ---------- ZIPs ----------
CLIENT=['CH185443_S06_County_Abstract_Index_20261020.ods',
        'CH185443_S06_County_Abstract_Index_20261020.xlsx',
        'CH185443_S06_County_Abstract_Index_20261020.pdf',
        'CH185443_S06_BLM_WYW-149799_Abstract_Index_20261020.ods',
        'CH185443_S06_BLM_WYW-149799_Abstract_Index_20261020.xlsx',
        'CH185443_S06_BLM_WYW-149799_Abstract_Index_20261020.pdf',
        'CH185443_S06_Abstract_Checklist_6-47N-76W_20261020.xlsx',
        'CH185443_S06_Abstract_Checklist_6-47N-76W_20261020.pdf',
        'CH185443_S06_6-47N-76W_Certification_Letter_UNSIGNED.docx',
        'CH185443_S06_6-47N-76W_Certification_Letter_UNSIGNED.pdf']
def sha(p):
    h=hashlib.sha256()
    with open(p,'rb') as fh:
        for ch in iter(lambda: fh.read(1<<20), b''): h.update(ch)
    return h.hexdigest()
man=[]
czip=os.path.join(PKG,'CH185443_S06_CLIENT_PACKAGE_20261020.zip')
with zipfile.ZipFile(czip,'w',zipfile.ZIP_DEFLATED) as z:
    for n in CLIENT:
        p=os.path.join(OUT,n); z.write(p,n); man.append({'file':n,'bytes':os.path.getsize(p),'sha256':sha(p)})
INTERNAL=[(qa_xlsx,'CH185443_S06_QA_REGISTER_INTERNAL_20261020.xlsx'),
          (os.path.join(SCR,'qa/QA_GATES.json'),'QA_GATES.json'),
          (os.path.join(SCR,'qa/S6_County_QA_register.json'),'S6_County_field_actions.json'),
          (os.path.join(SCR,'qa/S6_BLM_QA_register.json'),'S6_BLM_field_actions.json'),
          (os.path.join(SCR,'src/SRP_WYW149799.pdf'),'source/SRP_WYW149799_10-13-2026.pdf'),
          (os.path.join(SCR,'src/S6_County.ods'),'source/ORIGINAL_S6_County.ods'),
          (os.path.join(SCR,'src/S06_WYW149799_Combined_Parts_1-2.ods'),'source/ORIGINAL_S06_WYW149799_Combined_Parts_1-2.ods')]
iman=[]
izip=os.path.join(PKG,'CH185443_S06_QA_INTERNAL_20261020.zip')
with zipfile.ZipFile(izip,'w',zipfile.ZIP_DEFLATED) as z:
    for p,n in INTERNAL:
        z.write(p,n); iman.append({'file':n,'bytes':os.path.getsize(p),'sha256':sha(p)})
manifest={'project':'Combine-Harvest CH185443 / AFE 210275','section':'Section 6, T47N-R76W, Johnson County, Wyoming',
 'prepared':'10/20/2026','county_posted_thru':'10/6/2026','federal_posted_thru':'10/13/2026',
 'county_rows':69,'blm_rows':61,'status':'QUALIFIED - UNSIGNED - pending examiner acceptance',
 'client_zip':{'name':os.path.basename(czip),'bytes':os.path.getsize(czip),'sha256':sha(czip),'files':man},
 'qa_internal_zip':{'name':os.path.basename(izip),'bytes':os.path.getsize(izip),'sha256':sha(izip),'files':iman},
 'source_hashes':{'S6_County.ods':sha(os.path.join(SCR,'src/S6_County.ods')),
   'S06_WYW149799_Combined_Parts_1-2.ods':sha(os.path.join(SCR,'src/S06_WYW149799_Combined_Parts_1-2.ods')),
   'SRP_WYW149799.pdf':sha(os.path.join(SCR,'src/SRP_WYW149799.pdf')),
   'Overtime_Examples_11_13_15.zip':sha(os.path.join(SCR,'src/Overtime_Examples_11_13_15.zip'))}}
json.dump(manifest,open(os.path.join(PKG,'CH185443_S06_MANIFEST_20261020.json'),'w'),indent=1)
# gate 5: zip integrity
for z in (czip,izip):
    with zipfile.ZipFile(z) as zz:
        assert zz.testzip() is None, z
print('CLIENT ZIP',os.path.basename(czip),os.path.getsize(czip),'files',len(man))
print('QA ZIP    ',os.path.basename(izip),os.path.getsize(izip),'files',len(iman))
print('CRC OK for both')
