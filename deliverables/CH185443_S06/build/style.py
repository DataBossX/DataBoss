# -*- coding: utf-8 -*-
"""Donor-measured Overtime styling (Overtime_Examples_11_13_15.zip)."""
import datetime, re
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

F   = Font(name='Calibri', size=11, color='FF000000')
FB  = Font(name='Calibri', size=11, bold=True, color='FF000000')
_th = Side(style='thin', color='FF000000')
BORD= Border(left=_th, right=_th, top=_th, bottom=_th)
AL  = Alignment(horizontal='left', vertical='top', wrap_text=True)       # owner rule: uniform left/top
ALH = Alignment(horizontal='center', vertical='center', wrap_text=True)  # donor header: centered
ALM = Alignment(horizontal='left', vertical='top', wrap_text=False)
NOFILL = PatternFill(fill_type=None)
CM = 1/2.54

COUNTY_W_CM = [8.22854166666667,9.07520833333333,9.18104166666667,1.984375,1.95791666666667,
               2.06375,1.984375,9.445625,1.95791666666667]
BLM_W_CM    = [8.016875,8.334375,8.54604166666667,1.349375,0.978958333333333,2.06375,
               1.984375,1.95791666666667,8.78416666666667,0.846666666666667]

def to_date(s):
    m = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', (s or '').strip())
    if not m: return None
    return datetime.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))

def apply_page(ws, ncols, nrows, fit=True):
    ws.page_setup.orientation = 'landscape'
    if fit:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth  = 1
        ws.page_setup.fitToHeight = 0
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
    ws.page_margins.left = 0.7; ws.page_margins.right = 0.7      # donor
    ws.page_margins.top  = 0.3; ws.page_margins.bottom = 0.3     # donor
    ws.page_margins.header = 0.0; ws.page_margins.footer = 0.0
    ws.print_title_rows = '1:9'
    ws.print_area = 'A1:%s%d' % (get_column_letter(ncols), nrows)
    ws.freeze_panes = 'A10'

def set_widths(ws, widths_cm):
    for j, w in enumerate(widths_cm, 1):
        ws.column_dimensions[get_column_letter(j)].width = round(w*CM*96.0/7.0, 2)

def write_meta(ws, meta, ncols):
    for i,(lab,val) in enumerate(meta,1):
        c=ws.cell(i,1,lab); c.font=FB; c.alignment=ALM; c.fill=NOFILL
        d=to_date(val) if isinstance(val,str) else None
        c2=ws.cell(i,2, d if d else val); c2.font=F; c2.alignment=ALM; c2.fill=NOFILL
        if d: c2.number_format='M/D/YYYY'
        for col in range(3,ncols+1):
            cc=ws.cell(i,col); cc.font=F; cc.alignment=ALM; cc.fill=NOFILL
        ws.row_dimensions[i].height=15.0
    for col in range(1,ncols+1):
        cc=ws.cell(8,col); cc.font=F; cc.alignment=ALM; cc.fill=NOFILL
    ws.row_dimensions[8].height=15.0

def write_headers(ws, headers):
    for j,h in enumerate(headers,1):
        c=ws.cell(9,j,h); c.font=FB; c.alignment=ALH; c.border=BORD; c.fill=NOFILL
    ws.row_dimensions[9].height=15.75

def write_cell(ws, r, c, v, datecol=False, numcol=False):
    cell=ws.cell(r,c)
    if v in ('', None):
        cell.value=None
    elif datecol:
        d=to_date(v) if isinstance(v,str) else v
        if d: cell.value=d; cell.number_format='M/D/YYYY'
        else: cell.value=v
    else:
        cell.value=v
        if numcol and isinstance(v,int): cell.number_format='0'
    cell.font=F; cell.alignment=AL; cell.border=BORD; cell.fill=NOFILL
    return cell
