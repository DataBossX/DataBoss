#!/usr/bin/env python3
"""Isolated Section 19 federal PROPOSAL workbook (not a package) + QA receipt + hashes + zip."""
import os,sys,json,csv,hashlib,zipfile,subprocess,glob
from datetime import datetime,timezone
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side

OUT=sys.argv[1]; PREP="2026-10-19"; UNK="Unknown"
NOW=datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

HDR=["Document Type","Grantor","Grantee","File No","Page","Date of Doc","Rec Date","Comments","Legal Description","Part"]
NL=",\n"
# Document Type, Grantor, Grantee, Page, Date of Doc, Rec Date, Legal Description, Part
ROWS=[
 ["Oil and Gas Lease","United States of America","Yates Petroleum Corporation,\nYates Drilling Company,\nABO Petroleum Corporation,\nMyco Industries, Inc.",
  UNK,"2/1/1994",UNK,"T47N R76W Section 19: Lots 8, 9, 13, 14, 16, 20 (lease WYW-131225; record title also Sec 7 Lots 27, 28, 33, 34, Sec 30 Lot 5, T47N R77W Sec 33 Lot 11 - not Section 19 facts)",UNK],
 ["Communitization Agreement","United States of America","Aztec Gas & Oil Corp.",
  "82","9/2/1994",UNK,"T47N R76W Section 19: Lots 9 (SWNW) and 16 (NWSW), containing 35.93 acres and 38.18 acres, respectively, as corrected 9/14/2015; as originally executed SW/4 NW/4 (Lot 2) and NW/4 SW/4 (Lot 14), 40.21 and 39.98 acres; Sussex Formation only (CA WYW-136880, Tract 2)","1"],
 ["Correction Letter","United States of America","Nonsuch Natural Gas, Inc.",
  UNK,"9/14/2015",UNK,"T47N R76W Section 19: Lots 9 (SWNW) and 16 (NWSW), containing 35.93 acres and 38.18 acres, respectively (corrects CA WYW-136880 Tract 2)","1"],
 ["Designation of Successor Operator","United States of America","Nonsuch Natural Gas, Inc.",
  UNK,"4/4/2019",UNK,"CA WYW-136880 embracing Lots 9 and 16 of Section 19 and Lots 8, 9 and 16 of Section 24, \"T. 44 N., R. 76 W.\" as printed in the source recital (operative township T47N); approved effective 10/1/2011","1"],
 ["Assignment","Yates Petroleum Corporation","Yates Drilling Company,\nABO Petroleum Corporation,\nMyco Industries, Inc.",
  UNK,UNK,UNK,"Lease-wide assignment of 90% of record title in WYW-131225; approved 10/19/1994 effective 6/1/1994",UNK],
 ["Decision","United States of America",UNK,
  UNK,"7/10/1998",UNK,"Conformed to survey/resurvey, effective 8/17/1995 (lease WYW-131225, lease-wide)",UNK],
 ["Decision","United States of America",UNK,
  UNK,"7/17/2010",UNK,"Conformed to survey/resurvey, effective 9/29/2005 (lease WYW-131225, lease-wide)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-168834, Tract 01, 38.1800 acres, 49.862900% interest; committed effective 1/1/2006 (lease WYW-131225)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-164204, Tract 01, 40.4600 acres, 49.913600% interest; committed effective 9/1/2005; CLOSED, agreement terminated effective 3/6/2023 (lease WYW-131225)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-175812, Tract 02, 40.4900 acres, 49.963000% interest; committed effective 8/1/2007 (lease WYW-131225)",UNK],
 ["Unit Agreement","United States of America",UNK,
  UNK,"3/15/2001",UNK,"Unit WYW-150648X, 40.5900 acres; lease in unit, uncommitted (lease WYW-131225)",UNK],
 ["Unit Agreement","United States of America",UNK,
  "1",UNK,UNK,"North McRae Gap Unit WYW-179231X, approved 10/19/2010; Section 19 applicability and acreage Unknown; UNCORROBORATED - absent from Serial Register Page recapitulation (lease WYW-131225)","4"],
 ["Unit Agreement","United States of America",UNK,
  UNK,UNK,UNK,"Tractor (Deep) Unit WYW-191271X, committed 10/18/2021, 471.7600 acres fully committed (whole-lease, not Section 19 acreage); lease segregated into WYW-191554 (lease WYW-131225)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYWY-106695248, Tract 007, 40.5900 acres, 1.006200% interest; committed effective 5/1/2024, approved 12/22/2025 (lease WYW-131225)",UNK],
 ["Oil and Gas Lease","United States of America","Yates Petroleum Corporation, et al.",
  UNK,"10/1/1994",UNK,"T47N R76W Section 19: Lots 5, 6, 7, 10, 11, 12, 17 (lease WYW-133620; lease also covers Sec 18 Lots 21-36, Sec 30 Lot 12, T47N R77W Sec 23 Lots 3-6 - not Section 19 facts). Record title reads \"(No Records Found)\"",UNK],
 ["Lease Stipulation","United States of America",UNK,
  UNK,UNK,UNK,"T47N R76W Section 19: NE, E2NW (and Section 20: NW). Timing limitation stipulation February 1 to July 31, protecting sage grouse nesting habitat (lease WYW-133620)",UNK],
 ["Decision","United States of America",UNK,
  UNK,"11/1/1998",UNK,"Lease suspended, operations only, Section 17, payment required; suspension lifted 8/1/2003 (lease WYW-133620, lease-wide)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-168834, Tract 02, 38.3900 acres, 50.137100% interest; committed effective 1/1/2006 (lease WYW-133620)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-178851, Tract 02, 40.4500 acres, 50.130100% interest; effective 10/1/2009; CLOSED, agreement terminated effective 7/15/2020 (lease WYW-133620)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYW-164204, Tract 02, 40.6000 acres, 50.086400% interest; committed effective 9/1/2005; CLOSED, agreement terminated effective 3/6/2023 (lease WYW-133620)",UNK],
 ["Unit Agreement","United States of America",UNK,
  UNK,"3/15/2001",UNK,"Unit WYW-150648X, 161.4400 acres; lease in unit, uncommitted (lease WYW-133620)",UNK],
 ["Unit Agreement","United States of America",UNK,
  UNK,UNK,UNK,"Tractor (Deep) Unit WYW-191271X, committed 10/18/2021, 1,094.7600 acres fully committed (whole-lease, not Section 19 acreage); lease segregated into WYW-191555 (lease WYW-133620)",UNK],
 ["Communitization Agreement","United States of America",UNK,
  UNK,UNK,UNK,"CA WYWY-106808441, Tract 08, 161.4400 acres, 2.371400% interest; effective 4/1/2026; PENDING, not authorized (lease WYW-133620)",UNK],
 ["Decision","United States of America",UNK,
  UNK,"7/10/1998",UNK,"Conformed to survey/resurvey, effective 8/17/1995; further conformance 8/17/2010 effective 9/29/2005 (lease WYW-133620, lease-wide)",UNK],
]
META=[
 ["PENTERRA RESOURCES - CH185443 COMBINE-HARVEST"],
 ["SECTION 19, TOWNSHIP 47 NORTH, RANGE 76 WEST, 6th P.M., JOHNSON COUNTY, WYOMING"],
 ["FEDERAL INDEX - ISOLATED PROPOSAL WORKBOOK (NOT A PACKAGE, NOT A CLIENT DELIVERABLE)"],
 ["AFE 210275   |   Prepared October 19, 2026   |   Research cutoff: Unknown   |   Currentness: Unknown"],
 ["Source examiner: Claude (claude-opus-5). Role: Section 19 federal source examiner. Promotes nothing."],
 ["Every row is a PROPOSAL pending adjudication. Unknown means not proved by a cited original. File No is intentionally blank."],
 ["Typography named Calibri 11; rendered with Carlito 11 (metric-compatible substitute). Calibri not installed; MS Office parity NOT tested."],
]
wb=Workbook(); ws=wb.active; ws.title="Federal"
thin=Side(style="thin",color="FF000000"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
white=PatternFill("solid",fgColor="FFFFFFFF")
for i,m in enumerate(META,start=1):
    c=ws.cell(row=i,column=1,value=m[0]); c.font=Font(name="Calibri",size=11,bold=(i<=3))
for j,h in enumerate(HDR,start=1):
    c=ws.cell(row=9,column=j,value=h)
    c.font=Font(name="Calibri",size=11,bold=True); c.fill=white; c.border=bd
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
for r,row in enumerate(ROWS,start=10):
    dt,gr,ge,pg,dod,rec,legal,part=row
    vals=[dt,gr,ge,None,pg,dod,rec,None,legal,part]   # File No blank, Comments blank
    for j,v in enumerate(vals,start=1):
        c=ws.cell(row=r,column=j,value=v)
        c.font=Font(name="Calibri",size=11); c.fill=white; c.border=bd
        c.alignment=Alignment(vertical="top",wrap_text=True)
WIDTHS=[14,13,13,7,8,10,9,11,24,9]   # 118 Excel units: single column band on Letter landscape at 100%
for col,wd in zip("ABCDEFGHIJ",WIDTHS):
    ws.column_dimensions[col].width=wd
import math
def wrapped_lines(val,wd):
    if not val: return 1
    n=0
    for seg in str(val).split("\n"):
        n+=max(1,math.ceil(len(seg)/max(wd-1,3)))   # sum segments, do not take max
    return n
for r in range(10,10+len(ROWS)):
    lines=max(wrapped_lines(ws.cell(row=r,column=j).value,wd)
              for j,wd in enumerate(WIDTHS,start=1))
    ws.row_dimensions[r].height=lines*13.2+3        # no cap: body type is never shrunk to fit
ws.row_dimensions[9].height=32
ws.freeze_panes="A10"; ws.print_title_rows="9:9"
ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_LETTER
ws.sheet_properties.pageSetUpPr.fitToPage=False
ws.page_setup.scale=100
ws.page_margins.left=ws.page_margins.right=0.7
ws.page_margins.top=ws.page_margins.bottom=0.75
xlsx=f"{OUT}/S19_Federal_1_PROPOSAL_Claude_20261019.xlsx"; wb.save(xlsx)
print("rows:",len(ROWS))

for fmt in ("ods","pdf"):
    subprocess.run(["soffice","--headless","--convert-to",fmt,"--outdir",OUT,xlsx],
                   check=True,capture_output=True,timeout=300)
