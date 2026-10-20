"""ODS / XLSX / DOCX table extraction and honest data-row counting."""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Any

from .util import norm_compare, norm_text

NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "meta": "urn:oasis:names:tc:opendocument:xmlns:meta:1.0",
    "dc": "http://purl.org/dc/elements/1.1/",
}
MAX_REPEAT = 2000  # repeated empty rows/cols in ODS can claim 1M; cap non-empty repeats only


def _q(prefix: str, tag: str) -> str:
    return f"{{{NS[prefix]}}}{tag}"


def _safe_xml(data: bytes) -> ET.Element:
    if b"<!ENTITY" in data or b"<!DOCTYPE" in data[:2048] and b"[" in data[:4096]:
        raise ValueError("XML with DTD/entity declarations refused")
    return ET.fromstring(data)


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return norm_text(v)


def read_ods(data: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    meta: dict[str, Any] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = _safe_xml(zf.read("content.xml"))
        if "meta.xml" in zf.namelist():
            try:
                m = _safe_xml(zf.read("meta.xml"))
                for tag, key in (("dc:title", "title"), ("dc:creator", "creator"), ("meta:generator", "generator")):
                    p, t = tag.split(":")
                    el = m.find(f".//{_q(p, t)}")
                    if el is not None and el.text:
                        meta[key] = el.text
            except Exception:
                pass
    for tbl in root.iter(_q("table", "table")):
        rows: list[list[str]] = []
        for tr in tbl.iter(_q("table", "table-row")):
            rep = int(tr.get(_q("table", "number-rows-repeated"), "1"))
            cells: list[str] = []
            for tc in tr:
                if tc.tag not in (_q("table", "table-cell"), _q("table", "covered-table-cell")):
                    continue
                crep = int(tc.get(_q("table", "number-columns-repeated"), "1"))
                vt = tc.get(_q("office", "value-type"))
                if vt == "date":
                    txt = (tc.get(_q("office", "date-value")) or "")[:10]
                elif vt in ("float", "percentage", "currency"):
                    txt = _cell_text(float(tc.get(_q("office", "value"), "0")))
                else:
                    txt = "\n".join("".join(p.itertext()) for p in tc.iter(_q("text", "p")))
                txt = norm_text(txt)
                cells.extend([txt] * (crep if (txt or crep <= 64) else 1))
            while cells and not cells[-1]:
                cells.pop()
            if not cells and rep > 1:
                rows.append([])  # collapse huge empty repeats into a single blank row
                continue
            for _ in range(min(rep, MAX_REPEAT)):
                rows.append(list(cells))
        while rows and not rows[-1]:
            rows.pop()
        tables.append({"name": tbl.get(_q("table", "name"), ""), "rows": rows})
    return tables, meta


def read_xlsx(data: bytes, load_format: bool = False) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=not load_format, data_only=True)
    tables: list[dict[str, Any]] = []
    fmt: dict[str, Any] = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()
            rows.append(cells)
        while rows and not rows[-1]:
            rows.pop()
        tables.append({"name": ws.title, "rows": rows})
        if load_format:
            hdr = next((c for c in ws[1]), None) if ws.max_row else None
            fmt[ws.title] = {
                "freeze": ws.freeze_panes,
                "col_widths": [round(ws.column_dimensions[k].width or 0, 1) for k in sorted(ws.column_dimensions)[:4]],
                "header_bold": bool(hdr is not None and hdr.font and hdr.font.b),
                "header_fill": getattr(getattr(hdr, "fill", None), "fgColor", None) and str(hdr.fill.fgColor.rgb) if hdr is not None else None,
            }
    p = wb.properties
    meta = {"creator": p.creator, "title": p.title, "created": str(p.created) if p.created else None,
            "modified": str(p.modified) if p.modified else None, "sheets": [ws.title for ws in wb.worksheets]}
    if fmt:
        meta["format"] = fmt
    wb.close()
    return tables, meta


def read_docx(data: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import docx

    d = docx.Document(io.BytesIO(data))
    tables = []
    for i, t in enumerate(d.tables):
        rows = []
        for r in t.rows:
            cells, prev = [], None
            for c in r.cells:
                if c._tc is prev:  # merged cell repeats; keep one copy
                    continue
                prev = c._tc
                cells.append(norm_text(c.text))
            while cells and not cells[-1]:
                cells.pop()
            rows.append(cells)
        tables.append({"name": f"table{i + 1}", "rows": rows})
    cp = d.core_properties
    sec = d.sections[0] if d.sections else None
    meta = {
        "author": cp.author, "title": cp.title, "created": str(cp.created) if cp.created else None,
        "paragraph_styles": sorted({p.style.name for p in d.paragraphs if p.style is not None}),
        "table_styles": sorted({t.style.name for t in d.tables if t.style is not None}),
        "orientation": str(sec.orientation) if sec else None,
        "margins_in": [round(x.inches, 2) for x in (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)] if sec else None,
        "table_count": len(d.tables),
    }
    return tables, meta


def read_tables(ext: str, data: bytes, load_format: bool = False) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if ext == ".ods":
        return read_ods(data)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(data, load_format)
    if ext == ".docx":
        return read_docx(data)
    raise ValueError(f"unsupported table extension {ext}")


# ---------------------------------------------------------------- row counting

def _schema_norms(cfg: dict[str, Any]) -> dict[str, list[str]]:
    return {"county": [norm_compare(h) for h in cfg["county_schema"]],
            "federal": [norm_compare(h) for h in cfg["federal_schema"]]}


def find_header(rows: list[list[str]], cfg: dict[str, Any]) -> tuple[int | None, str | None, float]:
    norms = _schema_norms(cfg)
    best: tuple[float, int, str] = (0.0, -1, "")
    for i, row in enumerate(rows[:15]):
        cells = {norm_compare(c) for c in row if c}
        if not cells:
            continue
        for kind, hs in norms.items():
            score = len(cells & set(hs)) / len(hs)
            if score > best[0]:
                best = (score, i, kind)
    if best[0] >= 0.6:
        return best[1], best[2], best[0]
    return None, None, best[0]


def kind_from_name(name: str, cfg: dict[str, Any]) -> str | None:
    for rule in cfg["table_kind_rules"]:
        if re.search(rule["regex"], name):
            return rule["kind"]
    return None


def count_rows(rows: list[list[str]], cfg: dict[str, Any], path: str = "") -> dict[str, Any]:
    """Classify every row. Only rows classed ``data`` count as instrument rows."""
    hidx, kind, score = find_header(rows, cfg)
    header_recognised = hidx is not None
    if hidx is None:
        hidx = next((i for i, r in enumerate(rows) if any(r)), None)
    header = rows[hidx] if hidx is not None else []
    hnorm = [norm_compare(c) for c in header]
    narrative = [re.compile(p) for p in cfg["narrative_patterns"]]
    excluded = {"empty": 0, "repeated_header": 0, "narrative_placeholder": 0, "total_row": 0, "pre_header": 0}
    data_rows: list[list[str]] = []
    data_rownos: list[int] = []
    for i, r in enumerate(rows):
        if hidx is None or i <= hidx:
            if i < (hidx if hidx is not None else 0) and any(r):
                excluded["pre_header"] += 1
            elif not any(r) and i != hidx:
                excluded["empty"] += 1
            continue
        ne = [c for c in r if c]
        if not ne:
            excluded["empty"] += 1
            continue
        if [norm_compare(c) for c in r][: len(hnorm)] == hnorm[: len(r)] and len(ne) >= 2:
            excluded["repeated_header"] += 1
            continue
        joined = " ".join(ne)
        if len(ne) <= 2 and (any(p.search(joined) for p in narrative) or (len(ne) == 1 and len(joined) > 120)):
            excluded["narrative_placeholder"] += 1
            continue
        if len(ne) <= 2 and re.match(r"(?i)^(sub)?totals?\b", ne[0]):
            excluded["total_row"] += 1
            continue
        data_rows.append(r)
        data_rownos.append(i + 1)
    return {
        "header_row": (hidx + 1) if hidx is not None else None,
        "header": header, "kind_by_header": kind, "header_score": round(score, 3),
        "header_recognised": header_recognised, "raw_rows": len(rows),
        "data_rows": len(data_rows), "excluded": excluded,
        "rows": data_rows, "row_numbers": data_rownos,
    }
