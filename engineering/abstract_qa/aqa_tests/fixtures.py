"""Synthetic fixture builders. Fictional identifiers only - never real client data."""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from aqa.config import COUNTY_SCHEMA_DEFAULT as COUNTY, FEDERAL_SCHEMA_DEFAULT as FED  # noqa: E402


def make_pdf(path: Path | None, pages: int) -> bytes:
    """Minimal valid PDF built by hand so tests do not depend on a PDF writer library."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{3 + i} 0 R' for i in range(pages))}] /Count {pages} >>"]
    objs += ["<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>"] * pages
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    data = bytes(out)
    if path is not None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    return data


def make_png(path: Path) -> None:
    import struct
    import zlib

    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = zlib.compress(b"\x00\x00\x00\x00")
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0))
                     + chunk(b"IDAT", raw) + chunk(b"IEND", b""))


def make_xlsx(path: Path, rows: list[list], sheet: str = "Index", bold_header: bool = True) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    if bold_header:
        for c in ws[1]:
            c.font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_docx(path: Path, rows: list[list]) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Narrative paragraph that must never be counted as a row.")
    width = max(len(r) for r in rows)
    t = d.add_table(rows=0, cols=width)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def make_ods(path: Path, rows: list[list], name: str = "Index") -> None:
    def cell(v):
        if v == "" or v is None:
            return "<table:table-cell/>"
        return f'<table:table-cell office:value-type="string"><text:p>{escape(str(v))}</text:p></table:table-cell>'

    body = "".join("<table:table-row>" + "".join(cell(v) for v in r) + "</table:table-row>" for r in rows)
    # trailing huge empty repeat, as real ODS exports contain
    body += '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>'
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet>'
        f'<table:table table:name="{name}">{body}</table:table></office:spreadsheet></office:body></office:document-content>'
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", content)


def county_row(n, tract, itype, grantor, grantee, date, rec, pdf, page):
    return [n, tract, itype, grantor, grantee, date, rec, pdf, page]


def fed_row(serial, date, etype, assignor, assignee, tract, acres, status, pdf, page):
    return [serial, date, etype, assignor, assignee, tract, acres, status, pdf, page]


def base_config(tmp: Path, **kw) -> dict:
    cfg = {
        "source_roots": [str(tmp / "src")],
        "output_root": str(tmp / "eng_out"),
        "branch_groups": {"A": ["R41", "R43"], "B": ["R10", "R11"]},
        "donor_section": "R41",
        "county_default": "Example",
        "claimed_placements": None,
    }
    cfg.update(kw)
    return cfg
