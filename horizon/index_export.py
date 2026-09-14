"""Export Penterra-shaped workbook rows as index-reconciliation faces.

The exporter does not invent values, guess neighbouring document numbers, or
write back to the workbook. Datetime cells are emitted as M/D/YYYY; other
cells are stripped source text.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .index_reconciliation import (
    COUNT_KEYS,
    PACKET_SCHEMA_ID,
    PACKET_SCHEMA_VERSION,
    REQUIRED_FIELDS,
    SOURCE_NAMES,
)
from .isolated_delta import (
    PACKET_SCHEMA_ID as DELTA_SCHEMA_ID,
    PACKET_SCHEMA_VERSION as DELTA_SCHEMA_VERSION,
    REQUIRED_DELTA_KEYS,
    sha256_file,
)
from .project_manifest import ControlFileError
from .stable_key import StableKeyError, stable_key
from .workbook_qa import load_workbook_profile

DEFAULT_PROFILE = (
    Path(__file__).resolve().parent / "profiles" / "penterra_index_v1.json"
)


class IndexExportError(ValueError):
    """Raised when a workbook cannot be exported as source faces."""


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value.year}"
    if isinstance(value, date):
        return f"{value.month}/{value.day}/{value.year}"
    return str(value).strip()


def export_faces(
    workbook_path: Path,
    *,
    profile_path: Optional[Path] = None,
    source_name: str = "candidate_rows",
) -> List[Dict[str, object]]:
    workbook_path = workbook_path.expanduser().resolve()
    profile = load_workbook_profile(profile_path or DEFAULT_PROFILE)
    tables = profile.get("abstract_tables")
    if not isinstance(tables, list) or not tables:
        raise IndexExportError("Workbook profile has no abstract_tables")
    table = tables[0]
    sheet_name = str(table.get("sheet", ""))
    start_row = table.get("start_row")
    columns = table.get("columns")
    if (
        not sheet_name
        or type(start_row) is not int
        or not isinstance(columns, dict)
        or "instrument_number" not in columns
    ):
        raise IndexExportError("Workbook profile table is incomplete")
    digest = sha256_file(workbook_path)
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise IndexExportError(f"Workbook has no sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        faces: List[Dict[str, object]] = []
        seen: set[str] = set()
        for row_number in range(start_row, (sheet.max_row or start_row) + 1):
            document = cell_text(sheet[f"{columns['instrument_number']}{row_number}"].value)
            locator = cell_text(
                sheet[f"{columns.get('book_page', 'E')}{row_number}"].value
                if columns.get("book_page")
                else None
            )
            fields = {
                name: cell_text(sheet[f"{columns[name]}{row_number}"].value)
                if name in columns
                else ""
                for name in REQUIRED_FIELDS
            }
            if not document and not locator and not any(fields.values()):
                continue
            try:
                key = stable_key(document or None, locator or None)
            except StableKeyError as exc:
                raise IndexExportError(
                    f"{sheet_name}!{row_number} has no house stable key"
                ) from exc
            if key in seen:
                raise IndexExportError(
                    f"{sheet_name} has duplicate stable_key {key}"
                )
            seen.add(key)
            faces.append(
                {
                    "stable_key": key,
                    "source_sha256": digest,
                    "page": row_number,
                    "crop_id": f"{sheet_name}!{row_number}",
                    "fields": fields,
                }
            )
    finally:
        workbook.close()
    return faces


def build_index_packet(
    packet_id: str,
    *,
    master: Sequence[Dict[str, object]],
    pdf_index: Sequence[Dict[str, object]],
    handwritten_index: Sequence[Dict[str, object]],
    candidate_rows: Sequence[Dict[str, object]],
    orphan_allowlist: Optional[Sequence[Dict[str, object]]] = None,
) -> Dict[str, object]:
    if not packet_id.strip():
        raise IndexExportError("packet_id must be a non-empty string")
    payload = {
        "schema_id": PACKET_SCHEMA_ID,
        "schema_version": PACKET_SCHEMA_VERSION,
        "packet_id": packet_id.strip(),
        "master": list(master),
        "pdf_index": list(pdf_index),
        "handwritten_index": list(handwritten_index),
        "candidate_rows": list(candidate_rows),
        "expected_counts": {
            "master": len(master),
            "pdf_index": len(pdf_index),
            "handwritten_index": len(handwritten_index),
            "candidate_rows": len(candidate_rows),
        },
        "orphan_allowlist": list(orphan_allowlist or ()),
    }
    if set(payload["expected_counts"]) != set(COUNT_KEYS):
        raise IndexExportError("expected_counts keys drifted")
    if not any(payload[name] for name in SOURCE_NAMES):
        raise IndexExportError("At least one source index must have rows")
    return payload


def bind_delta_packet(
    proposals: Sequence[Dict[str, object]],
    *,
    source_workbook_sha256: str,
    packet_id: str,
) -> Dict[str, object]:
    digest = source_workbook_sha256.casefold()
    if len(digest) != 64 or any(character not in "0123456789abcdef" for character in digest):
        raise IndexExportError("source_workbook_sha256 is invalid")
    deltas: List[Dict[str, object]] = []
    for index, raw in enumerate(proposals):
        if not isinstance(raw, dict):
            raise IndexExportError(f"proposals[{index}] must be an object")
        confidence = raw.get("confidence")
        if confidence not in {"high", "medium"}:
            raise IndexExportError(
                f"proposals[{index}] confidence must be high or medium"
            )
        delta = {}
        for key in REQUIRED_DELTA_KEYS:
            if key not in raw:
                raise IndexExportError(f"proposals[{index}] missing {key}")
            delta[key] = raw[key]
        if delta["replace"] is not False:
            raise IndexExportError(
                f"proposals[{index}] replace must be false for loop-built deltas"
            )
        deltas.append(delta)
    if not deltas:
        raise IndexExportError("No medium/high proposals to bind")
    return {
        "schema_id": DELTA_SCHEMA_ID,
        "schema_version": DELTA_SCHEMA_VERSION,
        "packet_id": packet_id.strip(),
        "source_workbook_sha256": digest,
        "deltas": deltas,
    }


def refresh_candidate(
    packet: Dict[str, object],
    candidate_rows: Sequence[Dict[str, object]],
) -> Dict[str, object]:
    refreshed = dict(packet)
    refreshed["candidate_rows"] = list(candidate_rows)
    counts = dict(packet.get("expected_counts") or {})
    counts["candidate_rows"] = len(candidate_rows)
    refreshed["expected_counts"] = counts
    return refreshed


def _faces_or_empty(
    workbook: Optional[Path],
    *,
    profile_path: Optional[Path],
) -> List[Dict[str, object]]:
    if workbook is None:
        return []
    return export_faces(workbook, profile_path=profile_path)


def export_index_packet(
    packet_id: str,
    *,
    master: Optional[Path] = None,
    pdf_index: Optional[Path] = None,
    handwritten: Optional[Path] = None,
    candidate: Optional[Path] = None,
    profile_path: Optional[Path] = None,
    orphan_allowlist: Optional[Sequence[Dict[str, object]]] = None,
) -> Dict[str, object]:
    if not any((master, pdf_index, handwritten)):
        raise IndexExportError(
            "Index packet needs --master, --pdf-index, and/or --handwritten"
        )
    return build_index_packet(
        packet_id,
        master=_faces_or_empty(master, profile_path=profile_path),
        pdf_index=_faces_or_empty(pdf_index, profile_path=profile_path),
        handwritten_index=_faces_or_empty(handwritten, profile_path=profile_path),
        candidate_rows=_faces_or_empty(candidate, profile_path=profile_path),
        orphan_allowlist=orphan_allowlist,
    )


def load_orphan_allowlist(path: Path) -> List[Dict[str, object]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IndexExportError(f"Cannot read {path}: {exc}") from exc
    if isinstance(payload, dict):
        if "orphan_allowlist" not in payload:
            raise IndexExportError(
                f"{path} must be a JSON list or an object with orphan_allowlist"
            )
        payload = payload["orphan_allowlist"]
    if not isinstance(payload, list):
        raise IndexExportError("orphan allowlist must be a JSON list")
    allowlist: List[Dict[str, object]] = []
    for index, item in enumerate(payload):
        if not isinstance(item, dict):
            raise IndexExportError(f"orphan_allowlist[{index}] must be an object")
        allowlist.append(item)
    return allowlist


def _load_json(path: Path) -> Dict[str, object]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IndexExportError(f"Cannot read {path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise IndexExportError(f"{path} must contain a JSON object")
    return payload


def _packet_mode(args: argparse.Namespace) -> bool:
    return any(
        (
            args.master,
            args.pdf_index,
            args.handwritten,
            args.candidate,
            args.orphan_allowlist,
        )
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Export a Penterra index workbook as faces, or build a "
            "master/PDF/handwritten reconciliation packet."
        )
    )
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--packet-id", required=True)
    parser.add_argument("--profile", type=Path)
    parser.add_argument("--as", dest="source_name", default="candidate_rows")
    parser.add_argument("--master", type=Path)
    parser.add_argument("--pdf-index", dest="pdf_index", type=Path)
    parser.add_argument("--handwritten", type=Path)
    parser.add_argument("--candidate", type=Path)
    parser.add_argument("--orphan-allowlist", type=Path)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        args = build_parser().parse_args(argv)
        if _packet_mode(args):
            candidate = args.candidate or args.workbook
            packet = export_index_packet(
                args.packet_id,
                master=args.master,
                pdf_index=args.pdf_index,
                handwritten=args.handwritten,
                candidate=candidate,
                profile_path=args.profile,
                orphan_allowlist=(
                    load_orphan_allowlist(args.orphan_allowlist)
                    if args.orphan_allowlist
                    else None
                ),
            )
            args.output.write_text(
                json.dumps(packet, indent=2, sort_keys=True),
                encoding="utf-8",
            )
            print(
                json.dumps(
                    {
                        "output": str(args.output),
                        "packet_id": packet["packet_id"],
                        "expected_counts": packet["expected_counts"],
                    },
                    indent=2,
                )
            )
            return 0
        if args.workbook is None:
            raise IndexExportError(
                "--workbook is required unless building a three-index packet"
            )
        faces = export_faces(
            args.workbook,
            profile_path=args.profile,
            source_name=args.source_name,
        )
        args.output.write_text(
            json.dumps(
                {
                    "source_name": args.source_name,
                    "packet_id": args.packet_id,
                    "workbook": str(args.workbook.resolve()),
                    "faces": faces,
                },
                indent=2,
                sort_keys=True,
            ),
            encoding="utf-8",
        )
    except (OSError, ControlFileError, IndexExportError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(json.dumps({"output": str(args.output), "face_count": len(faces)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
