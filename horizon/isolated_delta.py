"""Apply source-proved field deltas to an isolated workbook copy.

The input workbook is never overwritten. New rows are not invented. A delta
is applied only when the workbook hash matches, the row key exists, the
field is allowed, and the value is non-empty source text. This is not
package release.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .examiner import require_named_examiner, write_new_json
from .project_manifest import ControlFileError
from .stable_key import StableKeyError, stable_key
from .workbook_qa import load_workbook_profile

PACKET_SCHEMA_ID = "dbx.source_proved_delta_packet"
PACKET_SCHEMA_VERSION = "1.0"
DRAFT_SCHEMA_ID = "dbx.source_proved_delta_draft"
DRAFT_SCHEMA_VERSION = "1.0"
DRAFT_STATUS = "UNAPPROVED_DRAFT"
TEMPLATE_SCHEMA_ID = "dbx.source_proved_delta_template"
TEMPLATE_SCHEMA_VERSION = "1.0"
RECEIPT_SCHEMA_ID = "dbx.isolated_delta_receipt"
RECEIPT_SCHEMA_VERSION = "1.0"
REQUIRED_PACKET_KEYS = {
    "schema_id",
    "schema_version",
    "packet_id",
    "source_workbook_sha256",
    "deltas",
}
REQUIRED_DELTA_KEYS = {
    "row_key",
    "field",
    "value",
    "source_sha256",
    "page",
    "crop_id",
    "replace",
}
ALLOWED_FIELDS = {
    "document_type",
    "grantor",
    "grantee",
    "recorded_date",
    "legal_description",
    "book_page",
    "document_date",
    "comments",
}
DEFAULT_PROFILE = (
    Path(__file__).resolve().parent / "profiles" / "penterra_index_v1.json"
)


class IsolatedDeltaError(ValueError):
    """Raised when a delta packet or target copy is unsafe."""


@dataclass(frozen=True)
class Delta:
    row_key: str
    lookup_key: str
    field: str
    value: str
    source_sha256: str
    page: int
    crop_id: str
    replace: bool


@dataclass
class DeltaOutcome:
    row_key: str
    field: str
    status: str
    message: str
    cell: str = ""


@dataclass
class IsolatedDeltaReceipt:
    generated_utc: str
    packet_id: str
    source_workbook: str
    source_workbook_sha256: str
    isolated_workbook: str
    applied: int
    rejected: int
    outcomes: List[DeltaOutcome]
    technical_pass: bool
    schema_id: str = RECEIPT_SCHEMA_ID
    schema_version: str = RECEIPT_SCHEMA_VERSION
    notes: List[str] = field(default_factory=lambda: [
        "The source workbook was not modified",
        "technical_pass is not package release",
    ])

    def to_dict(self) -> Dict[str, object]:
        return asdict(self)


def _is_sha256(value: object) -> bool:
    return (
        isinstance(value, str)
        and len(value) == 64
        and all(character in "0123456789abcdef" for character in value)
    )


def _require_text(
    value: object,
    label: str,
    *,
    allow_newline: bool = False,
) -> str:
    if not isinstance(value, str):
        raise IsolatedDeltaError(f"{label} must be a string")
    if (not allow_newline) and ("\n" in value or "\r" in value):
        raise IsolatedDeltaError(f"{label} must be a single-line string")
    text = value.strip()
    if not text:
        raise IsolatedDeltaError(f"{label} must be a non-empty string")
    return text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_row_key(raw: object) -> tuple[str, str]:
    """Return ``(packet_row_key, house_stable_key)``.

    Packet keys are ``<docno>|<book>-<page>``. A bare document number is
    treated as ``<docno>|`` and matches only a row whose book/page is empty.
    """
    text = _require_text(raw, "row_key")
    try:
        if "|" in text:
            document, locator = text.split("|", 1)
            lookup = stable_key(document or None, locator or None)
        else:
            lookup = stable_key(text, None)
    except StableKeyError as exc:
        raise IsolatedDeltaError(
            f"row_key is not a house stable key: {text}"
        ) from exc
    return text, lookup


def _value_ok(value: str) -> bool:
    if ",," in value:
        return False
    if value != value.strip():
        return False
    return True


def parse_delta_packet(payload: Dict[str, object]) -> tuple[str, str, List[Delta]]:
    if set(payload) != REQUIRED_PACKET_KEYS:
        raise IsolatedDeltaError("Delta packet has invalid top-level fields")
    if (
        payload["schema_id"] != PACKET_SCHEMA_ID
        or payload["schema_version"] != PACKET_SCHEMA_VERSION
    ):
        raise IsolatedDeltaError("Delta packet schema is invalid")
    packet_id = _require_text(payload["packet_id"], "packet_id")
    digest = payload["source_workbook_sha256"]
    if isinstance(digest, str):
        digest = digest.casefold()
    if not _is_sha256(digest):
        raise IsolatedDeltaError("source_workbook_sha256 is invalid")
    raw_deltas = payload["deltas"]
    if not isinstance(raw_deltas, list) or not raw_deltas:
        raise IsolatedDeltaError("deltas must be a non-empty list")
    deltas: List[Delta] = []
    for index, raw in enumerate(raw_deltas):
        if not isinstance(raw, dict) or set(raw) != REQUIRED_DELTA_KEYS:
            raise IsolatedDeltaError(f"deltas[{index}] has invalid fields")
        field_name = _require_text(raw["field"], "field")
        if field_name not in ALLOWED_FIELDS:
            raise IsolatedDeltaError(f"deltas[{index}] field is not allowed")
        allow_newline = field_name in {
            "grantor",
            "grantee",
            "legal_description",
        }
        value = _require_text(
            raw["value"],
            f"deltas[{index}] value",
            allow_newline=allow_newline,
        )
        if not _value_ok(value):
            raise IsolatedDeltaError(f"deltas[{index}] value is unsafe")
        source = raw["source_sha256"]
        if isinstance(source, str):
            source = source.casefold()
        if not _is_sha256(source):
            raise IsolatedDeltaError(f"deltas[{index}] source_sha256 is invalid")
        page = raw["page"]
        if type(page) is not int or page < 1:
            raise IsolatedDeltaError(f"deltas[{index}] page must be an integer >= 1")
        replace = raw["replace"]
        if type(replace) is not bool:
            raise IsolatedDeltaError(f"deltas[{index}] replace must be a boolean")
        try:
            row_key, lookup_key = parse_row_key(raw["row_key"])
        except IsolatedDeltaError as exc:
            raise IsolatedDeltaError(f"deltas[{index}] {exc}") from exc
        deltas.append(
            Delta(
                row_key=row_key,
                lookup_key=lookup_key,
                field=field_name,
                value=value.strip(),
                source_sha256=source,
                page=page,
                crop_id=_require_text(raw["crop_id"], "crop_id"),
                replace=replace,
            )
        )
    return packet_id, digest, deltas


def write_delta_packet(
    *,
    workbook: Path,
    deltas: Sequence[Dict[str, object]],
    output: Path,
    packet_id: str,
) -> Dict[str, object]:
    """Wrap examiner-held deltas around the current workbook hash.

    This does not invent legal, party, date, or document-type values.
    """
    resolved = workbook.expanduser()
    if not resolved.is_file():
        raise IsolatedDeltaError(f"workbook does not exist: {resolved}")
    token = packet_id.strip()
    if not token or "\n" in token:
        raise IsolatedDeltaError("packet_id must be a single-line string")
    if not isinstance(deltas, list) or not deltas:
        raise IsolatedDeltaError("deltas must be a non-empty list")
    packet = {
        "schema_id": PACKET_SCHEMA_ID,
        "schema_version": PACKET_SCHEMA_VERSION,
        "packet_id": token,
        "source_workbook_sha256": sha256_file(resolved),
        "deltas": list(deltas),
    }
    parse_delta_packet(packet)
    try:
        write_new_json(packet, output, "source-proved delta packet")
    except ValueError as exc:
        raise IsolatedDeltaError(str(exc)) from exc
    return packet


def write_delta_draft(
    *,
    workbook: Path,
    deltas: Sequence[Dict[str, object]],
    output: Path,
    packet_id: str,
) -> Dict[str, object]:
    """Hash-bind medium/high proposals. Does not apply or invent values."""
    resolved = workbook.expanduser().resolve()
    if not resolved.is_file():
        raise IsolatedDeltaError(f"workbook does not exist: {resolved}")
    token = packet_id.strip()
    if not token or "\n" in token:
        raise IsolatedDeltaError("packet_id must be a single-line string")
    if not isinstance(deltas, list) or not deltas:
        raise IsolatedDeltaError("delta draft requires medium/high proposals")
    parse_delta_packet(
        {
            "schema_id": PACKET_SCHEMA_ID,
            "schema_version": PACKET_SCHEMA_VERSION,
            "packet_id": token,
            "source_workbook_sha256": sha256_file(resolved),
            "deltas": list(deltas),
        }
    )
    draft = {
        "schema_id": DRAFT_SCHEMA_ID,
        "schema_version": DRAFT_SCHEMA_VERSION,
        "status": DRAFT_STATUS,
        "packet_id": token,
        "source_workbook_sha256": sha256_file(resolved),
        "operator": "",
        "deltas": list(deltas),
        "notes": [
            "UNAPPROVED_DRAFT cannot bind isolated_delta",
            "One-source blanks and conflicts are not in this draft",
            "Attest with a named examiner; Horizon does not invent values",
        ],
    }
    output = output.expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(draft, indent=2, sort_keys=True), encoding="utf-8")
    return draft


def attest_delta_draft(
    draft: Dict[str, object],
    *,
    workbook: Path,
    output: Path,
    operator: str,
) -> Dict[str, object]:
    """Promote a 2+ source draft. Horizon does not invent field values."""
    try:
        require_named_examiner(operator)
    except ValueError as exc:
        raise IsolatedDeltaError(str(exc)) from exc
    if (
        draft.get("schema_id") != DRAFT_SCHEMA_ID
        or draft.get("schema_version") != DRAFT_SCHEMA_VERSION
        or draft.get("status") != DRAFT_STATUS
    ):
        raise IsolatedDeltaError("source-proved delta draft schema is invalid")
    resolved = workbook.expanduser().resolve()
    if not resolved.is_file():
        raise IsolatedDeltaError(f"workbook does not exist: {resolved}")
    actual = sha256_file(resolved)
    expected = str(draft.get("source_workbook_sha256") or "").casefold()
    if actual != expected:
        raise IsolatedDeltaError(
            "Workbook hash does not match the delta draft; rebuild the "
            "draft against the current isolated workbook"
        )
    raw_deltas = draft.get("deltas")
    if not isinstance(raw_deltas, list) or not raw_deltas:
        raise IsolatedDeltaError("delta draft has no medium/high proposals")
    return write_delta_packet(
        workbook=resolved,
        deltas=raw_deltas,
        output=output,
        packet_id=str(draft.get("packet_id") or ""),
    )


def write_onesource_template(
    *,
    workbook: Path,
    rows: Sequence[Dict[str, object]],
    output: Path,
    packet_id: str,
) -> Dict[str, object]:
    """Write empty-value one-source/conflict rows. Does not invent fills."""
    resolved = workbook.expanduser().resolve()
    if not resolved.is_file():
        raise IsolatedDeltaError(f"workbook does not exist: {resolved}")
    token = packet_id.strip()
    if not token or "\n" in token:
        raise IsolatedDeltaError("packet_id must be a single-line string")
    if not rows:
        raise IsolatedDeltaError("one-source template requires remaining blanks")
    rows = list(rows)
    for index, raw in enumerate(rows):
        if not isinstance(raw, dict):
            raise IsolatedDeltaError(f"rows[{index}] must be an object")
        if raw.get("value"):
            raise IsolatedDeltaError(
                f"rows[{index}] value must stay empty until a writer holds "
                "source-proved text"
            )
    template = {
        "schema_id": TEMPLATE_SCHEMA_ID,
        "schema_version": TEMPLATE_SCHEMA_VERSION,
        "status": DRAFT_STATUS,
        "packet_id": token,
        "source_workbook_sha256": sha256_file(resolved),
        "operator": "",
        "deltas": list(rows),
        "notes": [
            "UNAPPROVED_DRAFT cannot bind isolated_delta",
            "value must stay empty until a writer holds source-proved text",
            "Do not copy source_values into value without a face",
        ],
    }
    output = output.expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(template, indent=2, sort_keys=True), encoding="utf-8"
    )
    return template


def attest_onesource_template(
    template: Dict[str, object],
    *,
    workbook: Path,
    output: Path,
    operator: str,
) -> Dict[str, object]:
    """Promote filled one-source rows. Horizon does not invent values."""
    try:
        require_named_examiner(operator)
    except ValueError as exc:
        raise IsolatedDeltaError(str(exc)) from exc
    if (
        template.get("schema_id") != TEMPLATE_SCHEMA_ID
        or template.get("schema_version") != TEMPLATE_SCHEMA_VERSION
        or template.get("status") != DRAFT_STATUS
    ):
        raise IsolatedDeltaError("one-source template schema is invalid")
    resolved = workbook.expanduser().resolve()
    if not resolved.is_file():
        raise IsolatedDeltaError(f"workbook does not exist: {resolved}")
    actual = sha256_file(resolved)
    expected = str(template.get("source_workbook_sha256") or "").casefold()
    if actual != expected:
        raise IsolatedDeltaError(
            "Workbook hash does not match the one-source template; rebuild "
            "it against the current isolated workbook"
        )
    raw_rows = template.get("deltas")
    if not isinstance(raw_rows, list):
        raise IsolatedDeltaError("one-source template deltas must be a list")
    filled: List[Dict[str, object]] = []
    for index, raw in enumerate(raw_rows):
        if not isinstance(raw, dict):
            raise IsolatedDeltaError(f"deltas[{index}] must be an object")
        value = raw.get("value")
        if not isinstance(value, str) or not value.strip():
            continue
        filled.append({key: raw.get(key) for key in REQUIRED_DELTA_KEYS})
    if not filled:
        raise IsolatedDeltaError(
            "Fill at least one source-proved value; Horizon does not invent them"
        )
    return write_delta_packet(
        workbook=resolved,
        deltas=filled,
        output=output,
        packet_id=str(template.get("packet_id") or ""),
    )


def apply_deltas(
    source_workbook: Path,
    output_workbook: Path,
    packet: Dict[str, object],
    *,
    profile_path: Optional[Path] = None,
) -> IsolatedDeltaReceipt:
    source_workbook = source_workbook.expanduser().resolve()
    output_workbook = output_workbook.expanduser().resolve()
    if output_workbook.exists():
        raise IsolatedDeltaError("Isolated output workbook must not already exist")
    if output_workbook == source_workbook:
        raise IsolatedDeltaError("Refusing to write over the source workbook")
    packet_id, expected_sha, deltas = parse_delta_packet(packet)
    actual_sha = sha256_file(source_workbook)
    if actual_sha != expected_sha:
        raise IsolatedDeltaError(
            "Workbook hash does not match the delta packet source_workbook_sha256"
        )
    profile = load_workbook_profile(profile_path or DEFAULT_PROFILE)
    tables = profile.get("abstract_tables")
    if not isinstance(tables, list) or not tables:
        raise IsolatedDeltaError("Workbook profile has no abstract_tables")
    table = tables[0]
    sheet_name = str(table.get("sheet", ""))
    start_row = table.get("start_row")
    columns = table.get("columns")
    if (
        not sheet_name
        or type(start_row) is not int
        or not isinstance(columns, dict)
        or "instrument_number" not in columns
    ):
        raise IsolatedDeltaError("Workbook profile table is incomplete")
    copied = False
    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(source_workbook, output_workbook)
        copied = True
        from openpyxl import load_workbook

        workbook = load_workbook(output_workbook)
        try:
            if sheet_name not in workbook.sheetnames:
                raise IsolatedDeltaError(f"Workbook has no sheet {sheet_name!r}")
            sheet = workbook[sheet_name]
            doc_column = str(columns["instrument_number"]).strip().upper()
            book_column = str(columns.get("book_page") or "").strip().upper()
            rows_by_key: Dict[str, int] = {}
            duplicate_keys: set[str] = set()
            for row_number in range(start_row, (sheet.max_row or start_row) + 1):
                document = sheet[f"{doc_column}{row_number}"].value
                locator = (
                    sheet[f"{book_column}{row_number}"].value if book_column else None
                )
                try:
                    key = stable_key(document, locator)
                except StableKeyError:
                    continue
                if key in rows_by_key:
                    duplicate_keys.add(key)
                else:
                    rows_by_key[key] = row_number
            outcomes: List[DeltaOutcome] = []
            applied = 0
            for delta in deltas:
                if delta.lookup_key in duplicate_keys:
                    raise IsolatedDeltaError(
                        f"row_key {delta.row_key} matches more than one index row"
                    )
                row_number = rows_by_key.get(delta.lookup_key)
                if row_number is None:
                    raise IsolatedDeltaError(
                        f"unknown row_key {delta.row_key}; refusing to invent a row"
                    )
                if delta.field not in columns:
                    raise IsolatedDeltaError(
                        f"profile has no column for field {delta.field}"
                    )
                column = str(columns[delta.field]).strip().upper()
                cell = f"{column}{row_number}"
                current = sheet[cell].value
                populated = current is not None and bool(str(current).strip())
                if populated and str(current).strip() == delta.value:
                    outcomes.append(
                        DeltaOutcome(
                            row_key=delta.row_key,
                            field=delta.field,
                            status="already_present",
                            message="cell already matches the source-proved value",
                            cell=cell,
                        )
                    )
                    continue
                if populated and not delta.replace:
                    raise IsolatedDeltaError(
                        f"{cell} already has a different value; "
                        "set replace true only with source proof"
                    )
                sheet[cell] = delta.value
                applied += 1
                outcomes.append(
                    DeltaOutcome(
                        row_key=delta.row_key,
                        field=delta.field,
                        status="applied",
                        message="wrote source-proved value to isolated copy",
                        cell=cell,
                    )
                )
            workbook.save(output_workbook)
        finally:
            workbook.close()
    except Exception:
        if copied and output_workbook.exists():
            output_workbook.unlink()
        raise
    return IsolatedDeltaReceipt(
        generated_utc=datetime.now(timezone.utc).isoformat(),
        packet_id=packet_id,
        source_workbook=str(source_workbook),
        source_workbook_sha256=actual_sha,
        isolated_workbook=str(output_workbook),
        applied=applied,
        rejected=0,
        outcomes=outcomes,
        technical_pass=True,
    )


def _load_json(path: Path) -> Dict[str, object]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IsolatedDeltaError(f"Cannot read {path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise IsolatedDeltaError(f"{path} must contain a JSON object")
    return payload


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Apply or write source-proved deltas. Does not invent values."
    )
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--attest", action="store_true")
    parser.add_argument("--from-draft", type=Path)
    parser.add_argument("--from-template", type=Path)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--packet", type=Path)
    parser.add_argument("--deltas", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt", type=Path)
    parser.add_argument("--packet-id")
    parser.add_argument("--operator")
    parser.add_argument("--profile", type=Path)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        args = build_parser().parse_args(argv)
        if args.attest:
            if args.from_draft is not None and args.from_template is not None:
                raise IsolatedDeltaError(
                    "--attest accepts only one of --from-draft or --from-template"
                )
            if args.operator is None or (
                args.from_draft is None and args.from_template is None
            ):
                raise IsolatedDeltaError(
                    "--attest requires --operator and --from-draft or "
                    "--from-template; Horizon does not invent field values"
                )
            if args.from_template is not None:
                packet = attest_onesource_template(
                    _load_json(args.from_template),
                    workbook=args.workbook,
                    output=args.output,
                    operator=args.operator,
                )
            else:
                packet = attest_delta_draft(
                    _load_json(args.from_draft),
                    workbook=args.workbook,
                    output=args.output,
                    operator=args.operator,
                )
            print(
                json.dumps(
                    {
                        "output": str(args.output),
                        "source_workbook_sha256": packet["source_workbook_sha256"],
                        "delta_count": len(packet["deltas"]),
                        "packages_complete": False,
                    },
                    indent=2,
                )
            )
            return 0
        if args.write:
            if args.deltas is None or args.packet_id is None:
                raise IsolatedDeltaError(
                    "--write requires --deltas and --packet-id; Horizon "
                    "does not invent field values"
                )
            raw = _load_json(args.deltas)
            deltas = raw.get("deltas") if isinstance(raw, dict) and "deltas" in raw else raw
            if not isinstance(deltas, list):
                raise IsolatedDeltaError("deltas file must be a list or {\"deltas\": [...]}")
            packet = write_delta_packet(
                workbook=args.workbook,
                deltas=deltas,
                output=args.output,
                packet_id=args.packet_id,
            )
            print(
                json.dumps(
                    {
                        "output": str(args.output),
                        "source_workbook_sha256": packet["source_workbook_sha256"],
                        "delta_count": len(packet["deltas"]),
                        "packages_complete": False,
                    },
                    indent=2,
                )
            )
            return 0
        if args.packet is None or args.receipt is None:
            raise IsolatedDeltaError("Pass --packet and --receipt to apply, or --write")
        receipt = apply_deltas(
            args.workbook,
            args.output,
            _load_json(args.packet),
            profile_path=args.profile,
        )
        args.receipt.write_text(
            json.dumps(receipt.to_dict(), indent=2, sort_keys=True),
            encoding="utf-8",
        )
    except (OSError, ControlFileError, IsolatedDeltaError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(
        json.dumps(
            {
                "receipt": str(args.receipt),
                "isolated_workbook": receipt.isolated_workbook,
                "applied": receipt.applied,
                "rejected": receipt.rejected,
                "technical_pass": receipt.technical_pass,
            },
            indent=2,
        )
    )
    return 0 if receipt.technical_pass else 2


if __name__ == "__main__":
    raise SystemExit(main())
