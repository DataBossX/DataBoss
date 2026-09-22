"""Apply the Penterra print-layout contract to an isolated workbook copy.

The source workbook is never overwritten. This writes Letter / landscape /
print-title settings from the workbook profile. It is not native Excel Print
Preview and not package release.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .isolated_delta import sha256_file
from .project_manifest import ControlFileError
from .workbook_qa import load_workbook_profile

RECEIPT_SCHEMA_ID = "dbx.print_layout_repair_receipt"
RECEIPT_SCHEMA_VERSION = "1.0"
DEFAULT_PROFILE = (
    Path(__file__).resolve().parent / "profiles" / "penterra_index_v1.json"
)
PAPERSIZE_LETTER = 1


class PrintLayoutRepairError(ValueError):
    """Raised when isolated print-layout repair is unsafe."""


@dataclass
class PrintLayoutRepairReceipt:
    generated_utc: str
    source_workbook: str
    source_workbook_sha256: str
    isolated_workbook: str
    applied: List[str]
    technical_pass: bool
    schema_id: str = RECEIPT_SCHEMA_ID
    schema_version: str = RECEIPT_SCHEMA_VERSION
    notes: List[str] = field(
        default_factory=lambda: [
            "The source workbook was not modified",
            "This is not native Excel Print Preview",
            "technical_pass is not package release",
        ]
    )

    def to_dict(self) -> Dict[str, object]:
        return asdict(self)


def _title_rows(value: object) -> str:
    text = str(value or "").replace("$", "")
    if "!" in text:
        text = text.split("!", 1)[1]
    return text


def repair_print_layout(
    source_workbook: Path,
    output_workbook: Path,
    *,
    profile_path: Optional[Path] = None,
) -> PrintLayoutRepairReceipt:
    source_workbook = source_workbook.expanduser().resolve()
    output_workbook = output_workbook.expanduser().resolve()
    if output_workbook.exists():
        raise PrintLayoutRepairError("Isolated output workbook must not already exist")
    if output_workbook == source_workbook:
        raise PrintLayoutRepairError("Refusing to write over the source workbook")
    profile = load_workbook_profile(profile_path or DEFAULT_PROFILE)
    rules = profile.get("abstract_print_layout")
    if not isinstance(rules, list) or not rules:
        raise PrintLayoutRepairError("Workbook profile has no abstract_print_layout")
    digest = sha256_file(source_workbook)
    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_workbook, output_workbook)
    from openpyxl import load_workbook

    workbook = load_workbook(output_workbook)
    applied: List[str] = []
    try:
        for index, rule in enumerate(rules):
            if not isinstance(rule, dict) or "sheet" not in rule:
                raise PrintLayoutRepairError(
                    f"abstract_print_layout[{index}] is invalid"
                )
            sheet_name = str(rule["sheet"])
            if sheet_name not in workbook.sheetnames:
                raise PrintLayoutRepairError(f"Workbook has no sheet {sheet_name!r}")
            sheet = workbook[sheet_name]
            if rule.get("orientation"):
                sheet.page_setup.orientation = str(rule["orientation"])
                applied.append(f"{sheet_name}.orientation")
            if "paper_size" in rule:
                paper = rule["paper_size"]
                if paper != PAPERSIZE_LETTER:
                    raise PrintLayoutRepairError(
                        "Isolated repair only writes US Letter (paper_size 1)"
                    )
                sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
                applied.append(f"{sheet_name}.paper_size")
            if rule.get("print_title_rows"):
                sheet.print_title_rows = _title_rows(rule["print_title_rows"])
                applied.append(f"{sheet_name}.print_title_rows")
            if "fit_to_width" in rule:
                sheet.page_setup.fitToWidth = int(rule["fit_to_width"])
                applied.append(f"{sheet_name}.fit_to_width")
            if "fit_to_height" in rule:
                sheet.page_setup.fitToHeight = int(rule["fit_to_height"])
                applied.append(f"{sheet_name}.fit_to_height")
            if rule.get("fit_to_page") is True:
                sheet.sheet_properties.pageSetUpPr.fitToPage = True
                applied.append(f"{sheet_name}.fit_to_page")
        workbook.save(output_workbook)
    except Exception:
        workbook.close()
        if output_workbook.exists():
            output_workbook.unlink()
        raise
    workbook.close()
    return PrintLayoutRepairReceipt(
        generated_utc=datetime.now(timezone.utc).isoformat(),
        source_workbook=str(source_workbook),
        source_workbook_sha256=digest,
        isolated_workbook=str(output_workbook),
        applied=applied,
        technical_pass=bool(applied),
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Write Penterra Letter print layout onto an isolated copy."
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    parser.add_argument("--profile", type=Path)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        args = build_parser().parse_args(argv)
        receipt = repair_print_layout(
            args.workbook,
            args.output,
            profile_path=args.profile,
        )
        args.receipt.write_text(
            json.dumps(receipt.to_dict(), indent=2, sort_keys=True),
            encoding="utf-8",
        )
    except (OSError, ControlFileError, PrintLayoutRepairError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(
        json.dumps(
            {
                "receipt": str(args.receipt),
                "isolated_workbook": receipt.isolated_workbook,
                "applied": receipt.applied,
                "technical_pass": receipt.technical_pass,
            },
            indent=2,
        )
    )
    return 0 if receipt.technical_pass else 2


if __name__ == "__main__":
    raise SystemExit(main())
