"""Project a Penterra index workbook into tract-ledger and occurrence packets.

This does not invent cells, guess neighbouring document numbers, or mutate
the workbook. Blank book/page stays blank. Bare document numbers stay bare.
Excel row numbers are locators, not recorded-page claims.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .index_export import IndexExportError, cell_text
from .isolated_delta import sha256_file
from .occurrence_build import RISK_BARE, RISK_CHECKABLE
from .occurrence_ledger import (
    PACKET_SCHEMA_ID as OCCURRENCE_SCHEMA_ID,
    PACKET_SCHEMA_VERSION as OCCURRENCE_SCHEMA_VERSION,
    parse_occurrence_packet,
)
from .reextraction_gate import (
    EXPORT_SCHEMA_ID,
    EXPORT_SCHEMA_VERSION,
    parse_ledger_export,
)
from .stable_key import StableKeyError, stable_key
from .workbook_qa import load_workbook_profile

DEFAULT_PROFILE = (
    Path(__file__).resolve().parent / "profiles" / "penterra_index_v1.json"
)


class WorkbookLedgerError(ValueError):
    """Raised when a workbook cannot be projected into a ledger."""


def _load_table(workbook_path: Path, profile_path: Optional[Path]) -> tuple[object, Dict[str, object], str]:
    workbook_path = workbook_path.expanduser().resolve()
    if not workbook_path.is_file():
        raise WorkbookLedgerError(f"workbook does not exist: {workbook_path}")
    profile = load_workbook_profile(profile_path or DEFAULT_PROFILE)
    tables = profile.get("abstract_tables")
    if not isinstance(tables, list) or not tables:
        raise WorkbookLedgerError("Workbook profile has no abstract_tables")
    table = tables[0]
    sheet_name = str(table.get("sheet", ""))
    start_row = table.get("start_row")
    columns = table.get("columns")
    if (
        not sheet_name
        or type(start_row) is not int
        or not isinstance(columns, dict)
        or "instrument_number" not in columns
    ):
        raise WorkbookLedgerError("Workbook profile table is incomplete")
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise WorkbookLedgerError(f"Workbook has no sheet {sheet_name!r}")
    return workbook, table, sha256_file(workbook_path)


def iter_index_rows(
    workbook_path: Path,
    *,
    profile_path: Optional[Path] = None,
) -> tuple[str, List[Dict[str, object]]]:
    workbook, table, digest = _load_table(workbook_path, profile_path)
    sheet_name = str(table["sheet"])
    start_row = int(table["start_row"])
    columns = table["columns"]
    assert isinstance(columns, dict)
    rows: List[Dict[str, object]] = []
    seen: set[str] = set()
    try:
        sheet = workbook[sheet_name]
        for row_number in range(start_row, (sheet.max_row or start_row) + 1):
            docno = cell_text(sheet[f"{columns['instrument_number']}{row_number}"].value)
            bookpage = cell_text(
                sheet[f"{columns['book_page']}{row_number}"].value
                if columns.get("book_page")
                else None
            )
            grantor = cell_text(
                sheet[f"{columns['grantor']}{row_number}"].value
                if columns.get("grantor")
                else None
            )
            grantee = cell_text(
                sheet[f"{columns['grantee']}{row_number}"].value
                if columns.get("grantee")
                else None
            )
            rec_date = cell_text(
                sheet[f"{columns['recorded_date']}{row_number}"].value
                if columns.get("recorded_date")
                else None
            )
            doc_date = cell_text(
                sheet[f"{columns['document_date']}{row_number}"].value
                if columns.get("document_date")
                else None
            )
            legal = cell_text(
                sheet[f"{columns['legal_description']}{row_number}"].value
                if columns.get("legal_description")
                else None
            )
            doc_type = cell_text(
                sheet[f"{columns['document_type']}{row_number}"].value
                if columns.get("document_type")
                else None
            )
            if not any((docno, bookpage, grantor, grantee, rec_date, legal, doc_type)):
                continue
            try:
                key = stable_key(docno or None, bookpage or None)
            except StableKeyError as exc:
                raise WorkbookLedgerError(
                    f"{sheet_name}!{row_number} has no house stable key"
                ) from exc
            if key in seen:
                raise WorkbookLedgerError(
                    f"{sheet_name} has duplicate stable_key {key}"
                )
            seen.add(key)
            rows.append(
                {
                    "row_id": f"wb{row_number:04d}",
                    "docno": docno,
                    "bookpage": bookpage,
                    "rec_date": rec_date,
                    "doc_date": doc_date,
                    "grantor": grantor,
                    "grantee": grantee,
                    "page": row_number,
                    "crop_id": f"{sheet_name}!{row_number}",
                    "stable_key": key,
                    "source_sha256": digest,
                }
            )
    finally:
        workbook.close()
    if not rows:
        raise WorkbookLedgerError("Workbook has no populated index rows")
    return digest, rows


def workbook_to_tract_export(
    workbook_path: Path,
    *,
    packet_id: str,
    profile_path: Optional[Path] = None,
) -> Dict[str, object]:
    token = packet_id.strip()
    if not token or "\n" in token:
        raise WorkbookLedgerError("packet_id must be a single-line string")
    _digest, rows = iter_index_rows(workbook_path, profile_path=profile_path)
    export = {
        "schema_id": EXPORT_SCHEMA_ID,
        "schema_version": EXPORT_SCHEMA_VERSION,
        "packet_id": token,
        "rows": [
            {
                "row_id": item["row_id"],
                "docno": item["docno"],
                "bookpage": item["bookpage"],
                "rec_date": item["rec_date"],
                "doc_date": item["doc_date"],
                "grantor": item["grantor"],
                "grantee": item["grantee"],
                "page": item["page"],
                "crop_id": item["crop_id"],
            }
            for item in rows
        ],
    }
    parse_ledger_export(export)
    return export


def workbook_to_occurrence_packet(
    workbook_path: Path,
    *,
    packet_id: str,
    profile_path: Optional[Path] = None,
) -> Dict[str, object]:
    token = packet_id.strip()
    if not token or "\n" in token:
        raise WorkbookLedgerError("packet_id must be a single-line string")
    _digest, rows = iter_index_rows(workbook_path, profile_path=profile_path)
    occurrences: List[Dict[str, object]] = []
    ledger: List[Dict[str, str]] = []
    candidates: List[Dict[str, object]] = []
    seen: set[str] = set()
    for item in rows:
        key = str(item["stable_key"])
        digest = str(item["source_sha256"])
        rec_date = str(item["rec_date"])
        grantor = str(item["grantor"])
        grantee = str(item["grantee"])
        checkable = bool(item["bookpage"]) or bool(rec_date and (grantor or grantee))
        occurrence = {
            "occurrence_id": str(item["row_id"]),
            "stable_key": key,
            "source_sha256": digest,
            "page": item["page"],
            "crop_id": str(item["crop_id"]),
            "fields": {
                "grantor": grantor,
                "grantee": grantee,
                "recorded_date": rec_date,
                "party": grantor or grantee,
                "date_role": "recorded" if rec_date else "",
            },
            "risk": RISK_CHECKABLE if checkable else RISK_BARE,
        }
        occurrences.append(occurrence)
        if key not in seen:
            seen.add(key)
            ledger.append({"stable_key": key, "source_sha256": digest})
            candidates.append(
                {
                    "row_id": str(item["row_id"]),
                    "stable_key": key,
                    "source_sha256": digest,
                    "fields": occurrence["fields"],
                }
            )
    packet = {
        "schema_id": OCCURRENCE_SCHEMA_ID,
        "schema_version": OCCURRENCE_SCHEMA_VERSION,
        "packet_id": token,
        "occurrences": occurrences,
        "unique_key_ledger": ledger,
        "candidate_rows": candidates,
        "allowlist": [],
    }
    parse_occurrence_packet(packet)
    return packet


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Project a Penterra workbook into tract-ledger and occurrence "
            "packets. Does not invent cells."
        )
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--packet-id", required=True)
    parser.add_argument("--export", type=Path)
    parser.add_argument("--occurrence", type=Path)
    parser.add_argument("--profile", type=Path)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        args = build_parser().parse_args(argv)
        if args.export is None and args.occurrence is None:
            raise WorkbookLedgerError("Pass --export and/or --occurrence")
        if args.export is not None:
            export = workbook_to_tract_export(
                args.workbook,
                packet_id=args.packet_id,
                profile_path=args.profile,
            )
            args.export.write_text(
                json.dumps(export, indent=2, sort_keys=True),
                encoding="utf-8",
            )
        if args.occurrence is not None:
            packet = workbook_to_occurrence_packet(
                args.workbook,
                packet_id=args.packet_id,
                profile_path=args.profile,
            )
            args.occurrence.write_text(
                json.dumps(packet, indent=2, sort_keys=True),
                encoding="utf-8",
            )
    except (OSError, WorkbookLedgerError, IndexExportError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(
        json.dumps(
            {
                "workbook": str(args.workbook),
                "export": str(args.export) if args.export else None,
                "occurrence": str(args.occurrence) if args.occurrence else None,
                "packages_complete": False,
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
