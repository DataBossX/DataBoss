"""Report style gaps against the Section 11 grid. Findings omit cell values."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime

from openpyxl.worksheet.worksheet import Worksheet

from .apply import find_index, is_checklist
from .profile import (
    CENTER_NO_WRAP,
    CHECKLIST_SCALE,
    CHECKLIST_WIDTHS,
    DATE_FORMAT,
    DATE_HEADERS,
    FONT_NAME,
    FONT_SIZE,
    HEADER_ROW_HEIGHT,
    MARGIN_BOTTOM,
    MARGIN_LEFT,
    MARGIN_RIGHT,
    MARGIN_TOP,
    MAX_DATA_ROW_HEIGHT,
    META_ROW_HEIGHT,
    PAPER_LETTER,
)


@dataclass(frozen=True)
class Finding:
    sheet: str
    cell: str
    code: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


def _close(actual: float | None, expected: float, tolerance: float = 0.05) -> bool:
    if actual is None:
        return False
    return abs(actual - expected) <= tolerance


def _font_ok(cell) -> bool:
    font = cell.font
    name = (font.name or "").lower()
    size = font.size
    return name == FONT_NAME.lower() and size == FONT_SIZE


def _audit_index(ws: Worksheet) -> list[Finding]:
    found = find_index(ws)
    assert found is not None
    _kind, header_row, headers, widths = found
    findings: list[Finding] = []
    if ws.freeze_panes:
        findings.append(Finding(ws.title, "A1", "freeze_panes"))
    if ws.print_title_rows or ws.print_title_cols:
        findings.append(Finding(ws.title, "A1", "print_titles"))
    if (ws.page_setup.orientation or "").lower() != "landscape":
        findings.append(Finding(ws.title, "A1", "orientation"))
    if str(ws.page_setup.paperSize or "") not in {str(PAPER_LETTER), "1"}:
        findings.append(Finding(ws.title, "A1", "paper_size"))
    if ws.page_setup.fitToWidth != 1:
        findings.append(Finding(ws.title, "A1", "fit_width"))
    margins = ws.page_margins
    if not _close(margins.left, MARGIN_LEFT) or not _close(margins.right, MARGIN_RIGHT):
        findings.append(Finding(ws.title, "A1", "horizontal_margin"))
    if not _close(margins.top, MARGIN_TOP) or not _close(margins.bottom, MARGIN_BOTTOM):
        findings.append(Finding(ws.title, "A1", "vertical_margin"))
    for index, width in enumerate(widths, start=1):
        letter = chr(ord("A") + index - 1)
        actual = ws.column_dimensions[letter].width
        if not _close(actual, width, 0.02):
            findings.append(Finding(ws.title, letter, "column_width"))
    for row in range(1, header_row):
        height = ws.row_dimensions[row].height
        if height is not None and not _close(height, META_ROW_HEIGHT, 0.2):
            findings.append(Finding(ws.title, f"A{row}", "meta_row_height"))
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col)
            if cell.value in (None, ""):
                continue
            if not _font_ok(cell):
                findings.append(Finding(ws.title, cell.coordinate, "font"))
            if cell.border and cell.border.left and cell.border.left.style:
                findings.append(Finding(ws.title, cell.coordinate, "meta_border"))
    header_height = ws.row_dimensions[header_row].height
    if not _close(header_height, HEADER_ROW_HEIGHT, 0.2):
        findings.append(Finding(ws.title, f"A{header_row}", "header_row_height"))
    last_row = header_row
    for row in range(header_row, (ws.max_row or header_row) + 1):
        if any(ws.cell(row, col).value not in (None, "") for col in range(1, len(headers) + 1)):
            last_row = row
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, col)
        if not (cell.font and cell.font.bold and _font_ok(cell)):
            findings.append(Finding(ws.title, cell.coordinate, "header_font"))
        if (cell.alignment.horizontal or "") != "center":
            findings.append(Finding(ws.title, cell.coordinate, "header_align"))
        if not (cell.border and cell.border.left and cell.border.left.style == "thin"):
            findings.append(Finding(ws.title, cell.coordinate, "border"))
    for row in range(header_row + 1, last_row + 1):
        height = ws.row_dimensions[row].height or 0
        if height > MAX_DATA_ROW_HEIGHT + 0.2:
            findings.append(Finding(ws.title, f"A{row}", "row_too_tall"))
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row, col)
            if not _font_ok(cell):
                findings.append(Finding(ws.title, cell.coordinate, "font"))
            if not (cell.border and cell.border.left and cell.border.left.style == "thin"):
                findings.append(Finding(ws.title, cell.coordinate, "border"))
            if cell.fill and cell.fill.patternType:
                findings.append(Finding(ws.title, cell.coordinate, "fill"))
            if header in CENTER_NO_WRAP:
                if cell.alignment.horizontal != "center" or cell.alignment.wrap_text:
                    findings.append(Finding(ws.title, cell.coordinate, "id_alignment"))
                if not cell.alignment.shrink_to_fit:
                    findings.append(Finding(ws.title, cell.coordinate, "shrink_to_fit"))
            elif header in DATE_HEADERS:
                if cell.alignment.horizontal != "center" or not cell.alignment.wrap_text:
                    findings.append(Finding(ws.title, cell.coordinate, "date_alignment"))
                if isinstance(cell.value, datetime) and cell.number_format != DATE_FORMAT:
                    findings.append(Finding(ws.title, cell.coordinate, "date_format"))
            else:
                if (cell.alignment.horizontal or "") not in {"left", "general"} or not cell.alignment.wrap_text:
                    findings.append(Finding(ws.title, cell.coordinate, "text_alignment"))
    return findings


def _audit_checklist(ws: Worksheet) -> list[Finding]:
    findings: list[Finding] = []
    if (ws.page_setup.orientation or "").lower() != "landscape":
        findings.append(Finding(ws.title, "A1", "orientation"))
    if ws.page_setup.scale != CHECKLIST_SCALE:
        findings.append(Finding(ws.title, "A1", "checklist_scale"))
    if str(ws.page_setup.paperSize or "") not in {str(PAPER_LETTER), "1"}:
        findings.append(Finding(ws.title, "A1", "paper_size"))
    margins = ws.page_margins
    if not _close(margins.left, MARGIN_LEFT) or not _close(margins.top, MARGIN_TOP):
        findings.append(Finding(ws.title, "A1", "margin"))
    last_col = min(ws.max_column or 1, len(CHECKLIST_WIDTHS))
    for index, width in enumerate(CHECKLIST_WIDTHS[:last_col], start=1):
        letter = chr(ord("A") + index - 1)
        if not _close(ws.column_dimensions[letter].width, width, 0.02):
            findings.append(Finding(ws.title, letter, "column_width"))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=last_col):
        height = ws.row_dimensions[row[0].row].height or 0
        if height > MAX_DATA_ROW_HEIGHT + 0.2:
            findings.append(Finding(ws.title, f"A{row[0].row}", "row_too_tall"))
        for cell in row:
            if cell.value in (None, ""):
                continue
            if not _font_ok(cell):
                findings.append(Finding(ws.title, cell.coordinate, "font"))
            if not (cell.border and cell.border.left and cell.border.left.style == "thin"):
                findings.append(Finding(ws.title, cell.coordinate, "border"))
            if cell.fill and cell.fill.patternType:
                findings.append(Finding(ws.title, cell.coordinate, "fill"))
    return findings


def audit_workbook(wb) -> list[Finding]:
    findings: list[Finding] = []
    recognized = False
    for ws in wb.worksheets:
        if find_index(ws):
            recognized = True
            findings.extend(_audit_index(ws))
        elif is_checklist(ws):
            recognized = True
            findings.extend(_audit_checklist(ws))
    if not recognized:
        findings.append(Finding("", "", "unrecognized"))
    return findings
