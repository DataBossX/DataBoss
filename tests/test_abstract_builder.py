"""Tests for abstract_builder: normalizers (standing-notes examples), renderer geometry, QA gates, census."""
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from abstract_builder import donor as D
from abstract_builder.normalize import (norm_book_page, norm_comment, norm_date, norm_doc_no, norm_int, norm_legal,
                                        norm_parties, normalize_blm_row, normalize_county_row)
from abstract_builder.qa import check_index
from abstract_builder.render import IndexMeta, write_index


@pytest.mark.parametrize("src,exp", [("004928", 4928), (52035, 52035), ("Unknown", None), ("", None), ("#104303", 104303)])
def test_doc_no(src, exp):
    assert norm_doc_no(src)[0] == exp


def test_doc_no_non_numeric_kept_and_logged():
    v, n = norm_doc_no("RW21-266")
    assert v == "RW21-266" and any("NON_NUMERIC" in x for x in n)


@pytest.mark.parametrize("src,exp", [
    ("228-21-22", "0228-0021-0022"), ("0228-0025/0028", "0228-0025-0028"), ("0296-0095 through 0269", "0296-0095-0269"),
    ("86A-83-0071", "86A83-0071"), ("87A42-264A", "87A42-264A"), ("A060-0046", "A060-0046"), ("67MR-350", "67MR-0350"),
    ("R/W 32-623", "R/W32-0623"), ("0224 (Mining)-0001", "0224 (Mining)-0001"), ("Book 224 Page 769", "0224-0769"),
    ("399-142", "0399-0142")])
def test_book_page(src, exp):
    assert norm_book_page(src)[0] == exp


def test_dates_never_invented():
    assert norm_date("1985")[0] is None
    assert norm_date("9/1985")[0] is None
    assert norm_date("Unknown")[0] is None
    assert norm_date("2/8/2022")[0] == date(2022, 2, 8)
    assert norm_date("September 27, 1985")[0] == date(1985, 9, 27)


def test_parties():
    assert norm_parties("Alpha Petroleum Corporation (releasing party)")[0] == "Alpha Petroleum Corporation"
    assert norm_parties("Beta Meadow Resources L.L.C.")[0] == "Beta Meadow Resources, LLC"
    assert norm_parties("A Corp\nB Corp")[0] == "A Corp,\nB Corp"
    assert norm_parties("Gamma Partners II LLP; Delta Resources Assets LLC")[0] == "Gamma Partners II LLP,\nDelta Resources Assets, LLC"
    assert norm_parties("Epsilon Exploration Corporation, a Delaware corporation")[0].endswith("a Delaware corporation")
    assert norm_parties("John A. Smith, Jane B. Smith, HW")[0] == "John A. Smith, Jane B. Smith, HW"
    assert norm_parties("Unknown")[0] is None
    assert norm_parties("Wells Fargo Bank, N.A. (mortgagee)")[0] == "Wells Fargo Bank, N.A."
    assert norm_parties("Operator (Operator)")[0] == "Operator (Operator)"


def test_legal_house_format():
    f = lambda v, s=19: norm_legal(v, s, 47, 76)[0]
    assert f("N/2, Section 19, T47N-R76W, Johnson County, Wyoming, aol") == "N/2 of 19-47N-76W, aol"
    assert f("Lots 3,4,5,6,7 (containing 160 acres) 6th P.M., and other lands", 6) == "Lots 3, 4, 5, 6, 7 of 6-47N-76W, aol"
    assert f("E2W2 of 31-47N-76W, insofar as to the Parkman formation", 31) == "E/2 W/2 of 31-47N-76W"
    assert f("All", 31) == "All of 31-47N-76W"
    assert f("outside lands only") == "outside lands only"
    assert f("Lots 5, 6 of Section 31, Township 47 North, Range 76 West, well locations NE/4 NE/4, aol", 31) == "Lots 5, 6 of 31-47N-76W, aol"
    assert norm_legal("Unknown")[0] is None


def test_legal_never_invents_section():
    v, n = norm_legal("N/2")           # no section context given
    assert v == "N/2" and any("LEGAL_FORMAT_REVIEW" in x for x in n)


def test_blm_casewide_newline_and_aol():
    v = norm_legal("Lots 5, 7, 10, 12 of 13-47N-77W\nLots 2, 4, 5, 7 of 24-47N-77W, aol", blm_casewide=True)[0]
    assert v == "Lots 5, 7, 10, 12 of 13-47N-77W,\nLots 2, 4, 5, 7 of 24-47N-77W, aol"


def test_comment_only_missing_image():
    assert norm_comment("Missing Image")[0] == "missing image"
    assert norm_comment("see Exhibit B")[0] is None
    assert norm_comment("")[0] is None


def test_page_part_single_integer():
    assert norm_int("p. 12-14", "PAGE")[0] == 12
    assert norm_int(81, "PAGE")[0] == 81
    assert norm_int("", "PART")[0] is None
    assert norm_int(0, "PART")[0] is None


def test_blm_row_clears_file_no_and_comments():
    r = normalize_blm_row(["Assignment", "A LLC", "B LLC", "F-1", "68", "12/30/2016", "", "note", "Lots 1 of 6-47N-76W", "2"])
    assert r.values[3] is None and r.values[7] is None
    assert r.values[4] == 68 and r.values[9] == 2 and r.values[5] == date(2016, 12, 30)


def test_county_row_end_to_end():
    r = normalize_county_row(["Mortgage", "Zeta Bank (mortgagee)", "Eta Pipeline L.L.C.", "004928", "228-21-22",
                              "2/8/2022", "February 9, 2022", "N/2, Section 19, T47N-R76W, Johnson County", "Missing image"], 19, 47, 76)
    assert r.values == ["Mortgage", "Zeta Bank", "Eta Pipeline, LLC", 4928, "0228-0021-0022", date(2022, 2, 8),
                        date(2022, 2, 9), "N/2 of 19-47N-76W", "missing image"]


def _meta():
    return IndexMeta("Example", "T47N-R76W Section 6: All", date(2026, 10, 20), None, date(2026, 10, 6))


def test_render_county_passes_qa(tmp_path: Path):
    rows = [normalize_county_row(["Deed", "A\nB", "C LLC", "123", "399-142", "2/8/2022", "2/9/2022", "All of 6-47N-76W", None], 6, 47, 76).values]
    x = write_index(rows, _meta(), "county", tmp_path / "c.xlsx")
    assert check_index(x) == []
    ws = openpyxl.load_workbook(x).active
    assert [ws.cell(9, c).value for c in range(1, 10)] == D.COUNTY_HEADERS
    assert ws["D10"].value == 123 and ws["E10"].value == "0399-0142" and ws["F10"].number_format == "m/d/yyyy"
    assert ws["I10"].border.left.style == "thin"                    # blank cells inside used rows are bordered
    assert ws.row_dimensions[10].height is None                     # auto-fit, never fixed
    assert ws["A10"].font.name == "Calibri" and ws["A10"].font.size == 11


def test_render_blm_columns(tmp_path: Path):
    rows = [normalize_blm_row(["Assignment", "A", "B", None, 68, "12/30/2016", None, None, "Lots 1 of 13-47N-77W", 2]).values]
    x = write_index(rows, IndexMeta("Example", "WYW-000001", date(2026, 10, 20), None, date(2026, 10, 13)), "blm", tmp_path / "b.xlsx")
    assert check_index(x) == []
    ws = openpyxl.load_workbook(x).active
    assert [ws.cell(9, c).value for c in range(1, 11)] == D.BLM_HEADERS
    assert ws["E10"].value == 68 and ws["J10"].value == 2 and ws["D10"].value is None and ws["H10"].value is None


def test_qa_catches_unknown_and_bad_font(tmp_path: Path):
    rows = [["Deed", "A", "Unknown", 1, "0001-0001", date(2020, 1, 1), date(2020, 1, 2), "All of 6-47N-76W", None]]
    x = write_index(rows, _meta(), "county", tmp_path / "c.xlsx")
    assert any(f.startswith("UNKNOWN_IN_CLIENT_CELL") for f in check_index(x))
    wb = openpyxl.load_workbook(x); ws = wb.active
    from openpyxl.styles import Font
    ws["A10"].font = Font(name="Arial", size=11)
    wb.save(x)
    assert any(f.startswith("FONT:") for f in check_index(x))


def test_donor_print_setup(tmp_path: Path):
    x = write_index([], _meta(), "county", tmp_path / "c.xlsx")
    ws = openpyxl.load_workbook(x).active
    pm = ws.page_margins
    assert (pm.left, pm.right, pm.top, pm.bottom) == (0.7, 0.7, 0.3, 0.3)
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
