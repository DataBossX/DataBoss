"""Deterministic abstract-table, reconciliation, and print-layout QA tests."""

from pathlib import Path

import openpyxl
import pytest

from horizon.workbook_qa import CheckResult, QAReport, inspect_workbook


CHECKS = [
    "abstract_required_fields",
    "abstract_counts",
    "abstract_key_reconciliation",
    "abstract_print_layout",
]
HEADERS = [
    "Instrument Number",
    "Document Type",
    "Grantor",
    "Grantee",
    "Recorded Date",
    "Legal Description",
]
ROWS = [
    [
        "2024-001",
        "Mineral Deed",
        "Alpha LLC",
        "Beta LLC",
        "2024-01-02",
        "Section 15, T45N, R76W",
    ],
    [
        "2024-002",
        "Assignment",
        "Beta LLC",
        "Gamma LLC",
        "2024-02-03",
        "Section 15, T45N, R76W",
    ],
]


def _table_profile(table_id: str, sheet: str) -> dict:
    return {
        "id": table_id,
        "sheet": sheet,
        "header_row": 1,
        "start_row": 2,
        "key_field": "instrument_number",
        "key_normalization": "alnum_upper",
        "columns": {
            "instrument_number": "A",
            "document_type": "B",
            "grantor": "C",
            "grantee": "D",
            "recorded_date": "E",
            "legal_description": "F",
        },
        "expected_headers": {
            "instrument_number": ["Instrument Number", "Instrument No."],
            "document_type": "Document Type",
            "grantor": "Grantor",
            "grantee": "Grantee",
            "recorded_date": "Recorded Date",
            "legal_description": "Legal Description",
        },
        "required_fields": [
            "instrument_number",
            "document_type",
            "grantor",
            "grantee",
            "recorded_date",
            "legal_description",
        ],
        "expected_rows": 2,
        "expected_unique_keys": 2,
    }


def _profile() -> dict:
    return {
        "abstract_tables": [
            _table_profile("master", "Master"),
            _table_profile("index", "Index"),
        ],
        "abstract_key_reconciliations": [
            {
                "left_table": "master",
                "right_table": "index",
                "mode": "exact",
            }
        ],
        "abstract_print_layout": [
            {
                "sheet": sheet,
                "orientation": "landscape",
                "paper_size": 1,
                "fit_to_width": 1,
                "fit_to_height": 0,
                "fit_to_page": True,
                "print_area": "$A$1:$F$3",
                "print_title_rows": "$1:$1",
                "freeze_panes": "A2",
            }
            for sheet in ("Master", "Index")
        ],
    }


def _make_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    master = workbook.active
    master.title = "Master"
    index = workbook.create_sheet("Index")
    for worksheet in (master, index):
        worksheet.append(HEADERS)
        for row in ROWS:
            worksheet.append(row)
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_area = "A1:F3"
        worksheet.print_title_rows = "1:1"
        worksheet.freeze_panes = "A2"
    workbook.save(path)
    workbook.close()


def _check(report: QAReport, check_id: str) -> CheckResult:
    return next(item for item in report.checks if item.check_id == check_id)


def test_complete_abstract_tables_pass_all_gates(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)

    report = inspect_workbook(candidate, CHECKS, profile=_profile())

    assert report.score.technical_pass
    assert all(check.passed for check in report.checks)
    assert _check(report, "abstract_required_fields").metrics["fields_checked"] == 24
    assert _check(report, "abstract_counts").metrics["unique_keys_counted"] == 4
    assert _check(report, "abstract_key_reconciliation").metrics[
        "comparisons_checked"
    ] == 1
    assert _check(report, "abstract_print_layout").metrics[
        "assertions_checked"
    ] == 16


def test_blank_required_legal_cell_blocks_candidate(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    workbook = openpyxl.load_workbook(candidate)
    workbook["Master"]["F2"] = None
    workbook.save(candidate)
    workbook.close()

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=_profile()
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_required_field_blank"
        and finding.sheet == "Master"
        and finding.cell == "F2"
        for finding in check.findings
    )


def test_authoritative_count_mismatch_blocks_candidate(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["expected_rows"] = 3

    report = inspect_workbook(candidate, ["abstract_counts"], profile=profile)

    check = _check(report, "abstract_counts")
    assert not check.passed
    assert any(
        finding.code == "abstract_count_mismatch" for finding in check.findings
    )


def test_duplicate_key_blocks_counts_and_reconciliation(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    workbook = openpyxl.load_workbook(candidate)
    workbook["Index"]["A3"] = "2024-001"
    workbook.save(candidate)
    workbook.close()

    report = inspect_workbook(
        candidate,
        ["abstract_counts", "abstract_key_reconciliation"],
        profile=_profile(),
    )

    assert any(
        finding.code == "abstract_key_duplicate"
        for check in report.checks
        for finding in check.findings
    )
    assert not report.score.technical_pass


def test_master_index_key_set_mismatch_blocks_candidate(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    workbook = openpyxl.load_workbook(candidate)
    workbook["Index"]["A3"] = "2024-999"
    workbook.save(candidate)
    workbook.close()

    report = inspect_workbook(
        candidate, ["abstract_key_reconciliation"], profile=_profile()
    )

    codes = {
        finding.code
        for finding in _check(report, "abstract_key_reconciliation").findings
    }
    assert codes == {
        "abstract_keys_missing_left",
        "abstract_keys_missing_right",
    }


def test_print_layout_mismatch_blocks_candidate(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    workbook = openpyxl.load_workbook(candidate)
    workbook["Master"].page_setup.orientation = "portrait"
    workbook.save(candidate)
    workbook.close()

    report = inspect_workbook(candidate, ["abstract_print_layout"], profile=_profile())

    check = _check(report, "abstract_print_layout")
    assert not check.passed
    assert any(
        finding.code == "abstract_layout_mismatch"
        and "orientation" in finding.message
        for finding in check.findings
    )


def test_missing_gate_rules_fail_closed(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)

    report = inspect_workbook(candidate, CHECKS, profile={})

    assert not report.score.technical_pass
    assert all(check.status == "not_evaluated" for check in report.checks)


def test_malformed_required_fields_profile_fails(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["required_fields"].append("missing_column")

    report = inspect_workbook(candidate, ["abstract_required_fields"], profile=profile)

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        for finding in check.findings
    )


def test_required_fields_string_cannot_be_silently_skipped(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["required_fields"] = "instrument_number"

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        and finding.sheet == "Master"
        for finding in check.findings
    )


@pytest.mark.parametrize("invalid_end_row", [1_048_577, 3.5, True])
def test_invalid_excel_row_bound_is_a_blocking_finding(
    tmp_path, invalid_end_row
):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["end_row"] = invalid_end_row

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        for finding in check.findings
    )


def test_unknown_print_setting_is_not_ignored(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_print_layout"][0]["orientaton"] = "landscape"

    report = inspect_workbook(
        candidate, ["abstract_print_layout"], profile=profile
    )

    check = _check(report, "abstract_print_layout")
    assert not check.passed
    assert check.findings[0].code == "abstract_layout_profile_invalid"
    assert "orientaton" in check.findings[0].message


def test_boolean_cannot_match_numeric_paper_size(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_print_layout"][0]["paper_size"] = True

    report = inspect_workbook(
        candidate, ["abstract_print_layout"], profile=profile
    )

    check = _check(report, "abstract_print_layout")
    assert not check.passed
    assert check.findings[0].code == "abstract_layout_profile_invalid"
    assert "paper_size" in check.findings[0].message


def test_non_string_header_expectation_is_profile_error(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["expected_headers"][
        "instrument_number"
    ] = False

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        and "non-empty string" in finding.message
        for finding in check.findings
    )


def test_column_beyond_xfd_is_profile_error(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["columns"]["legal_description"] = "XFE"

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        and finding.sheet == "Master"
        for finding in check.findings
    )


def test_unknown_abstract_table_setting_is_not_ignored(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["expected_headders"] = {
        "instrument_number": "Instrument Number"
    }

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        and "expected_headders" in finding.message
        for finding in check.findings
    )


def test_normalized_column_field_collision_is_profile_error(tmp_path):
    candidate = tmp_path / "candidate.xlsx"
    _make_workbook(candidate)
    profile = _profile()
    profile["abstract_tables"][0]["columns"]["instrument_number "] = "J"

    report = inspect_workbook(
        candidate, ["abstract_required_fields"], profile=profile
    )

    check = _check(report, "abstract_required_fields")
    assert not check.passed
    assert any(
        finding.code == "abstract_table_profile_invalid"
        for finding in check.findings
    )
