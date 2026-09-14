"""Source-proved deltas write an isolated copy only."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from horizon.isolated_delta import (
    IsolatedDeltaError,
    apply_deltas,
    sha256_file,
    write_delta_packet,
)
from horizon.stable_key import stable_key


PENTERRA_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "Doc No",
    "Book-Page",
    "Date of Doc",
    "Rec Date",
    "Legal Description",
    "Comments",
]


def _penterra_index(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    for label in (
        "Index County",
        "Lands",
        "Date",
        "Starting Date",
        "Date Posted Thru",
        "Indexed By",
        "Project",
    ):
        sheet.append([label, "SYNTH"])
    sheet.append(PENTERRA_HEADERS)
    sheet.append(
        [
            "",
            "SYNTH ALPHA LLC",
            "SYNTH BETA LLC",
            "2026-00020",
            "1-2",
            "",
            "01/02/2026",
            "",
            "",
        ]
    )
    sheet.append(
        [
            "ASSIGNMENT",
            "SYNTH GAMMA LLC",
            "SYNTH DELTA LLC",
            "933403",
            "3-4",
            "01/01/2001",
            "02/02/2001",
            "SYNTH T45N R76W SEC 13",
            "",
        ]
    )
    workbook.save(path)
    workbook.close()


def _packet(source_sha: str, **overrides) -> dict:
    packet = {
        "schema_id": "dbx.source_proved_delta_packet",
        "schema_version": "1.0",
        "packet_id": "SYNTH-P13-LEGAL",
        "source_workbook_sha256": source_sha,
        "deltas": [
            {
                "row_key": "2026-00020|1-2",
                "field": "legal_description",
                "value": "SYNTH T45N R76W SEC 13",
                "source_sha256": "a" * 64,
                "page": 1,
                "crop_id": "legal",
                "replace": False,
            }
        ],
    }
    packet.update(overrides)
    return packet


def test_apply_fills_blank_required_cell_on_copy_only(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _penterra_index(source)
    source_sha = sha256_file(source)

    receipt = apply_deltas(source, output, _packet(source_sha))

    assert receipt.applied == 1
    assert receipt.rejected == 0
    assert receipt.technical_pass is True
    assert sha256_file(source) == source_sha
    source_wb = load_workbook(source, data_only=True)
    output_wb = load_workbook(output, data_only=True)
    assert source_wb["Index"]["H9"].value in (None, "")
    assert output_wb["Index"]["H9"].value == "SYNTH T45N R76W SEC 13"
    assert output_wb["Index"]["H10"].value == "SYNTH T45N R76W SEC 13"
    source_wb.close()
    output_wb.close()


def test_same_docno_different_book_page_targets_the_matching_row(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    for _ in range(7):
        sheet.append(["SYNTH"])
    sheet.append(PENTERRA_HEADERS)
    sheet.append(["", "A", "B", "100", "10-1", "", "1/1/2001", "", ""])
    sheet.append(["", "C", "D", "100", "10-2", "", "1/2/2001", "", ""])
    workbook.save(source)
    workbook.close()

    packet = _packet(
        sha256_file(source),
        deltas=[
            {
                "row_key": "100|10-2",
                "field": "legal_description",
                "value": "SYNTH SECOND ROW ONLY",
                "source_sha256": "b" * 64,
                "page": 2,
                "crop_id": "legal",
                "replace": False,
            }
        ],
    )
    receipt = apply_deltas(source, output, packet)
    assert receipt.applied == 1
    assert stable_key("100", "10-2") == "100|0010-0002"
    output_wb = load_workbook(output, data_only=True)
    assert output_wb["Index"]["H9"].value in (None, "")
    assert output_wb["Index"]["H10"].value == "SYNTH SECOND ROW ONLY"
    output_wb.close()


def test_refuse_overwrite_without_replace(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _penterra_index(source)
    packet = _packet(
        sha256_file(source),
        deltas=[
            {
                "row_key": "933403|3-4",
                "field": "legal_description",
                "value": "SYNTH OTHER LEGAL",
                "source_sha256": "c" * 64,
                "page": 1,
                "crop_id": "legal",
                "replace": False,
            }
        ],
    )
    with pytest.raises(IsolatedDeltaError, match="already has"):
        apply_deltas(source, output, packet)
    assert not output.exists()


def test_refuse_missing_source_hash_and_unknown_row(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _penterra_index(source)
    with pytest.raises(IsolatedDeltaError, match="hash"):
        apply_deltas(source, output, _packet("0" * 64))
    assert not output.exists()

    packet = _packet(sha256_file(source))
    packet["deltas"][0]["row_key"] = "2026-09999|9-9"
    with pytest.raises(IsolatedDeltaError, match="unknown row_key"):
        apply_deltas(source, output, packet)
    assert not output.exists()


def test_refuse_double_comma_and_instrument_number_rewrite(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _penterra_index(source)
    digest = sha256_file(source)
    bad_value = _packet(
        digest,
        deltas=[
            {
                "row_key": "2026-00020|1-2",
                "field": "legal_description",
                "value": "A,,B",
                "source_sha256": "d" * 64,
                "page": 1,
                "crop_id": "legal",
                "replace": False,
            }
        ],
    )
    with pytest.raises(IsolatedDeltaError, match="unsafe"):
        apply_deltas(source, output, bad_value)
    bad_field = _packet(
        digest,
        deltas=[
            {
                "row_key": "2026-00020|1-2",
                "field": "instrument_number",
                "value": "2026-00021",
                "source_sha256": "e" * 64,
                "page": 1,
                "crop_id": "docno",
                "replace": False,
            }
        ],
    )
    with pytest.raises(IsolatedDeltaError, match="not allowed"):
        apply_deltas(source, output, bad_field)
    assert not output.exists()


def test_write_delta_packet_binds_workbook_hash_without_inventing(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _penterra_index(source)
    output = tmp_path / "delta-packet.json"
    packet = write_delta_packet(
        workbook=source,
        deltas=_packet("0" * 64)["deltas"],
        output=output,
        packet_id="SYNTH-P13-LEGAL",
    )
    assert packet["source_workbook_sha256"] == sha256_file(source)
    isolated = tmp_path / "copy.xlsx"
    receipt = apply_deltas(source, isolated, packet)
    assert receipt.applied == 1
    with pytest.raises(IsolatedDeltaError, match="not allowed"):
        write_delta_packet(
            workbook=source,
            deltas=[
                {
                    "row_key": "2026-00020|1-2",
                    "field": "instrument_number",
                    "value": "2026-00021",
                    "source_sha256": "e" * 64,
                    "page": 1,
                    "crop_id": "docno",
                    "replace": False,
                }
            ],
            output=tmp_path / "bad.json",
            packet_id="SYNTH-BAD",
        )
