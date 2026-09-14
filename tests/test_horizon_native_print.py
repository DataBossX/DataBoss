"""Isolated Letter layout repair and native Excel Print Preview binding."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from horizon.isolated_delta import sha256_file
from horizon.native_print import NativePrintError, assess_native_print
from horizon.package_finish import run_finish
from horizon.print_layout_repair import PrintLayoutRepairError, repair_print_layout


PENTERRA_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "Doc No",
    "Book-Page",
    "Date of Doc",
    "Rec Date",
    "Legal Description",
    "Comments",
]


def _write_index(path: Path, *, a4: bool = False, titles: bool = False) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    for _ in range(7):
        sheet.append(["SYNTH"])
    sheet.append(PENTERRA_HEADERS)
    sheet.append(
        [
            "Mineral Deed",
            "SYNTH ALPHA LLC",
            "SYNTH BETA LLC",
            "2026-09901",
            "",
            "1/1/2026",
            "1/2/2026",
            "SYNTH TRACT 15-45N-76W",
            "",
        ]
    )
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4 if a4 else sheet.PAPERSIZE_LETTER
    if titles:
        sheet.print_title_rows = "1:8"
    workbook.save(path)
    workbook.close()


def _receipt(workbook_sha: str, **overrides: object) -> dict[str, object]:
    packet = {
        "schema_id": "dbx.native_print_receipt",
        "schema_version": "1.0",
        "packet_id": "SYNTH-P15-PRINT",
        "workbook_sha256": workbook_sha,
        "application": "Microsoft Excel",
        "host": "Windows",
        "paper_size": 1,
        "orientation": "landscape",
        "page_count": 2,
        "expected_page_count": 2,
        "print_titles": True,
        "print_area_set": True,
        "operator": "SYNTH OPERATOR",
    }
    packet.update(overrides)
    return packet


def test_print_layout_repair_writes_letter_landscape_on_copy_only(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _write_index(source, a4=True)
    source_sha = sha256_file(source)
    receipt = repair_print_layout(source, output)
    assert receipt.technical_pass is True
    assert sha256_file(source) == source_sha
    source_wb = openpyxl.load_workbook(source)
    copy_wb = openpyxl.load_workbook(output)
    assert int(source_wb["Index"].page_setup.paperSize) == 9
    assert int(copy_wb["Index"].page_setup.paperSize) == 1
    assert copy_wb["Index"].page_setup.orientation == "landscape"
    titles = str(copy_wb["Index"].print_title_rows or "").replace("$", "")
    assert titles.endswith("1:8")
    source_wb.close()
    copy_wb.close()
    with pytest.raises(PrintLayoutRepairError, match="already exist"):
        repair_print_layout(source, output)


def test_native_print_accepts_letter_excel_and_rejects_a4_or_libreoffice(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "index.xlsx"
    _write_index(workbook, titles=True)
    digest = sha256_file(workbook)
    ok = assess_native_print(_receipt(digest), workbook=workbook)
    assert ok.technical_pass is True
    assert ok.page_count == 2

    a4 = assess_native_print(_receipt(digest, paper_size=9), workbook=workbook)
    assert a4.technical_pass is False
    assert any("Letter" in issue for issue in a4.issues)

    soffice = assess_native_print(
        _receipt(digest, application="LibreOffice", host="Linux"),
        workbook=workbook,
    )
    assert soffice.technical_pass is False

    with pytest.raises(NativePrintError, match="hash"):
        assess_native_print(_receipt("0" * 64), workbook=workbook)

    four = assess_native_print(
        _receipt(digest, page_count=4, expected_page_count=2),
        workbook=workbook,
    )
    assert four.technical_pass is False
    assert any("page_count" in issue for issue in four.issues)


def test_finish_runner_repairs_layout_then_binds_native_receipt(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    isolated = tmp_path / "letter.xlsx"
    _write_index(source, a4=True)
    # Hash of the repaired copy is unknown until after the gate; bind after.
    first = run_finish(
        sections=[15],
        workbook=source,
        print_layout_output=isolated,
    )
    assert first.packages_complete is False
    assert first.technical_pass is True
    assert {gate.name for gate in first.gates} == {
        "print_layout_repair",
        "workbook_qa",
    }
    packet = tmp_path / "native.json"
    packet.write_text(
        json.dumps(_receipt(sha256_file(isolated))),
        encoding="utf-8",
    )
    second = run_finish(
        sections=[15],
        workbook=isolated,
        native_print_receipt=packet,
    )
    assert second.packages_complete is False
    assert second.technical_pass is True
    native = next(gate for gate in second.gates if gate.name == "native_print")
    assert native.detail["page_count"] == 2
