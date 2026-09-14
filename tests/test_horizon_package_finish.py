"""Package finish runner never promotes and chains existing gates."""

import hashlib
import json
from pathlib import Path

import openpyxl

from horizon.isolated_delta import sha256_file
from horizon.package_finish import main, run_finish
from horizon.workbook_qa import inspect_workbook, load_workbook_profile


PENTERRA_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "Doc No",
    "Book-Page",
    "Date of Doc",
    "Rec Date",
    "Legal Description",
    "Comments",
]
PENTERRA_PROFILE = (
    Path(__file__).resolve().parents[1]
    / "horizon"
    / "profiles"
    / "penterra_index_v1.json"
)


def _penterra_workbook(path: Path, *, blank_legal: bool = False) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    for label in (
        "Index County",
        "Lands",
        "Date",
        "Starting Date",
        "Date Posted Thru",
        "Indexed By",
        "Project",
    ):
        sheet.append([label, "SYNTH"])
    sheet.append(PENTERRA_HEADERS)
    sheet.append(
        [
            "Mineral Deed",
            "SYNTH ALPHA LLC",
            "SYNTH BETA LLC",
            "2026-09901",
            "",
            "1/1/2026",
            "1/2/2026",
            "" if blank_legal else "SYNTH TRACT 15-45N-76W",
            "",
        ]
    )
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.print_title_rows = "1:8"
    workbook.save(path)
    workbook.close()


def _sha(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _checkable_export() -> dict:
    return {
        "schema_id": "dbx.tract_ledger_export",
        "schema_version": "1.0",
        "packet_id": "SYNTH-FINISH-001",
        "rows": [
            {
                "row_id": "p01r01",
                "docno": "2026-09901",
                "bookpage": "",
                "rec_date": "1/2/2026",
                "doc_date": "",
                "grantor": "SYNTH SURVEYOR",
                "grantee": "The Public",
                "page": 1,
                "crop_id": "p01r01",
            }
        ],
    }


def _occurrence_packet() -> dict:
    key = "2026-09901|"
    digest = _sha("face-a")
    fields = {
        "party": "SYNTH SURVEYOR",
        "date_role": "recorded",
        "legal": "SYNTH TRACT",
    }
    return {
        "schema_id": "dbx.occurrence_packet",
        "schema_version": "1.0",
        "packet_id": "SYNTH-FINISH-OCC",
        "occurrences": [
            {
                "occurrence_id": "occ-001",
                "stable_key": key,
                "source_sha256": digest,
                "page": 1,
                "crop_id": "p01r01",
                "fields": fields,
                "risk": "novel",
            }
        ],
        "unique_key_ledger": [{"stable_key": key, "source_sha256": digest}],
        "candidate_rows": [
            {
                "row_id": "R6-001",
                "stable_key": key,
                "source_sha256": digest,
                "fields": fields,
            }
        ],
        "allowlist": [],
    }


def test_empty_run_is_blocked_and_never_complete() -> None:
    receipt = run_finish(sections=[15, 13, 11])
    assert receipt.packages_complete is False
    assert receipt.technical_pass is False
    assert receipt.gates == []
    assert any("Mount PC/Drive" in action for action in receipt.next_actions)


def test_chained_checkable_packets_pass_gates_but_not_packages(
    tmp_path: Path,
) -> None:
    export = tmp_path / "export.json"
    packet = tmp_path / "packet.json"
    plat = tmp_path / "plat.pdf"
    export.write_text(json.dumps(_checkable_export()), encoding="utf-8")
    packet.write_text(json.dumps(_occurrence_packet()), encoding="utf-8")
    plat.write_bytes(b"%PDF-1.1\n1 0 obj<<>>endobj\ntrailer<<>>\n")

    receipt = run_finish(
        sections=[15, 13, 11],
        tract_export=export,
        occurrence_packet=packet,
        public_plats=[f"campbell,45n,76w,{plat}"],
    )
    assert receipt.technical_pass is True
    assert receipt.packages_complete is False
    assert {gate.name for gate in receipt.gates} == {
        "reextraction",
        "occurrence_ledger",
        "public_cadastral",
    }
    assert all(gate.technical_pass for gate in receipt.gates)


def test_penterra_profile_blocks_blank_legal(tmp_path: Path) -> None:
    workbook = tmp_path / "index.xlsx"
    _penterra_workbook(workbook, blank_legal=True)
    report = inspect_workbook(
        workbook,
        ["abstract_required_fields"],
        profile=load_workbook_profile(PENTERRA_PROFILE),
    )
    assert not report.score.technical_pass
    assert any(
        finding.code == "abstract_required_field_blank"
        and finding.cell == "H9"
        for finding in report.findings
    )


def test_a4_index_fails_letter_print_layout(tmp_path: Path) -> None:
    workbook = tmp_path / "index.xlsx"
    _penterra_workbook(workbook)
    loaded = openpyxl.load_workbook(workbook)
    loaded["Index"].page_setup.paperSize = loaded["Index"].PAPERSIZE_A4
    loaded.save(workbook)
    loaded.close()
    receipt = run_finish(sections=[15], workbook=workbook)
    assert receipt.packages_complete is False
    assert receipt.technical_pass is False
    qa = next(gate for gate in receipt.gates if gate.name == "workbook_qa")
    assert qa.technical_pass is False
    assert any(
        finding["code"] == "abstract_layout_mismatch"
        and finding["cell"] == ""
        for finding in qa.detail["findings"]
    )


def test_workbook_gate_runs_on_isolated_penterra_index(tmp_path: Path) -> None:
    workbook = tmp_path / "index.xlsx"
    _penterra_workbook(workbook)
    receipt = run_finish(sections=[15], workbook=workbook)
    assert receipt.packages_complete is False
    assert receipt.technical_pass is True
    assert receipt.gates[0].name == "workbook_qa"
    assert receipt.gates[0].technical_pass is True


def test_delta_fill_then_qa_runs_on_copy_not_source(tmp_path: Path) -> None:
    source = tmp_path / "index.xlsx"
    isolated = tmp_path / "index-filled.xlsx"
    packet = tmp_path / "delta.json"
    _penterra_workbook(source, blank_legal=True)
    packet.write_text(
        json.dumps(
            {
                "schema_id": "dbx.source_proved_delta_packet",
                "schema_version": "1.0",
                "packet_id": "SYNTH-FINISH-DELTA",
                "source_workbook_sha256": sha256_file(source),
                "deltas": [
                    {
                        "row_key": "2026-09901|",
                        "field": "legal_description",
                        "value": "SYNTH TRACT 15-45N-76W",
                        "source_sha256": _sha("face-legal"),
                        "page": 1,
                        "crop_id": "legal",
                        "replace": False,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    receipt = run_finish(
        sections=[15],
        workbook=source,
        delta_packet=packet,
        delta_output=isolated,
    )
    assert receipt.packages_complete is False
    assert receipt.technical_pass is True
    assert {gate.name for gate in receipt.gates} == {
        "isolated_delta",
        "workbook_qa",
    }
    assert all(gate.technical_pass for gate in receipt.gates)
    source_wb = openpyxl.load_workbook(source, data_only=True)
    isolated_wb = openpyxl.load_workbook(isolated, data_only=True)
    assert source_wb["Index"]["H9"].value in (None, "")
    assert isolated_wb["Index"]["H9"].value == "SYNTH TRACT 15-45N-76W"
    source_wb.close()
    isolated_wb.close()


def test_cli_empty_run_writes_blocked_receipt(tmp_path: Path) -> None:
    output = tmp_path / "finish.json"
    result = main(["--output", str(output), "--section", "15"])
    assert result == 2
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["packages_complete"] is False
    assert payload["technical_pass"] is False
