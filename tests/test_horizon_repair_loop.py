"""Workbook export, proposal binding, and isolated repair loop."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from horizon.index_export import (
    IndexExportError,
    bind_delta_packet,
    build_index_packet,
    export_faces,
)
from horizon.isolated_delta import apply_deltas, sha256_file
from horizon.package_finish import run_finish
from horizon.repair_loop import RepairLoopError, run_repair_loop


PENTERRA_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "Doc No",
    "Book-Page",
    "Date of Doc",
    "Rec Date",
    "Legal Description",
    "Comments",
]


def _sha(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _write_index(
    path: Path,
    *,
    legal: str = "SYNTH TRACT 15-45N-76W",
    rec_date: object = "1/2/2026",
    extra_row: bool = False,
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    for _ in range(7):
        sheet.append(["SYNTH"])
    sheet.append(PENTERRA_HEADERS)
    sheet.append(
        [
            "Mineral Deed",
            "SYNTH ALPHA LLC",
            "SYNTH BETA LLC",
            "2026-09901",
            "",
            "1/1/2026",
            rec_date,
            legal,
            "",
        ]
    )
    if extra_row:
        sheet.append(
            [
                "Mineral Deed",
                "SYNTH ALPHA LLC",
                "SYNTH BETA LLC",
                "2026-09901",
                "",
                "1/1/2026",
                "1/2/2026",
                legal,
                "",
            ]
        )
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.print_title_rows = "1:8"
    workbook.save(path)
    workbook.close()


def _fields(legal: str = "SYNTH TRACT 15-45N-76W") -> dict[str, str]:
    return {
        "document_type": "Mineral Deed",
        "grantor": "SYNTH ALPHA LLC",
        "grantee": "SYNTH BETA LLC",
        "recorded_date": "1/2/2026",
        "legal_description": legal,
    }


def _face(source: str, legal: str = "SYNTH TRACT 15-45N-76W") -> dict[str, object]:
    return {
        "stable_key": "2026-09901|",
        "source_sha256": _sha(source),
        "page": 1,
        "crop_id": source,
        "fields": _fields(legal),
    }


def _packet() -> dict[str, object]:
    return build_index_packet(
        "SYNTH-REPAIR-001",
        master=[_face("master")],
        pdf_index=[_face("pdf")],
        handwritten_index=[],
        candidate_rows=[],
        orphan_allowlist=[
            {
                "stable_key": "2026-09901|",
                "present_in": ["master", "pdf_index"],
                "source_sha256": _sha("master"),
            }
        ],
    )


def test_export_faces_reads_penterra_row_and_excel_dates(tmp_path: Path) -> None:
    workbook = tmp_path / "index.xlsx"
    _write_index(workbook, rec_date=datetime(2026, 1, 2))
    faces = export_faces(workbook)
    assert len(faces) == 1
    assert faces[0]["stable_key"] == "2026-09901|"
    assert faces[0]["fields"]["recorded_date"] == "1/2/2026"
    assert faces[0]["fields"]["legal_description"] == "SYNTH TRACT 15-45N-76W"
    assert faces[0]["page"] == 9


def test_export_refuses_duplicate_stable_keys(tmp_path: Path) -> None:
    workbook = tmp_path / "dup.xlsx"
    _write_index(workbook, extra_row=True)
    with pytest.raises(IndexExportError, match="duplicate"):
        export_faces(workbook)


def test_bind_delta_packet_strips_confidence_and_applies(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _write_index(source, legal="")
    packet = bind_delta_packet(
        [
            {
                "row_key": "2026-09901|",
                "field": "legal_description",
                "value": "SYNTH TRACT 15-45N-76W",
                "source_sha256": _sha("master"),
                "page": 1,
                "crop_id": "legal",
                "replace": False,
                "confidence": "medium",
                "agreeing_sources": ["master", "pdf_index"],
            }
        ],
        source_workbook_sha256=sha256_file(source),
        packet_id="SYNTH-BIND",
    )
    assert set(packet["deltas"][0]) == {
        "row_key",
        "field",
        "value",
        "source_sha256",
        "page",
        "crop_id",
        "replace",
    }
    receipt = apply_deltas(source, output, packet)
    assert receipt.applied == 1
    loaded = openpyxl.load_workbook(output, data_only=True)
    assert loaded["Index"]["H9"].value == "SYNTH TRACT 15-45N-76W"
    loaded.close()
    with pytest.raises(IndexExportError, match="high or medium"):
        bind_delta_packet(
            [
                {
                    "row_key": "2026-09901|",
                    "field": "legal_description",
                    "value": "SYNTH TRACT 15-45N-76W",
                    "source_sha256": _sha("master"),
                    "page": 1,
                    "crop_id": "legal",
                    "replace": False,
                    "confidence": "low",
                }
            ],
            source_workbook_sha256=sha256_file(source),
            packet_id="SYNTH-BIND-LOW",
        )


def test_repair_loop_fills_blank_on_copy_across_passes(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    out_dir = tmp_path / "repair"
    _write_index(source, legal="")
    source_sha = sha256_file(source)
    receipt = run_repair_loop(
        workbook=source,
        output_dir=out_dir,
        packet_id="SYNTH-REPAIR-001",
        index_packet=_packet(),
        max_loops=10,
    )
    assert receipt.packages_complete is False
    assert receipt.technical_pass is True
    assert receipt.original_workbook == str(source.resolve())
    assert sha256_file(source) == source_sha
    assert receipt.remaining_conflicts == 0
    isolated = Path(receipt.final_workbook)
    assert receipt.final_workbook_sha256 == sha256_file(isolated)
    assert isolated != source.resolve()
    source_wb = openpyxl.load_workbook(source, data_only=True)
    final_wb = openpyxl.load_workbook(isolated, data_only=True)
    assert source_wb["Index"]["H9"].value in (None, "")
    assert final_wb["Index"]["H9"].value == "SYNTH TRACT 15-45N-76W"
    source_wb.close()
    final_wb.close()
    assert any(item.applied == 1 for item in receipt.passes)
    assert any(item.stop_reason == "no_eligible_proposals" for item in receipt.passes)


def test_repair_loop_does_not_apply_conflicts(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    out_dir = tmp_path / "repair"
    _write_index(source, legal="")
    packet = build_index_packet(
        "SYNTH-CONFLICT",
        master=[_face("master", "SYNTH TRACT A")],
        pdf_index=[_face("pdf", "SYNTH TRACT B")],
        handwritten_index=[],
        candidate_rows=[],
        orphan_allowlist=[
            {
                "stable_key": "2026-09901|",
                "present_in": ["master", "pdf_index"],
                "source_sha256": _sha("master"),
            }
        ],
    )
    receipt = run_repair_loop(
        workbook=source,
        output_dir=out_dir,
        packet_id="SYNTH-CONFLICT",
        index_packet=packet,
    )
    assert receipt.packages_complete is False
    assert receipt.technical_pass is False
    assert receipt.passes[0].applied == 0
    assert receipt.passes[0].stop_reason == "reconciliation_blocked"
    assert list(out_dir.glob("*.xlsx")) == []


def test_finish_runner_repair_loop_never_completes_packages(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    packet_path = tmp_path / "index.json"
    out_dir = tmp_path / "repair"
    _write_index(source, legal="")
    packet_path.write_text(json.dumps(_packet()), encoding="utf-8")
    receipt = run_finish(
        sections=[15],
        workbook=source,
        index_packet=packet_path,
        repair_dir=out_dir,
        max_loops=10,
    )
    assert receipt.packages_complete is False
    assert receipt.technical_pass is True
    names = {gate.name for gate in receipt.gates}
    assert "repair_loop" in names
    assert "workbook_qa" in names
    with pytest.raises(RepairLoopError, match="empty"):
        run_repair_loop(
            workbook=source,
            output_dir=out_dir,
            packet_id="SYNTH-REPAIR-001",
            index_packet=_packet(),
        )
