from __future__ import annotations

import zipfile
from decimal import Decimal
from pathlib import Path

import pytest

import grocery_report_pipeline as grp
from horizon.repair import _HAVE_LXML, repair_workbook


def _read_cell_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml")]
        return zf.read(names[0]).decode("utf-8")


def _make_xlsx_with_error_formula(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ok"
    ws["A2"] = "=#REF!"
    wb.save(path)
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/media/plat1.png", b"\x89PNG\r\n\x1a\nFAKEPLATDATA")
    tmp.replace(path)


def _facts_from_text(tmp_path: Path, name: str, text: str) -> list[grp.Fact]:
    output = tmp_path / "out"
    extracted = output / "extracted_text"
    extracted.mkdir(parents=True)
    text_file = f"{name}.txt"
    (extracted / text_file).write_text(text, encoding="utf-8")
    rec = grp.FileRec(
        path=str(tmp_path / name),
        rel_path=name,
        name=name,
        ext=".txt",
        size_bytes=len(text),
        modified="2026-01-01",
        sha256="x",
        likely_type="deed",
    )
    tr = grp.TextRec(
        rel_path=name,
        abs_path=rec.path,
        sha256="x",
        method="read-text",
        ocr_used=False,
        char_count=len(text),
        text_file=text_file,
        status="ok",
    )
    return grp.extract_facts([rec], {rec.path: tr}, {rec.path: ["deed"]}, output, grp.BuildLog())


@pytest.mark.skipif(not _HAVE_LXML, reason="lxml required for XML repair")
def test_error_formula_cannot_become_non_error_literal(tmp_path):
    src = tmp_path / "report.xlsx"
    _make_xlsx_with_error_formula(src)
    original = src.read_bytes()
    dest = tmp_path / "report_v002.xlsx"
    result = repair_workbook(src, dest)
    assert result.refused
    assert result.output is None
    assert result.repaired is False
    assert not dest.exists()
    assert src.read_bytes() == original
    assert result.defects or "refused" in result.error.lower()
    xml = _read_cell_xml(src)
    assert 't="e"' in xml or "#REF!" in xml


@pytest.mark.skipif(not _HAVE_LXML, reason="lxml required for XML repair")
def test_malformed_worksheet_xml_is_hard_defect(tmp_path):
    src = tmp_path / "broken.xlsx"
    _make_xlsx_with_error_formula(src)
    tmp = src.with_suffix(".bad.xlsx")
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = b"<worksheet><broken"
            zout.writestr(item, data)
    tmp.replace(src)
    dest = tmp_path / "promoted.xlsx"
    result = repair_workbook(src, dest)
    assert result.refused
    assert result.output is None
    assert not dest.exists()
    assert "malformed" in result.error.lower() or any("malformed" in d for d in result.defects)


def test_unlabeled_date_is_not_recording_date(tmp_path):
    text = (
        "SYNTHETIC TEST DOCUMENT -- NOT REAL TITLE DATA\n"
        "MINERAL DEED\n"
        "Grantor: Foo Corp\n"
        "Grantee: Bar Trust\n"
        "Effective Date: 2019-03-15\n"
        "Legal: Section 12, T7N, R63W\n"
    )
    facts = _facts_from_text(tmp_path, "exec.txt", text)
    assert facts
    fact = facts[0]
    assert "recording_date" not in fact.values
    assert fact.values.get("effective_date") == "2019-03-15"
    assert "missing-recording-date" in fact.review_flags
    issues = grp.validate(
        [],
        {},
        {},
        facts,
        {"conflicts": []},
        tmp_path / "out",
        grp.BuildLog(),
    )
    assert any(i["rule"] == "missing-recording-date" for i in issues)


def test_ordinary_decimals_parse_and_incomplete_set_does_not_false_sum(tmp_path):
    assert grp._DECIMAL_RX.search("decimal interest 0.5")
    assert grp._DECIMAL_RX.search("decimal interest 0.25")
    assert grp._DECIMAL_RX.search("decimal interest 0.125")
    assert grp._DECIMAL_RX.findall("Owner A decimal interest 0.5\nOwner B decimal interest 0.5") == ["0.5", "0.5"]

    schedule = (
        "OWNERSHIP / mineral owner decimal interest schedule\n"
        "Owner Acme Minerals LLC decimal interest 0.5\n"
        "Owner Sample Family Trust decimal interest 0.5\n"
    )
    assert grp._owner_set_is_complete(schedule, [Decimal("0.5"), Decimal("0.5")]) is True

    lone_deed = (
        "SYNTHETIC TEST DOCUMENT -- NOT REAL TITLE DATA\n"
        "WARRANTY DEED\nGrantor: A\nGrantee: B\n"
        "decimal interest 0.5\nLegal: Section 12, T7N, R63W\n"
    )
    assert grp._owner_set_is_complete(lone_deed, [Decimal("0.5")]) is False
    facts = _facts_from_text(tmp_path, "deed.txt", lone_deed)
    recon = grp.reconcile(facts, tmp_path / "out", grp.BuildLog())
    types = [row[0] for row in recon.get("conflicts", [])]
    assert "decimal-sum" not in types
