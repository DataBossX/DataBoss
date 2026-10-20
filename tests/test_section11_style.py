"""Synthetic checks for the Section 11 cell-style tool.

Fixtures are fictional. They do not carry a client index, a legal
description, or a file identifier.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from horizon.section11_style import (
    BLM_HEADERS,
    COUNTY_HEADERS,
    apply_section11_style,
    audit_section11_style,
    cm_to_excel_width,
    detect_kind,
    main,
)


def _county_book() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    labels = (
        "County:",
        "Lands:",
        "Date:",
        "Starting Date:",
        "Date Posted Thru:",
        "Indexed By:",
        "Project:",
    )
    for row, label in enumerate(labels, start=1):
        ws.cell(row, 1, label)
    ws["B1"] = "Example"
    ws["B3"] = datetime(2026, 10, 20)
    ws["B4"] = "Inception"
    for col, header in enumerate(COUNTY_HEADERS, start=1):
        ws.cell(9, col, header)
    ws["A10"] = "Oil and Gas Lease"
    ws["B10"] = "Ada Example"
    ws["C10"] = "Example Minerals, LLC"
    ws["D10"] = 1001
    ws["E10"] = "0001-0002"
    ws["F10"] = datetime(2020, 1, 2)
    ws["G10"] = datetime(2020, 1, 3)
    ws["H10"] = "All"
    ws["I10"] = None
    ws["A11"] = "Release"
    ws["I11"] = "missing image"
    return wb


def _blm_book() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    ws["A1"] = "County:"
    ws["B1"] = "Example"
    for col, header in enumerate(BLM_HEADERS, start=1):
        ws.cell(9, col, header)
    ws["A10"] = "Decision"
    ws["B10"] = "Bureau of Land Management"
    ws["C10"] = "Example Operator, LLC"
    ws["D10"] = None
    ws["E10"] = 4
    ws["F10"] = datetime(2024, 3, 1)
    ws["H10"] = None
    ws["I10"] = "All"
    ws["J10"] = 1
    return wb


def test_detects_county_blm_and_checklist():
    county = _county_book().active
    assert detect_kind(county) == "county"
    blm = _blm_book().active
    assert detect_kind(blm) == "blm"
    book = Workbook()
    sheet = book.active
    sheet.title = "Abstract checklist"
    sheet["A2"] = "County Index"
    assert detect_kind(sheet) == "checklist"


def test_apply_keeps_values_and_passes_audit():
    wb = _county_book()
    ws = wb.active
    before = [[ws.cell(r, c).value for c in range(1, 10)] for r in range(1, 12)]
    assert apply_section11_style(ws) == "county"
    after = [[ws.cell(r, c).value for c in range(1, 10)] for r in range(1, 12)]
    assert after == before
    assert ws["I10"].value is None
    assert ws["D10"].number_format == "0"
    assert ws["F10"].number_format == "m/d/yyyy"
    assert audit_section11_style(ws) == []


def test_blm_blank_file_number_is_styled_and_left_blank():
    ws = _blm_book().active
    apply_section11_style(ws, "blm")
    assert ws["D10"].value is None
    assert ws["D10"].alignment.horizontal == "center"
    assert ws["D10"].alignment.shrink_to_fit is True
    assert ws["D10"].border.left.style == "thin"
    assert ws["H10"].value is None
    assert ws["H10"].alignment.horizontal == "left"
    assert audit_section11_style(ws, "blm") == []


def test_messy_sheet_fails_audit_until_styled(tmp_path: Path):
    wb = _county_book()
    ws = wb.active
    ws["A10"].font = Font(name="Arial", size=14, color="FFFF0000")
    findings = audit_section11_style(ws, "county")
    assert findings
    assert any("A10" in str(item) for item in findings)
    apply_section11_style(ws, "county")
    path = tmp_path / "styled.xlsx"
    wb.save(path)
    reloaded = load_workbook(path)
    assert audit_section11_style(reloaded.active, "county") == []
    assert abs(reloaded.active.column_dimensions["A"].width - cm_to_excel_width(8.22854166666667)) < 0.02


def test_checklist_bands_match_section11():
    wb = Workbook()
    ws = wb.active
    ws.title = "Abstract checklist"
    ws["A2"] = "County Index"
    ws["A3"] = "Yes"
    ws["B9"] = "A narrative answer that is long enough to sit on the left"
    ws["E9"] = "SERIAL-1: 2\nSERIAL-2: 3"
    ws["G9"] = datetime(2026, 10, 20)
    apply_section11_style(ws, "checklist")
    assert ws["A2"].font.bold is True
    assert ws["A2"].alignment.vertical == "center"
    assert ws["B9"].alignment.horizontal == "left"
    assert ws["E9"].alignment.horizontal == "center"
    assert ws["G9"].number_format == "mm-dd-yy"
    assert ws["A2"].value == "County Index"
    assert audit_section11_style(ws, "checklist") == []


def test_cli_apply_and_audit(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    dest = tmp_path / "out.xlsx"
    _county_book().save(source)
    assert main(["apply", str(source), str(dest)]) == 0
    assert dest.exists()
    assert main(["audit", str(dest)]) == 0
    assert main(["audit", str(source)]) == 1
