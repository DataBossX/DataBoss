from pathlib import Path

from tools.penterra.image_census import account_path, account_paths
from tools.penterra.section_package import (
    COUNTY_HEADERS,
    FEDERAL_HEADERS,
    WORK_DATE,
    cell_text,
    is_ocr_fragment,
    purify_county_record,
    write_checklist,
    write_letter_xlsx,
    write_ods,
    zip_members,
)


def test_cell_text_formats_iso_datetime_like_values():
    assert cell_text("2026-09-22 00:00:00") == "9/22/2026"


def test_ocr_fragments_are_blanked_and_real_parties_remain():
    fragment = purify_county_record(
        72,
        ["Affidavit", "whether now owned or hereafter acquired by operation", "John Smith", "", "", "", "", "", ""],
        frozenset({72}),
    )
    kept = purify_county_record(
        10,
        ["Warranty Deed", "Fred Christensen", "John F. Christensen", "177712", "0019-0455", "", "", "W/2 W/2", ""],
        frozenset({72}),
    )
    assert fragment[1] == ""
    assert kept[1] == "Fred Christensen"
    assert is_ocr_fragment("agrees to warrant and defend title to the assigned rights of way")


def test_ods_header_contract_and_empty_federal_rows(tmp_path: Path):
    county = tmp_path / "county.ods"
    federal = tmp_path / "federal.ods"
    write_ods(
        county,
        {"lands": "T45N-R76W Section 12: All", "posted_thru": ""},
        COUNTY_HEADERS,
        [["Warranty Deed", "A", "B", "1", "", "", "", "Section 12", ""]],
    )
    write_ods(
        federal,
        {"lands": "WYW-051704", "posted_thru": ""},
        FEDERAL_HEADERS,
        [],
    )
    import zipfile

    with zipfile.ZipFile(county) as archive:
        content = archive.read("content.xml").decode("utf-8")
        styles = archive.read("styles.xml").decode("utf-8")
    assert "T45N-R76W Section 12: All" in content
    assert WORK_DATE in content
    assert "Warranty Deed" in content
    assert 'style:print-orientation="portrait"' in styles
    assert 'style:scale-to="100%"' in styles
    with zipfile.ZipFile(federal) as archive:
        federal_xml = archive.read("content.xml").decode("utf-8")
    assert "WYW-051704" in federal_xml
    assert federal_xml.count("<table:table-row>") == 9


def test_checklist_print_area_and_letter_titles(tmp_path: Path):
    checklist = tmp_path / "checklist.xlsx"
    letter = tmp_path / "letter.xlsx"
    write_checklist(
        checklist,
        {
            "row2": ["County Index"],
            "row3": ["Yes"],
            "row5": ["Abstract Excel Sheet"],
            "row6": ["Yes"],
            "row8": ["Certification Letter"],
            "row9": ["Yes"],
        },
    )
    write_letter_xlsx(
        letter,
        {"lands": "T45N-R76W Section 14: All"},
        COUNTY_HEADERS,
        [["Quit-Claim Deed", "A", "B", "90438", "", "", "", "E/2 SE/4", ""]],
    )
    from openpyxl import load_workbook

    check = load_workbook(checklist)
    sheet = check.active
    assert sheet.title == "Abstract checklist"
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.print_area == "'Abstract checklist'!$A$1:$K$9"
    isolated = load_workbook(letter)
    index = isolated.active
    assert index.page_setup.orientation == "landscape"
    assert index.print_title_rows == "$1:$8"
    assert index.cell(3, 2).value == WORK_DATE


def test_exact_n_zip_membership(tmp_path: Path):
    members = {}
    for name in (
        "county.ods",
        "fed1.ods",
        "fed2.ods",
        "checklist.xlsx",
        "cert.docx",
        "cert.pdf",
    ):
        path = tmp_path / name
        path.write_text("x")
        members[name] = path
    zip_path = tmp_path / "P12_45N-76W-12__SUPERSEDING_TURNIN_DATED_20260922_EXACT6.zip"
    zip_members(zip_path, members)
    import zipfile

    with zipfile.ZipFile(zip_path) as archive:
        assert len(archive.namelist()) == 6


def test_image_census_counts_pdf_pages_not_containers(tmp_path: Path):
    from reportlab.pdfgen import canvas

    pdf = tmp_path / "two-page.pdf"
    other = tmp_path / "notes.txt"
    writer = canvas.Canvas(str(pdf))
    writer.drawString(72, 720, "page 1")
    writer.showPage()
    writer.drawString(72, 720, "page 2")
    writer.save()
    other.write_text("not an image")
    pdf_rows = account_path(pdf)
    assert len(pdf_rows) == 2
    assert all(row["kind"] == "pdf_page" for row in pdf_rows)
    unlabeled = tmp_path / "book583-original"
    unlabeled.write_bytes(pdf.read_bytes())
    packet = account_paths([pdf, other, unlabeled])
    assert packet["image_count"] == 4
    assert packet["pdf_container_extra_images"] == 0
    assert packet["file_count"] == 3
