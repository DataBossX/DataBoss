"""Build dated Combine-Harvest index packages from a selected master JSON.

The master JSON stays outside the public repository. Pass its path explicitly.
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from format_indexes import (
    PREP_DATE,
    audit_workbook,
    build_database,
    client_metadata,
    load_master,
    write_ods,
    write_xlsx,
)

TOWNSHIP = {
    6: ("47N-76W", "T47N-R76W Section 6"),
    7: ("47N-76W", "T47N-R76W Section 7"),
    13: ("47N-77W", "T47N-R77W Section 13"),
    18: ("47N-76W", "T47N-R76W Section 18"),
    19: ("47N-76W", "T47N-R76W Section 19"),
    24: ("47N-77W", "T47N-R77W Section 24"),
    30: ("47N-76W", "T47N-R76W Section 30"),
    31: ("47N-76W", "T47N-R76W Section 31"),
}


def section_folder_name(section: int) -> str:
    township, _ = TOWNSHIP[section]
    return f"Section {section} package {township}"


def write_letter(path: Path, section: int, federal_files: list[str], county_rows: int) -> None:
    township, lands = TOWNSHIP[section]
    township_no, range_no = township.split("-")
    serials = ", ".join(federal_files) if federal_files else ""
    body = f"""Township {township_no.replace('N', '')} North, Range {range_no.replace('W', '')} West, 6th P.M.
Section {section}: as indexed
Johnson County, Wyoming

Contents of abstract:

Index of documents reflecting entries posted in the tract indices of the Johnson County Clerk, limited to the {county_rows} county applications in the accompanying index. No official search-through date is asserted.

{("Images and indexes of federal case files " + serials + ". Page numbers are the first physical page of each indexed occurrence. No official federal search-through date is asserted.") if serials else "No federal case workbook is selected for this section. That is not proof that no federal lease affects the section."}

Bureau of Land Management serial-register pages are accounted separately from operative index rows. A serial-register page is not a substitute for the case-file document.

Judgment, lien, mortgage, and civil-suit search not performed.

This package was prepared on 10/19/2026 solely as qualified review work product for Combine-Harvest CH185443. It is not a title opinion, not a guaranty of title, and not a certified abstract. It is limited to the materials reviewed and listed. The instruments have not been examined for legal sufficiency. This letter is unsigned.

Prepared 10/19/2026 by: Ryan Gille

_________________________________
Gille Energy, LLC
"""
    # A minimal Word document is enough for the package member.
    document = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <w:body><w:p><w:r><w:t xml:space="preserve">{_xml(body)}</w:t></w:r></w:p></w:body>
</w:document>"""
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
 <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
 <Default Extension="xml" ContentType="application/xml"/>
 <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)


def _xml(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def write_checklist(path: Path, section: int, county_rows: int, federal_files: list[str]) -> None:
    township, lands = TOWNSHIP[section]
    book = Workbook()
    sheet = book.active
    sheet.title = "Checklist"
    rows = [
        ("County", "Johnson"),
        ("Lands Covered", lands),
        ("Date completed", PREP_DATE),
        ("Starting Date", None),
        ("Date Certified/Verified", None),
        ("Landman Name", "Ryan Gille"),
        ("Project/Prospect", "Combine-Harvest CH185443 / AFE210275"),
        ("County Index", f"Formatted; {county_rows} selected applications"),
        ("County Documents", "Not certified complete against every handwritten appearance"),
        ("Plat", "N/A - no separate plat member in this package"),
        ("STR online search", "Reference route only; no authenticated cutoff"),
        ("Compare index to Uploads", "Open except where a named visual census receipt exists"),
        ("BLM File", ", ".join(federal_files) if federal_files else "No selected federal workbook"),
        ("SRP", "Administrative; excluded from operative rows"),
        ("State File", "Not included"),
        ("Lease Report if applicable", "Not asserted"),
        ("Reviewed SRP vs BLM file or State file", "Open where the case census is open"),
        ("Certification Letter", "Unsigned; no search-through date"),
        ("Status", "QUALIFIED REVIEW"),
    ]
    bold = Font(name="Calibri", size=11, bold=True)
    body = Font(name="Calibri", size=11)
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(index, 1, label).font = bold
        cell = sheet.cell(index, 2, value)
        cell.font = body
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if hasattr(value, "year"):
            cell.number_format = "m/d/yyyy"
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 88
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def status_markdown(master: dict, audits: list[dict]) -> str:
    by_section = defaultdict(list)
    for audit in audits:
        by_section[audit["section"]].append(audit)
    lines = [
        "# CH185443 Combine-Harvest status report",
        "",
        "Prepared 10/19/2026. Qualified review. This is not a certified abstract, title opinion, or guaranty of title.",
        "",
        "The selected indexes contain the applications below. A row is a section application. It is not, by itself, proof that every physical document in a case file has been listed.",
        "",
        "## Files filed against each section",
        "",
        "| Section | Township | County files | BLM files | BLM workbooks | Blank cells | Unknown cells | Qualified unknown | Unresolved |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    totals = Counterish()
    for item in master["summary"]:
        section = item["section"]
        township = TOWNSHIP[section][0]
        lines.append(
            f"| {section} | {township} | {item['county_rows']} | {item['federal_rows']} | {item['federal_workbooks']} | {item['substantive_display_blanks']} | {item['substantive_display_unknown']} | {item['substantive_display_qualified_unknown']} | {item['unresolved_evidence_cells']} |"
        )
        for key in (
            "county_rows",
            "federal_rows",
            "federal_workbooks",
            "substantive_display_blanks",
            "substantive_display_unknown",
            "substantive_display_qualified_unknown",
            "unresolved_evidence_cells",
        ):
            totals.add(key, item[key])
    lines.append(
        f"| TOTAL |  | {totals['county_rows']} | {totals['federal_rows']} | {totals['federal_workbooks']} | {totals['substantive_display_blanks']} | {totals['substantive_display_unknown']} | {totals['substantive_display_qualified_unknown']} | {totals['unresolved_evidence_cells']} |"
    )
    lines += [
        "",
        "Blank and unknown counts are substantive cells only. County Comments and BLM File No and Comments are excluded. Literal Unknown cells are the Unknown column. Qualified unknown cells are unresolved populated cells carried from the master summary. Unresolved is the evidence-gap total and is not reduced by clearing a header.",
        "",
        "## BLM files",
        "",
        "Documents are the operative index rows in that section workbook. Parts are the distinct numeric Part values on those rows. A workbook with no rows is retained. An unnumbered source is not given Part 1.",
        "",
        "| Section | BLM file | Documents | Parts | Part numbers | Blank cells | Unknown cells |",
        "| --- | --- | ---: | ---: | --- | ---: | ---: |",
    ]
    for audit in sorted(audits, key=lambda item: (item["section"], item["file"])):
        if audit["kind"] == "county":
            continue
        part_text = ", ".join(str(part) for part in audit["parts"]) if audit["parts"] else ""
        lines.append(
            f"| {audit['section']} | {audit['file']} | {audit['documents']} | {audit['part_count']} | {part_text} | {audit['blank_cells']} | {audit['unknown_cells']} |"
        )
    lines += [
        "",
        "## Display applied on 10/19/2026",
        "",
        "- Index Date is 10/19/2026. Instrument dates use m/d/yyyy.",
        "- Starting Date and Date Posted Thru are blank because no proved search period was supplied. The prior header text Available records and Unknown was not a search cutoff.",
        "- BLM File No and Comments stay as columns and are blank.",
        "- Page is the starting physical page, one number. Part is one number in the last column.",
        "- Parties and legal descriptions that already use a comma and a new line in the same cell are preserved.",
        "- Section legals that reach other lands keep , aol. Case-wide rows keep every proved leg in 47N-76W and 47N-77W.",
        "",
        "## What this pass does not close",
        "",
        "Section 6 now has a 7-page visual appearance census: 128 appearances, 117 locator groups, 36 candidates not in the selected or held snapshots, and 1 selected row without a matching appearance. Those 36 candidates were not added to the index.",
        "",
        "WYW153681 Part 2 now has a 79-page visual census: 29 source occurrences, of which 25 are substantive documents, plus 1 administrative document, 2 postal records, and 1 jacket. The other federal parts named in the 10/19/2026 master notes still need a full boundary census.",
        "",
        *_part2_import_lines(master),
        "Codex, Claude, and the Grok bot were not used. This package was built from the selected master rows and the two census receipts already on Drive.",
        "",
        "## Source notes that stay open",
        "",
        "WYW-145618 section 24 has five index rows. Four of them are stored as Part 1 on pages 164, 166, 170, and 188. The located WYW145618 PDF is Part 2 of 2 and has 35 pages. Those page numbers do not fall inside that PDF, so the rows were not renumbered.",
        "",
        "W-40637 and WYW-094604 section 13 each have one index row and a blank Part. They were left blank rather than called Part 1.",
        "",
        "WYW-094604 section 24 is the retained header-only workbook: 0 documents.",
        "",
        "### WYW153681 Part 2 documents, from the 79-page visual census",
        "",
        "29 occurrences: 25 substantive documents, 1 administrative checklist, 2 postal records, and 1 jacket. Jacket and administrative and postal rows stay out of the operative index.",
        "",
        "| ID | Pages | Category | Document |",
        "| --- | --- | --- | --- |",
        "| J001 | 1 | Jacket | Case-file jacket |",
        "| D002 | 2-5 | Substantive | Transfer of Operating Rights (Sublease) |",
        "| A006 | 6-7 | Administrative | Federal suspension/debarment eligibility checklist |",
        "| D008 | 8-10 | Substantive | Notice: Lease Accounts Transferred to Non-Terminable Status |",
        "| D011 | 11-14 | Substantive | Approval of Initial Niobrara Formation Participating Area A, Tractor (Deep) Unit |",
        "| D015 | 15-16 | Substantive | Notice — Leases Placed in Indefinite Term |",
        "| D017 | 17-18 | Substantive | Memorandum — First Production and Qualifying Well Determination |",
        "| D019 | 19-22 | Substantive | Transfer of Operating Rights (Sublease) |",
        "| D023 | 23-24 | Substantive | Memorandum — Approval of Tractor (Deep) Unit Agreement |",
        "| D025 | 25-26 | Substantive | Notice — Name Change Recognized |",
        "| D027 | 27 | Substantive | Certificate of Name Change |",
        "| D028 | 28-35 | Substantive | Transfer of Operating Rights (Sublease) |",
        "| D036 | 36-38 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D039 | 39-41 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D042 | 42-44 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D045 | 45-47 | Substantive | Assignment of Record Title Interest |",
        "| D048 | 48-50 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D051 | 51-53 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D054 | 54-56 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D057 | 57-59 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D060 | 60-62 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| D063 | 63-65 | Substantive | Transfer of Operating Rights (Sublease) and Overriding Royalty |",
        "| P066 | 66 | Postal | FedEx delivery confirmation email |",
        "| P067 | 67 | Postal | FedEx shipment receipt |",
        "| D068 | 68-69 | Substantive | BLM Notice — Mass Transfer Returned for Correction |",
        "| D070 | 70-72 | Substantive | Transfer of Operating Rights and Overriding Royalty, returned |",
        "| D073 | 73-75 | Substantive | Assignment of Record Title Interest |",
        "| D076 | 76 | Substantive | Authorization Letter — Remove Depth Wording |",
        "| D077 | 77-79 | Substantive | Transfer of Operating Rights (Sublease) |",
        "",
        "### Section 6 handwritten index",
        "",
        "All 7 physical pages of S06_Tract_Index.pdf were visually reviewed in the census already on Drive. 128 appearances reconcile to 117 locator groups: 69 selected groups, 11 held groups, 36 candidates not in the selected or held snapshots, and 1 cancelled entry. The 36 candidates were not added.",
        "",
    ]
    return "\n".join(lines) + "\n"


def _part2_import_lines(master: dict) -> list[str]:
    note = master.get("wyw153681_part2_import")
    if not note:
        return []
    held = ", ".join(note.get("not_placed_on_a_project_section") or []) or "none"
    return [
        f"The Part 2 census was applied to the eight section packages: {note.get('added_count', 0)} new section applications and {note.get('filled_count', 0)} received dates filled on rows that already had the same Part and starting Page. A document is added only when its reviewed legal names that project section. Depth and formation stay out of the client legal. Documents with no listed project-section lands stay in the census: {held}.",
        "",
    ]


class Counterish(dict):
    def add(self, key, value):
        self[key] = self.get(key, 0) + value

    def __getitem__(self, key):
        return dict.get(self, key, 0)


def build(master_path: Path, output: Path) -> Path:
    master = load_master(master_path)
    output.mkdir(parents=True, exist_ok=True)
    audits = []
    grouped = defaultdict(list)
    for workbook in master["workbooks"]:
        grouped[workbook["section"]].append(workbook)
        audits.append(audit_workbook(workbook))
    for section, books in grouped.items():
        folder = output / section_folder_name(section)
        folder.mkdir(parents=True, exist_ok=True)
        federal = [book["file"].split("_")[0] for book in books if book["kind"] != "county"]
        county_rows = next(book["row_count"] for book in books if book["kind"] == "county")
        township = TOWNSHIP[section][0]
        write_letter(folder / f"{section}-{township}_Certification_Letter.docx", section, federal, county_rows)
        write_checklist(
            folder / f"Abstract_Checklist_{section}-{township}.xlsx",
            section,
            county_rows,
            federal,
        )
        for book in books:
            stem = Path(book["file"]).stem
            write_xlsx(folder / f"{stem}.xlsx", book)
            write_ods(folder / f"{stem}.ods", book)
    build_database(output / "CH185443_inventory_20261019.sqlite", master, audits)
    report = status_markdown(master, audits)
    (output / "CH185443_Status_Report_20261019.md").write_text(report)
    (output / "audits.json").write_text(json.dumps(audits, indent=2))
    archive_path = output / "CH185443_Formatted_Section_Packages_20261019.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in output.rglob("*"):
            if path.is_file() and path != archive_path and path.suffix != ".zip":
                archive.write(path, path.relative_to(output).as_posix())
    return archive_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("master_json", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(build(args.master_json, args.output))


if __name__ == "__main__":
    main()
