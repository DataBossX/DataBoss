"""Format Combine-Harvest abstract indexes to the Overtime package layout.

This module does not fetch source images and does not invent instrument facts.
It applies display rules to an already selected row file:

- preparation date 10/19/2026, typed as m/d/yyyy
- unproved Starting Date and Date Posted Thru left blank
- Excel date serials shown as m/d/yyyy
- BLM File No and Comments columns kept and left blank
- Page and Part stored as integers
- parties and legal lines kept as comma-plus-newline cells
"""

from __future__ import annotations

import json
import sqlite3
import zipfile
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXCEL_EPOCH = datetime(1899, 12, 30)
PREP_DATE = datetime(2026, 10, 19)
DATE_FORMAT = "m/d/yyyy"
FONT_NAME = "Calibri"
FONT_SIZE = 11

COUNTY_WIDTHS = {
    "Document Type": 42,
    "Grantor": 32,
    "Grantee": 32,
    "Doc No": 16,
    "Book-Page": 16,
    "Date of Doc": 14,
    "Rec Date": 14,
    "Legal Description": 46,
    "Comments": 18,
}
FEDERAL_WIDTHS = {
    "Document Type": 46,
    "Grantor": 32,
    "Grantee": 32,
    "File No": 12,
    "Page": 10,
    "Date of Doc": 14,
    "Rec Date": 14,
    "Comments": 14,
    "Legal Description": 46,
    "Part": 10,
}

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HEADER_FILL = PatternFill("solid", fgColor="E7EEF5")
LABEL_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
BODY_FONT = Font(name=FONT_NAME, size=FONT_SIZE)
HEADER_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def excel_serial_to_date(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not (1 < value < 80000):
        return None
    return EXCEL_EPOCH + timedelta(days=int(value))


def format_date_text(value):
    found = excel_serial_to_date(value)
    if found is None:
        return None
    return f"{found.month}/{found.day}/{found.year}"


def is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def is_literal_unknown(value) -> bool:
    return isinstance(value, str) and value.strip().lower() == "unknown"


def substantive_headers(kind: str, headers: list[str]) -> set[str]:
    if kind == "county":
        return {h for h in headers if h != "Comments"}
    return {h for h in headers if h not in {"File No", "Comments"}}


def display_value(kind: str, header: str, value):
    if kind != "county" and header in {"File No", "Comments"}:
        return None
    if header in {"Date of Doc", "Rec Date"}:
        return excel_serial_to_date(value) if excel_serial_to_date(value) else (None if is_blank(value) else value)
    if header in {"Page", "Part"} and isinstance(value, float) and value.is_integer():
        return int(value)
    if is_blank(value):
        return None
    return value


def sort_key(kind: str, headers: list[str], row: dict, order: int):
    if kind == "county":
        return (0, row.get("new_row") or row.get("original_row") or order, order)
    values = row["values"]
    part = values[headers.index("Part")] if "Part" in headers else None
    page = values[headers.index("Page")] if "Page" in headers else None
    part_key = part if isinstance(part, (int, float)) and not isinstance(part, bool) else 10**9
    page_key = page if isinstance(page, (int, float)) and not isinstance(page, bool) else 10**9
    return (1, part_key, page_key, order)


def metadata_map(workbook: dict) -> dict[str, str]:
    found = {}
    for row in workbook.get("metadata") or []:
        if row and row[0]:
            found[str(row[0]).rstrip(":")] = "" if len(row) < 2 or row[1] is None else str(row[1])
    return found


def client_metadata(workbook: dict) -> list[tuple[str, object]]:
    raw = metadata_map(workbook)
    return [
        ("County:", raw.get("County") or "Johnson"),
        ("Lands:", raw.get("Lands") or ""),
        ("Date:", PREP_DATE),
        ("Starting Date:", None),
        ("Date Posted Thru:", None),
        ("Indexed By:", "Ryan Gille"),
        ("Project:", raw.get("Project") or "Combine-Harvest CH185443 / AFE210275"),
    ]


def line_count(value) -> int:
    if value is None:
        return 1
    if isinstance(value, datetime):
        return 1
    return max(1, str(value).count("\n") + 1)


def write_xlsx(path: Path, workbook: dict) -> None:
    headers = workbook["headers"]
    kind = workbook["kind"]
    widths = COUNTY_WIDTHS if kind == "county" else FEDERAL_WIDTHS
    book = Workbook()
    sheet = book.active
    sheet.title = "Index"
    label_font = LABEL_FONT
    for index, (label, value) in enumerate(client_metadata(workbook), start=1):
        sheet.cell(index, 1, label).font = label_font
        cell = sheet.cell(index, 2, value)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        if isinstance(value, datetime):
            cell.number_format = DATE_FORMAT
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(9, column, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
        sheet.column_dimensions[get_column_letter(column)].width = widths.get(header, 18)
    rows = sorted(
        enumerate(workbook["rows"]),
        key=lambda item: sort_key(kind, headers, item[1], item[0]),
    )
    for offset, (_, row) in enumerate(rows):
        excel_row = 10 + offset
        height = 18
        for column, header in enumerate(headers, start=1):
            raw = row["values"][column - 1] if column - 1 < len(row["values"]) else None
            value = display_value(kind, header, raw)
            cell = sheet.cell(excel_row, column, value)
            cell.font = BODY_FONT
            cell.alignment = CENTER if header in {"Page", "Part", "Doc No", "Book-Page"} else WRAP
            cell.border = THIN
            if isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
            if header in {"Page", "Part"} and isinstance(value, int):
                cell.number_format = "0"
            height = max(height, 15 * line_count(value) + 4)
        sheet.row_dimensions[excel_row].height = min(height, 120)
    sheet.freeze_panes = "A10"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = metadata_map(workbook).get("Lands", "")
    sheet.oddFooter.left.text = "Prepared 10/19/2026"
    sheet.print_title_rows = "9:9"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _ods_cell(value) -> str:
    if value is None or value == "":
        return "<table:table-cell/>"
    if isinstance(value, datetime):
        text = f"{value.month}/{value.day}/{value.year}"
        return (
            '<table:table-cell office:value-type="date" office:date-value="'
            + value.strftime("%Y-%m-%d")
            + '"><text:p>'
            + escape(text)
            + "</text:p></table:table-cell>"
        )
    if isinstance(value, int) and not isinstance(value, bool):
        return (
            f'<table:table-cell office:value-type="float" office:value="{value}">'
            f"<text:p>{value}</text:p></table:table-cell>"
        )
    lines = str(value).split("\n")
    paragraphs = "".join(f"<text:p>{escape(line)}</text:p>" for line in lines)
    return f'<table:table-cell office:value-type="string">{paragraphs}</table:table-cell>'


def write_ods(path: Path, workbook: dict) -> None:
    headers = workbook["headers"]
    kind = workbook["kind"]
    rows_xml = []
    for label, value in client_metadata(workbook):
        rows_xml.append(
            "<table:table-row>" + _ods_cell(label) + _ods_cell(value) + "</table:table-row>"
        )
    rows_xml.append("<table:table-row/>")
    rows_xml.append("<table:table-row>" + "".join(_ods_cell(h) for h in headers) + "</table:table-row>")
    ordered = sorted(
        enumerate(workbook["rows"]),
        key=lambda item: sort_key(kind, headers, item[1], item[0]),
    )
    for _, row in ordered:
        cells = []
        for column, header in enumerate(headers):
            raw = row["values"][column] if column < len(row["values"]) else None
            cells.append(_ods_cell(display_value(kind, header, raw)))
        rows_xml.append("<table:table-row>" + "".join(cells) + "</table:table-row>")
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 office:version="1.3">
 <office:automatic-styles>
  <style:style style:name="ce" style:family="table-cell">
   <style:text-properties style:font-name="Calibri" fo:font-size="11pt"/>
  </style:style>
 </office:automatic-styles>
 <office:body><office:spreadsheet>
  <table:table table:name="Index">{''.join(rows_xml)}</table:table>
 </office:spreadsheet></office:body>
</office:document-content>"""
    styles = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 office:version="1.3">
 <office:styles>
  <style:default-style style:family="table-cell">
   <style:text-properties style:font-name="Calibri" fo:font-size="11pt"/>
  </style:default-style>
 </office:styles>
</office:document-styles>"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet", compress_type=zipfile.ZIP_STORED)
        archive.writestr("content.xml", content)
        archive.writestr("styles.xml", styles)
        archive.writestr("META-INF/manifest.xml", manifest)


def audit_workbook(workbook: dict) -> dict:
    headers = workbook["headers"]
    kind = workbook["kind"]
    substantive = substantive_headers(kind, headers)
    blanks = Counter()
    unknowns = Counter()
    parts = []
    for row in workbook["rows"]:
        for header in headers:
            value = row["values"][headers.index(header)] if header in headers else None
            shown = display_value(kind, header, value)
            if header in substantive and is_blank(shown):
                blanks[header] += 1
            if header in substantive and is_literal_unknown(value):
                unknowns[header] += 1
        if kind != "county" and "Part" in headers:
            part = row["values"][headers.index("Part")]
            if isinstance(part, (int, float)) and not isinstance(part, bool):
                parts.append(int(part))
    return {
        "section": workbook["section"],
        "kind": kind,
        "file": workbook["file"],
        "documents": len(workbook["rows"]),
        "parts": sorted(set(parts)),
        "part_count": len(set(parts)),
        "blank_cells": sum(blanks.values()),
        "blank_by_column": dict(blanks),
        "unknown_cells": sum(unknowns.values()),
        "unknown_by_column": dict(unknowns),
    }


def build_database(path: Path, master: dict, audits: list[dict]) -> None:
    if path.exists():
        path.unlink()
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE section_summary (
            section INTEGER PRIMARY KEY,
            county_rows INTEGER,
            federal_rows INTEGER,
            federal_workbooks INTEGER,
            blank_cells INTEGER,
            unknown_cells INTEGER,
            qualified_unknown_cells INTEGER,
            unresolved_evidence_cells INTEGER
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE blm_files (
            section INTEGER,
            file TEXT,
            documents INTEGER,
            part_count INTEGER,
            parts TEXT,
            blank_cells INTEGER,
            unknown_cells INTEGER
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE index_rows (
            section INTEGER,
            kind TEXT,
            file TEXT,
            row_id TEXT,
            document_type TEXT,
            grantor TEXT,
            grantee TEXT,
            page TEXT,
            part TEXT,
            legal TEXT,
            blank_substantive INTEGER,
            unknown_substantive INTEGER
        )
        """
    )
    for item in master.get("summary") or []:
        connection.execute(
            """
            INSERT INTO section_summary VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                item["section"],
                item["county_rows"],
                item["federal_rows"],
                item["federal_workbooks"],
                item["substantive_display_blanks"],
                item["substantive_display_unknown"],
                item["substantive_display_qualified_unknown"],
                item["unresolved_evidence_cells"],
            ),
        )
    for audit in audits:
        if audit["kind"] == "county":
            continue
        connection.execute(
            "INSERT INTO blm_files VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                audit["section"],
                audit["file"],
                audit["documents"],
                audit["part_count"],
                ", ".join(str(part) for part in audit["parts"]),
                audit["blank_cells"],
                audit["unknown_cells"],
            ),
        )
    for workbook in master["workbooks"]:
        headers = workbook["headers"]
        kind = workbook["kind"]
        substantive = substantive_headers(kind, headers)

        def col(name, values):
            if name not in headers:
                return ""
            value = values[headers.index(name)]
            shown = display_value(kind, name, value)
            if shown is None:
                return ""
            if isinstance(shown, datetime):
                return format_date_text(value) or ""
            return str(shown)

        for row in workbook["rows"]:
            values = row["values"]
            blanks = 0
            unknowns = 0
            for header in headers:
                raw = values[headers.index(header)]
                shown = display_value(kind, header, raw)
                if header in substantive and is_blank(shown):
                    blanks += 1
                if header in substantive and is_literal_unknown(raw):
                    unknowns += 1
            connection.execute(
                "INSERT INTO index_rows VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    workbook["section"],
                    kind,
                    workbook["file"],
                    row.get("row_id") or "",
                    col("Document Type", values),
                    col("Grantor", values),
                    col("Grantee", values),
                    col("Page", values),
                    col("Part", values),
                    col("Legal Description", values),
                    blanks,
                    unknowns,
                ),
            )
    connection.commit()
    connection.close()


def load_master(path: Path) -> dict:
    return json.loads(path.read_text())
