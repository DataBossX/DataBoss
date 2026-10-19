import json
import sqlite3
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from build_package import build
from format_indexes import excel_serial_to_date, format_date_text, numeric_document_number
from openpyxl import load_workbook


class FormatTests(unittest.TestCase):
    def test_known_recording_date(self):
        self.assertEqual(format_date_text(24558), "3/27/1967")
        self.assertEqual(excel_serial_to_date(24524), datetime(1967, 2, 21))
        self.assertEqual(numeric_document_number("084033"), 84033)
        self.assertIsNone(numeric_document_number("Unknown"))

    def test_package_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            tmp_path = Path(directory)
            master = {
                "summary": [
                    {
                        "section": 24,
                        "county_rows": 1,
                        "federal_rows": 2,
                        "federal_workbooks": 1,
                        "substantive_display_blanks": 1,
                        "substantive_display_unknown": 1,
                        "substantive_display_qualified_unknown": 0,
                        "unresolved_evidence_cells": 2,
                    }
                ],
                "parent_drive_id": "parent",
                "parent_sha256": "abc",
                "workbooks": [
                    {
                        "section": 24,
                        "kind": "county",
                        "file": "47N-77W-24_Johnson_Co_Penterra_Abstract_Index.xlsx",
                        "metadata": [
                            ["County:", "Johnson"],
                            ["Lands:", "T47N-R77W Section 24: All"],
                            ["Date:", "2026-10-19T00:00:00.000Z"],
                            ["Starting Date:", "Available records"],
                            ["Date Posted Thru:", "Unknown"],
                            ["Indexed By:", "Ryan Gille"],
                            ["Project:", "Combine-Harvest CH185443 / AFE210275"],
                        ],
                        "headers": [
                            "Document Type", "Grantor", "Grantee", "Doc No", "Book-Page",
                            "Date of Doc", "Rec Date", "Legal Description", "Comments",
                        ],
                        "rows": [
                            {
                                "row_id": "c1",
                                "values": [
                                    "Deed",
                                    "Example Energy, LLC,\nSample Resources Corporation",
                                    "Unknown",
                                    "084033",
                                    "1-1",
                                    24524,
                                    24558,
                                    "All of 24-47N-77W, aol",
                                    "missing image",
                                ],
                            }
                        ],
                        "row_count": 1,
                    },
                    {
                        "section": 24,
                        "kind": "federal",
                        "file": "WYW-145618_Johnson_Co_Penterra_Abstract_Index.xlsx",
                        "metadata": [
                            ["County:", "Johnson"],
                            ["Lands:", "WYW-145618; T47N-R77W Section 24"],
                            ["Starting Date:", "Available records"],
                            ["Date Posted Thru:", "Unknown"],
                        ],
                        "headers": [
                            "Document Type", "Grantor", "Grantee", "File No", "Page",
                            "Date of Doc", "Rec Date", "Comments", "Legal Description", "Part",
                        ],
                        "rows": [
                            {
                                "row_id": "f2",
                                "values": ["Transfer", "Alpha, LLC", "Beta, LLC", "SHOULD-CLEAR", 13, None, None, "note", "outside lands only", 10],
                            },
                            {
                                "row_id": "f1",
                                "values": [
                                    "Transfer",
                                    "Alpha, LLC,\nBeta, LLC",
                                    "Gamma Inc.",
                                    "",
                                    2,
                                    24558,
                                    None,
                                    "",
                                    "Lots 2, 4 of 24-47N-77W,\nLots 1, 2 of 19-47N-76W, aol",
                                    2,
                                ],
                            },
                        ],
                        "row_count": 2,
                    },
                ],
            }
            master_path = tmp_path / "master.json"
            master_path.write_text(json.dumps(master))
            out = tmp_path / "out"
            build(master_path, out)
            folder = out / "Section 24 package 47N-77W"
            sheet = load_workbook(folder / "WYW-145618_Johnson_Co_Penterra_Abstract_Index.xlsx").active
            self.assertEqual(sheet["B3"].value, datetime(2026, 10, 19))
            self.assertEqual(sheet["B3"].number_format, "m/d/yyyy")
            self.assertIsNone(sheet["B4"].value)
            self.assertEqual(sheet["B5"].value, "Unknown")
            self.assertEqual(sheet["J9"].value, "Part")
            self.assertIsNone(sheet["D10"].value)
            self.assertIsNone(sheet["H10"].value)
            self.assertEqual(sheet["E10"].value, 2)
            self.assertEqual(sheet["J10"].value, 2)
            self.assertEqual(sheet["E11"].value, 13)
            self.assertEqual(sheet["J11"].value, 10)
            self.assertIn("\n", sheet["I10"].value)
            county_sheet = load_workbook(folder / "47N-77W-24_Johnson_Co_Penterra_Abstract_Index.xlsx").active
            self.assertEqual(county_sheet["D10"].value, 84033)
            self.assertEqual(county_sheet["I10"].value, "missing image")
            self.assertEqual(county_sheet["F10"].number_format, "m/d/yyyy")
            ods = zipfile.ZipFile(folder / "47N-77W-24_Johnson_Co_Penterra_Abstract_Index.ods")
            text = "".join(ET.fromstring(ods.read("content.xml")).itertext())
            self.assertIn("2/21/1967", text)
            self.assertIn("3/27/1967", text)
            self.assertNotIn("Available records", text)
            connection = sqlite3.connect(out / "CH185443_inventory_20261019.sqlite")
            documents, parts, blanks, unknowns = connection.execute(
                "SELECT documents, parts, blank_cells, unknown_cells FROM blm_files"
            ).fetchone()
            self.assertEqual((documents, parts, unknowns), (2, "2, 10", 0))
            self.assertGreaterEqual(blanks, 1)
            self.assertEqual(
                connection.execute("SELECT doc_no_source_text, doc_no_display FROM workbook_rows WHERE row_id = 'c1'").fetchone(),
                ("084033", "84033"),
            )
            self.assertTrue(connection.execute("SELECT status FROM package_versions").fetchone()[0].startswith("QUALIFIED"))
            archive = out / "Section_24_package_47N-77W_QUALIFIED_UNSIGNED_20261019.zip"
            self.assertIsNone(zipfile.ZipFile(archive).testzip())


if __name__ == "__main__":
    unittest.main()
