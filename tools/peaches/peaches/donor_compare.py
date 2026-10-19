"""Compare a candidate workbook's structure against a donor workbook.

All measurements are parser-level (ODS styles.xml / content.xml via the
standard library; XLSX sheet properties via openpyxl). Parser-level parity
is not a native-application print verification and is never labelled as one.

CLI: python -m peaches.donor_compare <candidate> <donor> <out_json>
"""
from __future__ import annotations

import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

from . import wb

STYLE = wb.NS["style"]
FO = wb.NS["fo"]


def _cell(grid, r, c):
    if 0 <= r < len(grid) and 0 <= c < len(grid[r]):
        return (grid[r][c] or "").strip()
    return ""


def _ods_structure(path):
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        content = ET.fromstring(zf.read("content.xml"))
        styles_root = None
        if "styles.xml" in names:
            styles_root = ET.fromstring(zf.read("styles.xml"))
    roots = [r for r in (styles_root, content) if r is not None]

    def _font_from(style_el):
        tp = style_el.find(wb.q("style", "text-properties"))
        if tp is None:
            return None, None
        return (tp.get(wb.q("style", "font-name")),
                tp.get(wb.q("fo", "font-size")))

    font_name = font_size = None
    for root in roots:  # 1) true default-style for table cells
        for ds in root.iter(wb.q("style", "default-style")):
            if ds.get(wb.q("style", "family")) == "table-cell":
                font_name, font_size = _font_from(ds)
                if font_name:
                    break
        if font_name:
            break
    if not font_name:  # 2) named "Default" cell style (LibreOffice donors)
        for root in roots:
            for st in root.iter(wb.q("style", "style")):
                if st.get(wb.q("style", "name")) == "Default" and \
                        st.get(wb.q("style", "family")) == "table-cell":
                    font_name, font_size = _font_from(st)
                    if font_name:
                        break
            if font_name:
                break
    font_name = font_name or "not-set"
    font_size = font_size or "not-set"

    # Page setup: aggregate across ALL page layouts (multi-sheet workbooks
    # carry several); report the distinct value sets, deterministically.
    orientations, margin_sets, scale_bits = set(), set(), set()
    for root in roots:
        for layout in root.iter(wb.q("style", "page-layout")):
            props = layout.find(wb.q("style", "page-layout-properties"))
            if props is None:
                continue
            o = props.get(wb.q("style", "print-orientation"))
            if o:
                orientations.add(o)
            m = [props.get(wb.q("fo", k)) for k in
                 ("margin-top", "margin-right", "margin-bottom",
                  "margin-left")]
            if any(v is not None for v in m):
                margin_sets.add("/".join(
                    v if v is not None else "?" for v in m))
            for key, val in props.attrib.items():
                if "scale" in key or "fit" in key:
                    scale_bits.add("%s=%s" % (key.split("}")[-1], val))
    orientation = ";".join(sorted(orientations)) or "not-set"
    margins = " | ".join(sorted(margin_sets)) or "not-set"
    scale = ";".join(sorted(scale_bits)) or "not-set"
    return font_name, font_size, orientation, margins, scale


def _xlsx_structure(path):
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_name, _grid = wb.choose_sheet(wb.read_xlsx_sheets(path))
        ws = book[sheet_name] if sheet_name in book.sheetnames \
            else book.worksheets[0]
        font_name, font_size = "not-set", "not-set"
        for cell in ws[9]:
            if cell.value:
                if cell.font is not None:
                    font_name = cell.font.name or font_name
                    font_size = (str(cell.font.size)
                                 if cell.font.size else font_size)
                break
        def _safe(getter):
            try:
                return getter()
            except Exception:
                return None

        ps = ws.page_setup
        orientation = _safe(lambda: ps.orientation) or "not-set"
        margin_vals = [_safe(lambda k=k: getattr(ws.page_margins, k))
                       for k in ("top", "right", "bottom", "left")]
        margins = "/".join(str(v) for v in margin_vals)
        scale_bits = []
        if _safe(lambda: ps.fitToPage):
            scale_bits.append("fitToPage=True")
        if _safe(lambda: ps.fitToWidth):
            scale_bits.append("fitToWidth=%s" % ps.fitToWidth)
        if _safe(lambda: ps.fitToHeight):
            scale_bits.append("fitToHeight=%s" % ps.fitToHeight)
        if _safe(lambda: ps.scale):
            scale_bits.append("scale=%s" % ps.scale)
        scale = ";".join(scale_bits) if scale_bits else "not-set"
        return font_name, font_size, orientation, margins, scale
    finally:
        book.close()


def extract_structure(path):
    _name, grid = wb.read_index_sheet(path)
    headers = [_cell(grid, 8, c) for c in range(10)]
    labels = []
    for r in range(7):
        first = ""
        for c in range(4):
            v = _cell(grid, r, c)
            if v:
                first = v
                break
        labels.append(first)
    row8_blank = all(not _cell(grid, 7, c) for c in range(12))
    first_data = None
    for r in range(9, len(grid)):
        if any(_cell(grid, r, c) for c in range(12)):
            first_data = r + 1
            break
    if str(path).lower().endswith(".ods"):
        font_name, font_size, orientation, margins, scale = _ods_structure(path)
        method = "parser-level: ODS styles.xml/content.xml"
    else:
        font_name, font_size, orientation, margins, scale = _xlsx_structure(path)
        method = "parser-level: openpyxl sheet properties"
    return {
        "file": os.path.basename(path),
        "method": method,
        "schema_headers": headers,
        "metadata_labels_rows1_7": labels,
        "row8_blank": row8_blank,
        "first_data_row": first_data,
        "default_font_name": font_name,
        "default_font_size": font_size,
        "page_orientation": orientation,
        "page_margins_trbl": margins,
        "page_scale": scale,
    }


def compare(candidate_path, donor_path):
    cand = extract_structure(candidate_path)
    donor = extract_structure(donor_path)
    props = ["schema_headers", "metadata_labels_rows1_7", "row8_blank",
             "first_data_row", "default_font_name", "default_font_size",
             "page_orientation", "page_margins_trbl", "page_scale"]
    properties = {}
    for prop in props:
        properties[prop] = {
            "status": "AGREES" if cand[prop] == donor[prop] else "DIFFERS",
            "candidate": cand[prop],
            "donor": donor[prop],
        }
    return {
        "candidate": cand["file"],
        "donor": donor["file"],
        "method": cand["method"] + " | " + donor["method"],
        "note": "Parser-level comparison only; not a native print check.",
        "properties": properties,
    }


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 3:
        print("usage: python -m peaches.donor_compare <candidate> <donor> "
              "<out_json>", file=sys.stderr)
        return 2
    result = compare(argv[0], argv[1])
    with open(argv[2], "w", encoding="utf-8") as fh:
        json.dump(result, fh, indent=2)
    agrees = sum(1 for p in result["properties"].values()
                 if p["status"] == "AGREES")
    print("donor_compare: %d/%d properties AGREE -> %s"
          % (agrees, len(result["properties"]), argv[2]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
