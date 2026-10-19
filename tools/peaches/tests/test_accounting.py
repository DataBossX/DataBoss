"""Synthetic, offline regressions for complete page and field accounting."""
import csv
import json
import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl
from peaches import boundaries, inventory, vision_queue, wb


def write_csv(path, rows, fields=None):
    with open(path, "w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def make_pdf(path, pages, text=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    # A visible solid pixel makes these real image-only pages, not text pages.
    pixel = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 2, 2))
    pixel.clear_with(128)
    picture = pixel.tobytes("png")
    for number in range(pages):
        page = doc.new_page(width=36, height=36)
        if text:
            page.insert_text((2, 12), "test", fontsize=6)
        else:
            page.insert_image(fitz.Rect(2, 2, 20, 20), stream=picture)
    doc.save(path)
    doc.close()


class AccountingFixture(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="page_accounting_")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.sources = self.root / "sources"
        self.sources.mkdir()
        self.out = self.root / "out"
        self.out.mkdir()
        self.inv = self.out / "inventory.csv"
        self.bounds = self.out / "boundaries.csv"

    def measure(self):
        inventory.run_inventory(str(self.sources), str(self.out))
        boundaries.run_boundaries(str(self.sources), str(self.bounds))

class PageAccountingTests(AccountingFixture):
    def test_five_scanned_pdfs_account_for_all_768_pages(self):
        for index, pages in enumerate((195, 79, 211, 195, 88), 1):
            make_pdf(self.sources / ("scan_%d.pdf" % index), pages)
        self.measure()
        queue = vision_queue.build_queue(self.inv, self.bounds)
        self.assertEqual(len(queue), 768)
        self.assertEqual(len({(r["source_sha256"], r["page"]) for r in queue}), 768)
        self.assertTrue(all(r["status"] == "UNREVIEWED" for r in queue))
        self.assertTrue(all(r["page_id"] == "%s:%s" % (r["source_sha256"], r["page"]) for r in queue))

    def test_aliases_deduplicate_content_and_same_names_do_not_collide(self):
        make_pdf(self.sources / "a" / "scan.pdf", 2)
        alias = self.sources / "b" / "copy.pdf"
        alias.parent.mkdir()
        shutil.copyfile(self.sources / "a" / "scan.pdf", alias)
        make_pdf(self.sources / "c" / "scan.pdf", 3, text=True)
        self.measure()
        queue = vision_queue.build_queue(self.inv, self.bounds)
        self.assertEqual(len(queue), 5)
        self.assertEqual(len({r["source_sha256"] for r in queue}), 2)
        self.assertEqual(sorted(len(json.loads(r["source_paths"])) for r in queue), [1, 1, 1, 2, 2])

    def test_unmeasured_pdf_records_an_issue_without_phantom_page(self):
        write_csv(self.inv, [{"path":"bad.pdf", "name":"bad.pdf", "ext":".pdf",
                             "bytes":"4", "sha256":"a"*64, "pages":"",
                             "native_text_pages":"", "error":"unreadable"}])
        write_csv(self.bounds, [], boundaries.FIELDS)
        issues = []
        self.assertEqual(vision_queue.build_queue(self.inv, self.bounds, issues=issues), [])
        self.assertEqual(issues[0]["code"], "UNMEASURED_PDF")
        with self.assertRaises(ValueError):
            vision_queue.build_queue(self.inv, self.bounds)

    def test_prior_visual_receipt_is_reused_but_ocr_is_never_visual(self):
        make_pdf(self.sources / "scan.pdf", 2)
        self.measure()
        sha = inventory.sha256_file(self.sources / "scan.pdf")
        review = self.out / "reviews.csv"
        common = {"source_sha256": sha, "status":"VISUALLY_REVIEWED",
                  "reviewer":"fixture-reviewer", "reviewed_at":"2026-01-01T00:00:00Z",
                  "evidence_ref":"synthetic-test-receipt", "disposition":"continuation"}
        write_csv(review, [
            dict(common, page=1, review_method="HUMAN_VISION"),
            dict(common, page=2, review_method="OCR_ONLY")])
        queue = vision_queue.build_queue(self.inv, self.bounds, review_csv=review)
        self.assertEqual(queue[0]["status"], "VISUALLY_REVIEWED")
        self.assertEqual(queue[1]["status"], "UNREVIEWED")
        self.assertEqual(queue[1]["review_method"], "")


class DatabaseAccountingTests(AccountingFixture):
    def setUp(self):
        super().setUp()
        # Imported lazily so queue regressions expose their behavioral failures.
        from peaches import accounting
        self.accounting = accounting
        make_pdf(self.sources / "scan.pdf", 2)
        self.measure()
        self.book = self.sources / "federal.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Index"
        for col, title in enumerate(wb.FEDERAL_HEADERS, 1):
            sheet.cell(9, col, title)
        values = ["Lease", "UNKNOWN", "UNKNOWN party", "", 1, "", "1/1/1901", "", "", 1]
        for col, value in enumerate(values, 1):
            sheet.cell(10, col, value)
        book.save(self.book)
        book.close()
        self.db = self.out / "accounting.sqlite"

    def run_import(self, **kwargs):
        return self.accounting.import_accounting(
            self.db, self.inv, self.bounds, self.sources,
            workbook_paths=[self.book], **kwargs)

    def test_expected_blanks_unknowns_and_unresolved_are_separate(self):
        report = self.run_import()
        self.assertEqual(report["source_instances"], 1)
        self.assertEqual(report["source_bodies"], 1)
        self.assertEqual(report["physical_pages"], 2)
        self.assertEqual(report["visually_reviewed_pages"], 0)
        self.assertEqual(report["cell_counts"], {
            "EXPECTED_BLANK": 2, "UNRESOLVED_BLANK": 2,
            "UNKNOWN_EXACT": 1, "UNKNOWN_QUALIFIED": 1,
            "POPULATED_UNVERIFIED": 4})
        self.assertEqual(report["workbook_data_rows"], 1)
        self.assertIsNone(report["source_document_count"])

    def test_idempotent_import_preserves_review_and_all_input_bytes(self):
        before = {str(p):inventory.sha256_file(p) for p in (self.book, self.sources/"scan.pdf", self.inv)}
        first = self.run_import()
        sha = inventory.sha256_file(self.sources / "scan.pdf")
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE pages SET status='VISUALLY_REVIEWED', review_method='HUMAN_VISION', reviewer='fixture', reviewed_at='2026-01-01T00:00:00Z', evidence_ref='fixture' WHERE source_sha256=? AND page=1", (sha,))
        expected = inventory.sha256_file(self.db)
        second = self.run_import(expected_db_sha=expected)
        self.assertEqual(second["physical_pages"], first["physical_pages"])
        self.assertEqual(second["visually_reviewed_pages"], 1)
        after = {str(p):inventory.sha256_file(p) for p in (self.book, self.sources/"scan.pdf", self.inv)}
        self.assertEqual(before, after)

    def test_writer_lock_and_stale_digest_refuse_without_modifying_database(self):
        self.run_import()
        original = inventory.sha256_file(self.db)
        with self.assertRaises(ValueError):
            self.run_import()
        with self.assertRaises(ValueError):
            self.run_import(expected_db_sha="0"*64)
        lock = Path(str(self.db) + ".writer.lock")
        lock.write_text("other writer")
        with self.assertRaises(FileExistsError):
            self.run_import(expected_db_sha=original)
        self.assertEqual(lock.read_text(), "other writer")
        lock.unlink()
        self.assertEqual(inventory.sha256_file(self.db), original)

    def test_conflicting_receipt_does_not_replace_prior_review(self):
        sha = inventory.sha256_file(self.sources / "scan.pdf")
        review = self.out / "reviews.csv"
        receipt = {"source_sha256":sha, "page":1, "status":"VISUALLY_REVIEWED",
                   "review_method":"HUMAN_VISION", "reviewer":"first-reviewer",
                   "reviewed_at":"2026-01-01T00:00:00Z", "evidence_ref":"first-receipt",
                   "disposition":"continuation"}
        write_csv(review, [receipt])
        self.run_import(review_csv=review)
        original = inventory.sha256_file(self.db)
        write_csv(review, [dict(receipt, reviewer="second-reviewer", evidence_ref="conflict")])
        with self.assertRaises(ValueError):
            self.run_import(expected_db_sha=original, review_csv=review)
        self.assertEqual(inventory.sha256_file(self.db), original)

    def test_changed_source_fails_before_database_replacement(self):
        self.run_import()
        original = inventory.sha256_file(self.db)
        with open(self.sources / "scan.pdf", "ab") as stream:
            stream.write(b"changed")
        with self.assertRaises(ValueError):
            self.run_import(expected_db_sha=original)
        self.assertEqual(inventory.sha256_file(self.db), original)


if __name__ == "__main__":
    unittest.main()
