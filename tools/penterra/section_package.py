"""Build Section 13/15-format Penterra packages without inventing title facts."""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from tools.penterra.document_type_purity import merge_legal, split_document_type

WORK_DATE = "9/22/2026"
WORK_DATE_LONG = "September 22, 2026"
WORK_DATE_ISO = "2026-09-22"
CLIENT = "Anschutz Exploration Corporation"
PROJECT = "Abstract 4576 Overtime"
INDEXER = "Ryan Gille"
COUNTY = "Campbell"
TOWNSHIP = "T45N-R76W"

COUNTY_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "Doc No",
    "Book-Page",
    "Date of Doc",
    "Rec Date",
    "Legal Description",
    "Comments",
]
FEDERAL_HEADERS = [
    "Document Type",
    "Grantor",
    "Grantee",
    "File No",
    "Page",
    "Date of Doc",
    "Rec Date",
    "Comments",
    "Legal Description",
    "Part",
]

# Saved-byte OCR sentence fragments identified on P12 R3. These are not parties.
P12_OCR_PARTY_ROWS = frozenset(
    {72, 82, 88, 95, 98, 121, 130, 136, 148, 170, 176, 177, 189, 190, 191}
)

CHECKLIST_WIDTHS = {
    1: 14.0,
    2: 18.0,
    3: 29.28515625,
    4: 30.140625,
    5: 38.0,
    6: 32.7109375,
    7: 28.0,
    8: 18.0,
    9: 17.7109375,
    10: 18.85546875,
    11: 15.85546875,
}
CHECKLIST_HEIGHTS = {2: 45.0, 3: 30.0, 5: 30.0, 6: 30.0, 8: 30.0, 9: 90.0}


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value.year}"
    text = str(value).strip()
    if text.endswith(" 00:00:00") and text[:10].count("-") == 2:
        year, month, day = text[:10].split("-")
        return f"{int(month)}/{int(day)}/{int(year)}"
    return text


def nonempty(value: Any) -> bool:
    return cell_text(value) != ""


def is_ocr_fragment(value: Any) -> bool:
    text = cell_text(value)
    if not text:
        return False
    lowered = text.lower()
    markers = (
        "whether now owned",
        "agrees to warrant",
        "rights under the lease",
        "reserves and excepts",
        "and mortgagee",
        "failure to comply",
        "or its successors and assigns, in, to, and under",
        "have agreed to amend",
        "predecessors in title",
        "presently existing and hereafter",
        "oil. gas, casinghead",
        "promissory note",
        "assigned rights of way",
        "comlli",
        "article",
        "does hereby release",
    )
    return any(marker in lowered for marker in markers)


def clean_party(value: Any, row_number: int, listed_rows: frozenset[int]) -> str:
    text = cell_text(value)
    if not text:
        return ""
    if row_number in listed_rows or is_ocr_fragment(text):
        return ""
    return text


def load_xlsx_rows(path: Path) -> list[list[Any]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    rows: list[list[Any]] = []
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=max(sheet.max_column or 1, 10), values_only=True):
        rows.append(list(row))
    return rows


def data_rows(rows: Sequence[Sequence[Any]], start: int = 10) -> list[list[Any]]:
    out: list[list[Any]] = []
    for index, row in enumerate(rows, start=1):
        if index < start:
            continue
        if any(nonempty(cell) for cell in row):
            out.append(list(row))
    return out


def write_ods(path: Path, header: dict[str, str], columns: Sequence[str], records: Sequence[Sequence[Any]]) -> None:
    """Write a 7-row header + blank row 8 + header row 9 + data from row 10."""
    cells: list[list[str]] = [
        ["County:", header.get("county", COUNTY)],
        ["Lands:", header["lands"]],
        ["Date:", header.get("date", WORK_DATE)],
        ["Starting Date:", header.get("starting_date", "Inception")],
        ["Date Posted Thru:", header.get("posted_thru", "")],
        ["Indexed By:", header.get("indexed_by", INDEXER)],
        ["Project:", header.get("project", PROJECT)],
        [""],
        list(columns),
    ]
    for record in records:
        cells.append([cell_text(record[i]) if i < len(record) else "" for i in range(len(columns))])

    def xml_cell(text: str) -> str:
        if text == "":
            return '<table:table-cell/>'
        return (
            f'<table:table-cell office:value-type="string">'
            f'<text:p>{escape(text)}</text:p></table:table-cell>'
        )

    row_xml = []
    for row in cells:
        row_xml.append(
            "<table:table-row>"
            + "".join(xml_cell(value) for value in row)
            + "</table:table-row>"
        )
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:body>
  <office:spreadsheet>
   <table:table table:name="Index">
    {''.join(row_xml)}
   </table:table>
  </office:spreadsheet>
 </office:body>
</office:document-content>
"""
    styles = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:automatic-styles>
  <style:page-layout style:name="pm1">
   <style:page-layout-properties fo:page-width="8.5in" fo:page-height="11in"
    fo:margin-top="0.3in" fo:margin-bottom="0.3in" fo:margin-left="0.7in"
    fo:margin-right="0.7in" style:print-orientation="portrait" style:scale-to="100%"/>
  </style:page-layout>
 </office:automatic-styles>
 <office:master-styles>
  <style:master-page style:name="Default" style:page-layout-name="pm1"/>
 </office:master-styles>
</office:document-styles>
"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>
</manifest:manifest>
"""
    meta = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0">
 <office:meta><meta:generator>DataBossX Penterra package builder</meta:generator></office:meta>
</office:document-meta>
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.writestr("content.xml", content)
        archive.writestr("styles.xml", styles)
        archive.writestr("meta.xml", meta)


def write_checklist(path: Path, answers: dict[str, Sequence[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Abstract checklist"
    blocks = (
        (2, answers["row2"]),
        (3, answers["row3"]),
        (5, answers["row5"]),
        (6, answers["row6"]),
        (8, answers["row8"]),
        (9, answers["row9"]),
    )
    for row_number, values in blocks:
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
    for column, width in CHECKLIST_WIDTHS.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_number, height in CHECKLIST_HEIGHTS.items():
        sheet.row_dimensions[row_number].height = height
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins = PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75, header=0.3, footer=0.3)
    sheet.print_area = "A1:K9"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_letter_xlsx(path: Path, header: dict[str, str], columns: Sequence[str], records: Sequence[Sequence[Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Index"
    meta = [
        ("County:", header.get("county", COUNTY)),
        ("Lands:", header["lands"]),
        ("Date:", header.get("date", WORK_DATE)),
        ("Starting Date:", header.get("starting_date", "Inception")),
        ("Date Posted Thru:", header.get("posted_thru", "")),
        ("Indexed By:", header.get("indexed_by", INDEXER)),
        ("Project:", header.get("project", PROJECT)),
    ]
    for index, (label, value) in enumerate(meta, start=1):
        sheet.cell(index, 1, label)
        sheet.cell(index, 2, value)
    for column, name in enumerate(columns, start=1):
        sheet.cell(9, column, name)
    for offset, record in enumerate(records):
        for column in range(1, len(columns) + 1):
            sheet.cell(10 + offset, column, cell_text(record[column - 1]) if column - 1 < len(record) else "")
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:8"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_cert_docx(path: Path, paragraphs: Sequence[str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_cert_pdf(path: Path, paragraphs: Sequence[str]) -> None:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        topMargin=0.85 * inch,
        bottomMargin=0.85 * inch,
    )
    style = ParagraphStyle(
        "cert",
        fontName="Times-Roman",
        fontSize=11,
        leading=14,
        spaceAfter=8,
    )
    flow = []
    for text in paragraphs:
        flow.append(Paragraph(escape(text).replace("\n", "<br/>"), style))
        flow.append(Spacer(1, 4))
    document.build(flow)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buffer.getvalue())


def zip_members(path: Path, members: dict[str, Path]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, source in members.items():
            archive.write(source, arcname=name)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, default=str) + "\n")


def purify_county_record(row_number: int, record: Sequence[Any], ocr_rows: frozenset[int]) -> list[str]:
    padded = list(record) + [""] * (9 - len(record))
    raw_type = cell_text(padded[0])
    clean_type, qualification = split_document_type(raw_type)
    legal = merge_legal(cell_text(padded[7]), qualification)
    return [
        clean_type,
        clean_party(padded[1], row_number, ocr_rows),
        clean_party(padded[2], row_number, ocr_rows),
        cell_text(padded[3]),
        cell_text(padded[4]),
        cell_text(padded[5]),
        cell_text(padded[6]),
        legal,
        cell_text(padded[8]),
    ]


def count_populated(records: Sequence[Sequence[Any]], columns: int) -> dict[str, int]:
    populated = {index: 0 for index in range(columns)}
    for record in records:
        for index in range(columns):
            if index < len(record) and nonempty(record[index]):
                populated[index] += 1
    return populated
